"""Assemble the appeal report deck from a loaded study."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.services.reportgen.charts import add_grouped_lift_chart, add_lift_chart, wrap_label
from app.services.reportgen.composite import PACK_ASPECT, composite_pack, contain_size, element_preview, file_aspect, gallery_frame
from app.services.reportgen.study import (
    Study,
    category_story,
    classification_story,
    cover_headline,
    insight_cards,
    legal_picks,
    load_study,
    strongest_silos,
)
from app.services.reportgen.theme import (
    BLUSH,
    BLUSH_CARD,
    BLUSH_MID,
    FONT,
    GRAY,
    NAVY,
    PINK,
    ROSE,
    SLIDE_H,
    SLIDE_W,
    TEAL,
    WHITE,
    add_footer,
    add_rect,
    add_text,
    apply_leading,
    bar_color,
    fitted_size,
    fmt_lift,
    fmt_pct,
    lift_color,
    set_background,
    style_cell,
)

ASSETS = Path(__file__).resolve().parent / "template_assets"
TRACK = RGBColor(0xF4, 0xF0, 0xF2)
RGB_AXIS = NAVY


def build_report(
    json_path: str | Path,
    output_path: str | Path,
    download_images: bool = True,
    logo_path: str | Path | None = None,
    study_type: str | None = None,
    cache_dir: str | Path | None = None,
    prepared_by: str | None = None,
) -> Path:
    _FOOTER_NOTES.clear()
    study = load_study(
        json_path,
        cache_dir=cache_dir,
        download_images=download_images,
        study_type=study_type,
    )
    if logo_path:
        study.brand_logo = Path(logo_path)
    if prepared_by:
        study.prepared_by = prepared_by
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    presentation.core_properties.title = study.title
    presentation.core_properties.subject = "Design Element Appeal Report"
    presentation.core_properties.category = study.model_name

    slides = []
    slides.append(("cover", _cover(presentation, study)))
    slides.append(("divider", _divider(presentation, "Inputs")))
    slides.append(("content", _snapshot(presentation, study)))
    slides.append(("content", _method(presentation, study)))
    for slide in _inventory(presentation, study):
        slides.append(("content", slide))
    slides.append(("content", _deliverables(presentation, study)))
    slides.append(("content", _reading(presentation, study)))
    slides.append(("divider", _divider(presentation, "Key Findings")))
    slides.append(("content", _insights(presentation, study)))
    slides.append(("content", _base_appeal(presentation, study)))
    slides.append(("divider", _divider(presentation, "Combination Winners")))
    slides.append(("content", _build(presentation, study)))
    slides.append(("content", _build_low(presentation, study)))
    slides.append(("content", _build_compare(presentation, study)))
    slides.append(("divider", _divider(presentation, "Category Winners")))
    slides.append(("content", _roadmap(presentation, study)))
    for slide in _category_winner_slides(presentation, study):
        slides.append(("content", slide))
    slides.append(("divider", _divider(presentation, "Element Winners")))
    for category in study.categories:
        for slide in _category_slides(presentation, study, category):
            slides.append(("content", slide))
    slides.append(("divider", _divider(presentation, "Appendix")))
    slides.append(("content", _appendix_intro(presentation, study)))
    for slide in _lift_tables(presentation, study):
        slides.append(("content", slide))
    slides.append(("divider", _divider(presentation, "Element Assets")))
    for slide in _asset_slides(presentation, study):
        slides.append(("content", slide))
    slides.append(("cover", _thanks(presentation, study)))

    total = len(slides)
    for index, (kind, slide) in enumerate(slides, start=1):
        if kind == "content":
            add_footer(slide, _FOOTER_NOTES.get(id(slide)) or _footer_note(study), index, total)

    presentation.save(str(destination))
    return destination


def _blank(presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


_FOOTER_NOTES: dict[int, str] = {}


def _footer_note(study: Study) -> str:
    sample = f"n={study.respondents} respondents" if study.respondents else f"n={study.base_size}"
    return (
        f"{sample}.  {study.model_name} model ({study.rating_bands}): "
        "element lifts show movement above or below the model baseline."
    )


def _visual(study: Study) -> str:
    """How elements are drawn. Hybrid stays on the layer layout until its data exists."""
    kind = (study.study_type or "layer").lower()
    if kind == "grid":
        return "grid"
    if kind == "text":
        return "text"
    return "layer"


_PREPARED_BY_DEFAULT = "J Brown Fitterman | jbrown@mindgenomicsassociates.com"


def _prepared_by_line(study: Study) -> str:
    credit = (study.prepared_by or "").strip() or _PREPARED_BY_DEFAULT
    return f"Prepared by: {credit}"


def _cover_marks(slide, study: Study) -> str:
    """Left-column lockup shared by the cover and the thank-you slide. MGA is always present."""
    mga = ASSETS / "cover_Picture_17.png"
    if study.brand_logo and study.brand_logo.exists():
        _place_fit(slide, study.brand_logo, 0.55, 3.02, 3.45, 0.50)
        if mga.exists():
            slide.shapes.add_picture(str(mga), Inches(1.377), Inches(3.915), Inches(1.795), Inches(1.795))
        return f"{study.title} Appeal Report"
    rexona = ASSETS / "cover_Picture_15.png"
    if rexona.exists() and "rexona" in study.title.lower():
        slide.shapes.add_picture(str(rexona), Inches(1.377), Inches(2.120), Inches(1.795), Inches(1.496))
        if mga.exists():
            slide.shapes.add_picture(str(mga), Inches(1.377), Inches(3.915), Inches(1.795), Inches(1.795))
        return "Rexona x MGA Design Element Appeal Report"
    if mga.exists():
        slide.shapes.add_picture(str(mga), Inches(1.377), Inches(3.915), Inches(1.795), Inches(1.795))
    return f"{study.title} Appeal Report"


def _cover(presentation, study: Study):
    """Same cover as the Rexona x MGA appeal report: marks left, title block right."""
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    report_name = _cover_marks(slide, study)

    _hairline(slide, 4.2464, 1.3478, 4.2464, 1.3478 + 4.7826)
    _hairline(slide, 4.5996, 3.75, 4.5996 + 6.8551, 3.75)

    headline = cover_headline(study)
    headline_size = 28 if len(headline) > 70 else 32 if len(headline) > 40 else 36
    headline_height = 1.55 if len(headline) > 40 else 1.313
    add_text(
        slide,
        headline,
        4.600, 1.975, 8.481, headline_height,
        size=headline_size, bold=True, color=NAVY,
    )
    when = study.launched_label or ""
    add_text(
        slide,
        f"{report_name}  |  {when}".strip(" |"),
        4.600, 4.244, 7.2, 0.37,
        size=16, color=NAVY,
    )
    add_text(
        slide,
        _prepared_by_line(study),
        4.600, 4.628, 7.64, 0.37,
        size=16, color=NAVY,
    )
    add_text(slide, "Mind Genomics Associates, Inc", 4.600, 5.012, 4.2, 0.37, size=16, color=NAVY)
    add_text(slide, study.title, 4.600, 5.711, 8.4, 0.55, size=16, color=NAVY)
    return slide


def _hairline(slide, x1, y1, x2, y2, color: RGBColor = RGBColor(0xF2, 0xC6, 0xCF), width: float = 0.75):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def _divider(presentation, title: str):
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    add_rect(slide, 0, 0, 0.18, SLIDE_H, PINK)
    add_text(slide, title, 0.85, 2.85, 11, 1.1, size=54, bold=True, color=NAVY)
    add_rect(slide, 0.9, 4.15, 1.4, 0.06, TEAL)
    return slide


def _snapshot(presentation, study: Study):
    """Inputs opener: one card, filled from the study file. No country codes."""
    ink = RGBColor(0x17, 0x22, 0x38)
    card_fill = RGBColor(0xFB, 0xEF, 0xF2)
    rail = RGBColor(0xF3, 0xC9, 0xD4)

    slide = _blank(presentation)
    set_background(slide, WHITE)
    add_rect(slide, 0, 0, 0.39, SLIDE_H, rail)

    _flush(add_text(
        slide, "The study in this report",
        0.70, 0.52, 11.7, 0.55, size=34, bold=True, color=ink,
    ))
    _hairline(slide, 0.72, 1.20, 0.72 + 1.38, 1.20, rail, 1.4)
    _flush(add_text(
        slide, "Scores in the pages that follow come from this file.",
        0.72, 1.48, 10, 0.32, size=15.5, italic=True, color=ink,
    ))

    left, top, width, height = 0.88, 2.02, 11.55, 2.30
    add_rect(slide, left, top, width, height, card_fill, rounded=True, radius=0.026)
    add_rect(slide, left, top, width, 0.14, TEAL)
    _flush(add_text(
        slide, study.title,
        left + 0.35, top + 0.40, width - 0.7, 0.85, size=22, bold=True, color=ink,
    ))
    _flush(add_text(
        slide, f"{study.silo_count} silos  ·  {study.element_count} elements",
        left + 0.35, top + 1.32, width - 0.7, 0.32, size=15, color=ink,
    ))
    _flush(add_text(
        slide, _respondent_line(study),
        left + 0.35, top + 1.72, width - 0.7, 0.30, size=14, color=TEAL,
    ))
    return slide


def _respondent_line(study: Study) -> str:
    count = study.respondents or study.base_size
    if len(study.gender_distribution) == 1:
        gender = study.gender_distribution[0][0].lower()
        return f"{count} {gender} respondents"
    return f"{count} respondents"


def _flush(box):
    frame = box.text_frame
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    for paragraph in frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
    return box


def _draw_fitted(
    slide,
    text,
    x,
    y,
    w,
    max_h,
    size,
    min_size,
    color,
    bold=False,
    max_lines=2,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    one_line_floor: float | None = None,
) -> float:
    """Draw text inside a box, shrinking the font until the heading stays in bounds.

    Returns the height actually used, so the next line can start below it.
    """
    chosen, height = fitted_size(
        text, w, max_h, size, min_size, bold=bold, max_lines=max_lines, one_line_floor=one_line_floor,
    )
    if height <= 0:
        return 0.0
    box = add_text(
        slide, text, x, y, w, height,
        size=chosen, bold=bold, color=color, align=align, anchor=anchor,
    )
    _flush(box)
    apply_leading(box, chosen)
    return height


def _heading_block(slide, eyebrow, headline, detail, x, y, w, headline_size=26, detail_size=14) -> float:
    """Eyebrow, heading, and subheading stacked so a wrap never paints over the line below."""
    _draw_fitted(slide, eyebrow, x, y, w, 0.22, size=12, min_size=9, color=TEAL, bold=True, max_lines=1)
    head_y = y + 0.24
    head_h = _draw_fitted(
        slide, headline, x, head_y, w, 0.80,
        size=headline_size, min_size=15, color=NAVY, bold=True, max_lines=2, one_line_floor=18,
    )
    detail_y = head_y + max(head_h, 0.30) + 0.08
    detail_h = _draw_fitted(
        slide, detail, x, detail_y, w, 0.52,
        size=detail_size, min_size=11, color=GRAY, max_lines=2, one_line_floor=12,
    )
    return detail_y + max(detail_h, 0.22)


def _method(presentation, study: Study):
    """Same method slide as the main report: five numbered rows under a short intro."""
    ink = RGBColor(0x17, 0x22, 0x38)
    card_fill = RGBColor(0xFB, 0xEF, 0xF2)
    rail = RGBColor(0xF3, 0xC9, 0xD4)
    counts = [len(category.elements) for category in study.categories]
    low, high = (min(counts), max(counts)) if counts else (0, 0)
    variation = f"{low}–{high}" if low != high else str(low)

    slide = _blank(presentation)
    set_background(slide, WHITE)
    add_rect(slide, 0, 0, 0.39, SLIDE_H, rail)
    _flush(add_text(slide, "How the method works", 0.72, 0.48, 11.5, 0.52, size=34, bold=True, color=ink))
    _hairline(slide, 0.72, 1.12, 0.72 + 1.38, 1.12, rail, 1.4)
    _flush(add_text(
        slide,
        "Every respondent sees their own mix of design variations, and we learn what individually moves the appeal score.",
        0.72, 1.28, 11.6, 0.32, size=15, italic=True, color=ink,
    ))

    silos = _name_list(category.label for category in study.categories)
    seen = {
        "text": ("Each respondent sees a set of statements", "Each task is a unique mix of statements, never quite the same as anyone else’s."),
        "grid": ("Each respondent sees a grid of images", "Each task shows one image from each set, never quite the same as anyone else’s."),
    }.get(_visual(study), ("Each respondent sees a set of pack designs", "Each pack is a unique combination of variations — never quite the same as anyone else’s."))
    steps = [
        ("01", "We slice the design into unique elements", f"Each element is one design choice — {silos}."),
        ("02", f"Each element has {variation} possible variations", "Like the small choices a designer would make between the options in a silo."),
        ("03", seen[0], seen[1]),
        ("04", "They rate each one", "The same rating scale is used on every task. The variation does the work."),
        ("05", "We learn what drives the rating", "The model separates the baseline from the movement each element creates, on its own."),
    ]
    row_h, gap = 0.88, 0.08
    for index, (number, title, body) in enumerate(steps):
        top = 1.72 + index * (row_h + gap)
        add_rect(slide, 0.72, top, 11.90, row_h, card_fill, rounded=True, radius=0.083)
        circle = 0.64
        cy = top + (row_h - circle) / 2
        _badge(slide, number, 0.90, cy, circle, 14)
        _flush(add_text(slide, title, 1.72, top + 0.14, 10.5, 0.30, size=16, bold=True, color=ink))
        _flush(add_text(slide, body, 1.72, top + 0.46, 10.5, 0.30, size=13, color=ink))
    return slide


def _inventory(presentation, study: Study):
    """Silo cards, matching the main report: letter, name, count, and a short sample."""
    ink = RGBColor(0x1E, 0x2D, 0x44)
    mute = RGBColor(0x62, 0x6F, 0x7B)
    card_fill = RGBColor(0xFA, 0xF7, 0xF2)
    rail = RGBColor(0xEB, 0xC6, 0xD1)
    bar = RGBColor(0x00, 0x7B, 0x7F)
    per_slide = 16
    slides = []
    chunks = [study.categories[index:index + per_slide] for index in range(0, len(study.categories), per_slide)]
    cols, card_w, card_h, gap = 4, 2.78, 1.10, 0.13
    for page, chunk in enumerate(chunks, start=1):
        slide = _blank(presentation)
        set_background(slide, WHITE)
        add_rect(slide, 0, 0, 0.42, SLIDE_H, rail)
        suffix = f"  ·  {page} of {len(chunks)}" if len(chunks) > 1 else ""
        _flush(add_text(slide, f"Design elements{suffix}", 0.70, 0.48, 10, 0.52, size=34, bold=True, color=ink))
        _flush(add_text(
            slide,
            f"{study.silo_count} design silos, {study.element_count} element variations in total.",
            0.70, 1.12, 10, 0.28, size=14, color=mute,
        ))
        for index, category in enumerate(chunk):
            left = 0.70 + (index % cols) * (card_w + gap)
            top = 1.58 + (index // cols) * (card_h + gap)
            add_rect(slide, left, top, card_w, card_h, card_fill)
            add_rect(slide, left, top, 0.10, card_h, bar)
            _flush(add_text(
                slide, category.code, left + 0.18, top + 0.10, 0.42, 0.42,
                size=26, bold=True, color=bar,
            ))
            name = _flush(add_text(
                slide, category.label, left + 0.62, top + 0.16, 1.68, 0.28,
                size=13, bold=True, color=ink,
            ))
            name.text_frame.word_wrap = False
            _flush(add_text(
                slide, str(len(category.elements)), left + card_w - 0.48, top + 0.16, 0.36, 0.28,
                size=14, bold=True, color=mute, align=PP_ALIGN.RIGHT,
            ))
            _flush(add_text(
                slide, _option_sample(category, study), left + 0.18, top + 0.64, card_w - 0.32, 0.36,
                size=11, color=mute,
            ))
        _flush(add_text(
            slide,
            "Each card shows the silo letter, the element count, and a sample of the options.",
            0.70, 6.62, 11.5, 0.24, size=11, color=mute,
        ))
        slides.append(slide)
    return slides


def _option_sample(category, study: Study | None = None, limit: int = 42) -> str:
    if study is not None and _visual(study) == "text":
        count = len(category.elements)
        return f"{count} statement" if count == 1 else f"{count} statements"
    labels = [element.label for element in category.elements]
    shown = []
    for label in labels:
        trial = shown + [label]
        rest = len(labels) - len(trial)
        text = ", ".join(trial) if rest == 0 else ", ".join(trial) + f", +{rest}"
        if len(text) > limit and shown:
            break
        shown = trial
    rest = len(labels) - len(shown)
    if not shown:
        return ""
    if rest:
        return ", ".join(shown) + f", +{rest}"
    return ", ".join(shown)


def _deliverables(presentation, study: Study):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, "WHAT YOU GET BACK")
    question = {
        "text": "Which statements drive appeal?",
        "grid": "Which images drive appeal?",
    }.get(_visual(study), "What design elements drive pack appeal?")
    add_text(slide, question, 0.62, 0.72, 12, 0.7, size=30, bold=True, color=NAVY)
    base_body = (
        f"The {study.model_name} intercept. Starting appeal before a single statement is added."
        if _visual(study) == "text"
        else f"The {study.model_name} intercept. The pack’s starting appeal before a single element is added."
    )
    cards = [
        ("Combination winners", "The strongest option in each group, shown together, against a weaker combination of the same groups."),
        ("Category winners", "The winning element in every category, so the same read works one category at a time."),
        ("Element winners", "Every element inside its category, with the score used in the charts."),
        ("Base appeal", base_body),
    ]
    for index, (title, body) in enumerate(cards):
        left = 0.62 + (index % 2) * 6.3
        top = 1.75 + (index // 2) * 2.35
        add_rect(slide, left, top, 6.0, 2.1, BLUSH_CARD, rounded=True)
        add_text(slide, f"0{index + 1}", left + 0.3, top + 0.28, 1.2, 0.4, size=18, bold=True, color=TEAL)
        add_text(slide, title, left + 0.3, top + 0.75, 5.4, 0.4, size=20, bold=True, color=NAVY)
        add_text(slide, body, left + 0.3, top + 1.25, 5.4, 0.65, size=14, color=GRAY)
    return slide


def _reading(presentation, study: Study):
    """Numbered reading guide and coefficient table, as on the main report."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    rail = RGBColor(0xF2, 0xC6, 0xCF)
    add_rect(slide, 0, 0, 0.37, SLIDE_H, rail)
    _flush(add_text(
        slide,
        "Reading The Scores – What The Coefficients Mean",
        0.53, 0.38, 12.2, 0.55,
        size=28, bold=True, color=NAVY,
    ))
    _hairline(slide, 0.66, 0.98, 2.16, 0.98, rail, 1.25)

    _number(slide, "1", 1.68, 1.30)
    _flush(add_text(slide, "Element Lift Scores", 2.10, 1.22, 8, 0.38, size=18, bold=True, color=NAVY))
    _flush(add_text(
        slide,
        "Element lifts show the estimated movement each design choice creates above or below the model baseline.",
        2.10, 1.64, 7.4, 0.7, size=14, color=NAVY,
    ))

    _number(slide, "2", 1.67, 4.00)
    _flush(add_text(slide, "Explore Combinations", 2.10, 3.92, 4.2, 0.38, size=18, bold=True, color=NAVY))
    _flush(add_text(
        slide,
        "These scores can be combined to explore multiple design scenarios.",
        2.10, 4.34, 4.4, 0.7, size=14, color=NAVY,
    ))

    _number(slide, "3", 7.39, 4.00)
    _flush(add_text(slide, "Reading The Scale", 7.82, 3.92, 4.4, 0.38, size=18, bold=True, color=NAVY))
    _flush(add_text(slide, "Coefficient Interpretation", 7.82, 4.34, 4.4, 0.32, size=14, color=NAVY))

    rule = RGBColor(0x27, 0x31, 0x42)
    _hairline(slide, 1.90, 2.15, 1.90, 3.50, rule, 2)
    _hairline(slide, 1.90, 3.50, 6.06, 3.50, rule, 2)
    _hairline(slide, 6.06, 3.50, 6.06, 4.85, rule, 2)

    strong = RGBColor(0xDE, 0xF1, 0xD3)
    rows = [
        ("Coefficient Scale", "Interpretation", True, WHITE),
        ("+10 or higher", "Strong lift", False, strong),
        ("+5 to +9.9", "Moderate lift", False, WHITE),
        ("+2 to +4.9", "Small lift", False, WHITE),
        ("Below +2", "Minimal / weak lift", False, WHITE),
    ]
    frame = slide.shapes.add_table(len(rows), 2, Inches(7.81), Inches(4.78), Inches(4.42), Inches(2.01))
    table = frame.table
    table.columns[0].width = Inches(2.01)
    table.columns[1].width = Inches(2.41)
    for index, (left, right, header, band) in enumerate(rows):
        table.rows[index].height = Inches(0.40)
        for column, text in enumerate((left, right)):
            style_cell(
                table.cell(index, column), text,
                size=14, bold=header, color=NAVY,
                fill=band if column == 0 else WHITE,
                border="273142", border_width="12700",
            )
    return slide


def _badge(slide, label, x, y, size, font):
    """Navy circle with the number centered in the shape, not in a separate box."""
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    dot.fill.solid()
    dot.fill.fore_color.rgb = NAVY
    dot.line.fill.background()
    _shape_label(dot, label, font, WHITE, bold=True, wrap=False)


def _shape_label(shape, text, size, color, bold=False, wrap=True, align=PP_ALIGN.CENTER):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = wrap
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    inset = Inches(0.08) if wrap else Emu(0)
    frame.margin_left = inset
    frame.margin_right = inset
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def _number(slide, label, x, y):
    _badge(slide, label, x, y, 0.32, 11)


def _insights(presentation, study: Study):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, "KEY FINDINGS")
    add_text(slide, "Where appeal moves", 0.62, 0.72, 12, 0.48, size=30, bold=True, color=NAVY)
    add_text(
        slide,
        "The read starts with the winning combination, then the winner in each category, then every element in that category.",
        0.62, 1.25, 12, 0.35, size=15, color=GRAY,
    )
    for index, (title, body) in enumerate(insight_cards(study)):
        top = 1.8 + index * 1.2
        add_rect(slide, 0.62, top, 12.1, 1.08, BLUSH_CARD if index % 2 == 0 else BLUSH, rounded=True)
        add_text(slide, f"{index + 1:02d}", 0.82, top, 0.7, 1.08, size=20, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, 1.7, top + 0.12, 10.6, 0.36, size=16, bold=True, color=NAVY)
        add_text(slide, body, 1.7, top + 0.5, 10.6, 0.46, size=13, color=GRAY)
    return slide


def _base_appeal(presentation, study: Study):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, "MODEL OVERVIEW: BASE APPEAL")
    add_text(slide, "Base appeal before element lift", 0.62, 0.72, 12, 0.48, size=30, bold=True, color=NAVY)
    add_text(
        slide,
        f"Base appeal is this model’s starting {study.model_name} score ({study.rating_bands}) before individual design elements add or subtract lift.",
        0.62, 1.28, 12, 0.45, size=15, color=GRAY,
    )
    add_rect(slide, 0.62, 2.05, 6.3, 4.4, BLUSH_MID, rounded=True)
    add_text(slide, "BASE APPEAL", 0.95, 2.35, 5.6, 0.3, size=13, bold=True, color=TEAL)
    add_text(slide, study.title, 0.95, 2.8, 5.6, 0.9, size=16, color=NAVY)
    if study.base_appeal is None:
        add_text(slide, "No intercept", 0.95, 3.85, 5.6, 0.8, size=36, bold=True, color=NAVY)
        add_text(slide, "Element scores are the lift on their own", 0.95, 4.75, 5.6, 0.55, size=14, color=GRAY)
    else:
        add_text(slide, fmt_pct(study.base_appeal), 0.95, 3.75, 5.6, 1.15, size=72, bold=True, color=NAVY)
        add_text(slide, "Model baseline  ·  intercept", 0.95, 5.15, 5.6, 0.35, size=14, color=GRAY)
    add_text(slide, f"Base size {study.base_size}", 0.95, 5.55, 5.6, 0.3, size=14, bold=True, color=NAVY)

    add_rect(slide, 7.2, 2.05, 5.5, 4.4, NAVY, rounded=True)
    add_text(slide, "HOW TO USE IT", 7.55, 2.35, 4.9, 0.3, size=13, bold=True, color=PINK)
    add_text(
        slide,
        "Each element chart is movement around this number, not a new total by itself.",
        7.55, 2.85, 4.85, 1.0, size=16, color=WHITE,
    )
    add_text(
        slide,
        "A build-up adds the baseline and one lift from each silo. That sum is a model scenario. It is not capped at 0 or 100.",
        7.55, 4.05, 4.85, 1.3, size=16, color=WHITE,
    )
    return slide


def _roadmap(presentation, study: Study):
    """Four design areas, spaced like the main report's roadmap."""
    ink = RGBColor(0x1F, 0x2A, 0x44)
    mute = RGBColor(0x65, 0x74, 0x8B)
    blush = RGBColor(0xF7, 0xEC, 0xF1)
    pink = RGBColor(0xE8, 0xCF, 0xDA)
    edge = RGBColor(0xE4, 0xE9, 0xF0)
    leaders = strongest_silos(study, 4)
    areas = "four" if len(leaders) == 4 else str(len(leaders))

    slide = _blank(presentation)
    set_background(slide, WHITE)
    add_rect(slide, 0, 0, 0.42, SLIDE_H, pink)

    _flush(add_text(slide, "CATEGORY WINNERS", 0.78, 0.36, 6.4, 0.24, size=12, bold=True, color=ink))
    _flush(add_text(
        slide,
        f"Appeal moves most clearly\nin {areas} design areas",
        0.78, 0.68, 6.7, 1.15, size=30, bold=True, color=ink,
    ))
    add_rect(slide, 7.85, 0.62, 4.85, 1.22, blush, rounded=True, radius=0.1)
    _flush(add_text(
        slide,
        "These are the silos where one element, or the gap inside the silo, moves appeal the most.",
        8.08, 0.78, 4.42, 0.9, size=14, color=ink, anchor=MSO_ANCHOR.MIDDLE,
    ))

    _flush(add_text(slide, "ROADMAP FOR THE READ", 0.78, 2.15, 6, 0.24, size=11, bold=True, color=mute))

    card_w, card_h, gap = 2.92, 2.55, 0.16
    origin_x, origin_y = 0.78, 2.48
    for index, category in enumerate(leaders):
        left = origin_x + index * (card_w + gap)
        add_rect(slide, left, origin_y, card_w, card_h, WHITE, rounded=True, radius=0.06, line=edge)
        _flush(add_text(
            slide, f"{index + 1:02d}", left + 0.18, origin_y + 0.12, 2.5, 0.22,
            size=12, bold=True, color=TEAL,
        ))
        _flush(add_text(
            slide, category.label, left + 0.18, origin_y + 0.36, 2.56, 0.52,
            size=16, bold=True, color=ink,
        ))
        body = _flush(add_text(
            slide, _roadmap_line(category, study), left + 0.18, origin_y + 0.96, 2.56, 1.42,
            size=12, color=mute,
        ))
        body.text_frame.word_wrap = True

    row_w = len(leaders) * card_w + max(0, len(leaders) - 1) * gap
    names = _name_list(category.label for category in leaders)
    bar_top = origin_y + card_h + 0.46
    add_rect(slide, origin_x, bar_top, row_w, 0.78, blush, rounded=True, radius=0.08)
    _flush(add_text(
        slide,
        f"The clearest appeal movement is concentrated in {names}.",
        origin_x, bar_top, row_w, 0.78,
        size=15, color=ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    ))
    return slide


def _roadmap_line(category, study: Study | None = None) -> str:
    best, worst = category.best, category.worst
    if study is not None and _visual(study) == "text":
        return f"{best.label}\n{fmt_lift(best.lift)}"
    lead = f"{best.label} leads at {fmt_lift(best.lift)}."
    if worst.lift < 0:
        return f"{lead}\n{worst.label} pulls the other way."
    if best.lift <= 0:
        return f"{best.label} is least negative at {fmt_lift(best.lift)}."
    return f"{lead}\nLow end is {worst.label}."


def _name_list(names) -> str:
    items = list(names)
    if not items:
        return "these design areas"
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + f", and {items[-1]}"


def _category_winner_slides(presentation, study: Study):
    """The winning element in every category, strongest category first.

    One classification question is shown on the same card. Several questions
    stay on the element slides, where each question already has its own chart.
    """
    ranked = sorted(study.categories, key=lambda category: category.best.lift, reverse=True)
    question = study.classifications[0] if len(study.classifications) == 1 else None
    if _visual(study) == "text":
        per_page = 5
    elif question is not None:
        per_page = 4
    else:
        per_page = 5
    in_pack = {element.name for element in _pack_picks(study, "best")}
    pages = _pages(ranked, per_page)
    return [
        _category_winner_slide(presentation, study, page, index, len(pages), question, in_pack)
        for index, page in enumerate(pages, start=1)
    ]


def _category_winner_slide(presentation, study, categories, page: int, pages: int, question, in_pack: set[str]):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, f"CATEGORY WINNERS{_page_mark(page, pages)}")
    add_text(slide, "The winning element in each category", 0.48, 0.52, 12.3, 0.4, size=26, bold=True, color=NAVY)
    if question is not None:
        detail = f"Ranked by the total sample. The notes are the winners for “{question.text.rstrip('?')}”."
    elif len(study.classifications) > 1:
        detail = "Ranked by coefficient. Group cuts for each category are in Element Winners."
    else:
        detail = "Ranked by coefficient. The mark notes a winner that is not on the highest-sum pack."
    add_text(slide, detail, 0.48, 0.98, 12.3, 0.32, size=14, color=GRAY)

    top, bottom, gap = 1.46, 6.82, 0.1
    count = max(len(categories), 1)
    row_h = (bottom - top - gap * (count - 1)) / count
    for index, category in enumerate(categories):
        _winner_card(
            slide, study, category, top + index * (row_h + gap), row_h, question, in_pack,
        )
    return slide


def _winner_card(slide, study, category, top, height, question, in_pack: set[str]):
    """One category: a steady image well, the winner, and the group notes."""
    x, width = 0.42, 12.5
    card = RGBColor(0xFB, 0xF6, 0xF8)
    add_rect(slide, x, top, width, height, card, rounded=True, radius=0.1)
    winner = category.best
    visual = _visual(study) != "text"
    text_x = x + 0.28
    if visual:
        well = min(height - 0.24, 1.08)
        well_x = x + 0.14
        well_y = top + (height - well) / 2
        add_rect(slide, well_x, well_y, well, well, WHITE, rounded=True, radius=0.08)
        art = _element_on_pack(study, winner) or _raw_image(winner)
        if art:
            inset = 0.08
            _place_fit(slide, art, well_x + inset, well_y + inset, well - inset * 2, well - inset * 2)
        text_x = well_x + well + 0.22

    score_w = 1.2
    text_w = x + width - score_w - 0.2 - text_x
    off_pack = bool(study.blocked) and winner.name not in in_pack
    picks = _segment_winners(category, question) if question is not None else []
    agree = len({best.label for _segment, best, _score in picks}) == 1 if picks else False

    label_w = text_w - (1.55 if off_pack else 0)
    label_y = top + 0.14 if picks else top + (height - 0.62) / 2
    _flush(add_text(
        slide, category.label.upper(), text_x, label_y, label_w, 0.24,
        size=11, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE,
    ))
    if off_pack:
        pill_w, pill_h = 1.38, 0.26
        pill_x = text_x + text_w - pill_w
        add_rect(slide, pill_x, label_y, pill_w, pill_h, RGBColor(0xFD, 0xF0, 0xF3), rounded=True, radius=0.08)
        _flush(add_text(
            slide, "Off the pack", pill_x, label_y, pill_w, pill_h,
            size=10, bold=True, color=ROSE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        ))
    _flush(add_text(
        slide, winner.label, text_x, label_y + 0.28, text_w, 0.36 if picks else 0.34,
        size=16 if _visual(study) == "text" else 18, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    ))
    if picks and agree:
        who = "Both groups agree" if len(picks) == 2 else "Every group agrees"
        _flush(add_text(
            slide, who, text_x, top + height - 0.4, text_w, 0.24,
            size=12, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
        ))
    elif picks:
        _winner_chips(slide, picks, text_x, top + height - 0.46, text_w, 0.32)
    _flush(add_text(
        slide, fmt_lift(winner.lift), x + width - score_w - 0.12, top, score_w, height,
        size=22, bold=True, color=lift_color(winner.lift), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
    ))


def _segment_winners(category, question):
    picks = []
    for segment in question.segments:
        best = max(
            category.elements,
            key=lambda element, segment=segment: float(element.segment_lifts.get(segment.key, element.lift)),
        )
        raw = best.segment_lifts.get(segment.key)
        score = float(raw) if raw is not None else best.lift
        picks.append((segment, best, score))
    return picks


def _winner_chips(slide, picks, x, y, width, height):
    """One note per group when the groups do not share a winner."""
    count = len(picks)
    gap = 0.08
    chip_w = (width - gap * (count - 1)) / count
    for index, (segment, best, score) in enumerate(picks):
        left = x + index * (chip_w + gap)
        add_rect(slide, left, y, chip_w, height, WHITE, rounded=True, radius=0.06)
        _flush(add_text(
            slide, f"{segment.label}   ·   {best.label}   {fmt_lift(score)}",
            left + 0.1, y, chip_w - 0.16, height,
            size=11, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        ))


# Nine rows keep a wrapped name clear of the next bar. The gallery is 3×3.
_FINDINGS_PER_PAGE = 9
# Four grouped rows match the classification readout: a 2×2 gallery and bars with air.
_GROUPED_PER_PAGE = 4
_ASSET_PER_PAGE = 10
_ASSET_PER_ROW = 5


def _pages(items, size: int):
    return [items[index:index + size] for index in range(0, len(items), size)] or [[]]


def _grouped_pages(items):
    """Pages of four. A leftover single row is pulled back so a page is never one bar."""
    pages = _pages(items, _GROUPED_PER_PAGE)
    if len(pages) > 1 and len(pages[-1]) == 1:
        pages[-1].insert(0, pages[-2].pop())
    return pages


def _page_mark(index: int, count: int) -> str:
    if count <= 1:
        return ""
    return f"   ·   {index} of {count}"


def _category_slides(presentation, study: Study, category):
    """One findings slide per page when the silo is longer than the chart can hold.

    Each classification question gets its own chart and legend. Questions are
    different cuts, so their answers are never mixed on one legend.
    """
    if _visual(study) == "text":
        return _text_category_slides(presentation, study, category)
    questions = list(study.classifications) or [None]
    slides = []
    for question in questions:
        if question is None:
            ordered = sorted(category.elements, key=lambda element: element.lift, reverse=True)
        else:
            ordered = sorted(
                category.elements,
                key=lambda element: max(
                    (element.segment_lifts.get(segment.key, element.lift) for segment in question.segments),
                    default=element.lift,
                ),
                reverse=True,
            )
        pages = _pages(ordered, _FINDINGS_PER_PAGE) if question is None else _grouped_pages(ordered)
        slides.extend(
            _category_slide(presentation, study, category, page, index, len(pages), question)
            for index, page in enumerate(pages, start=1)
        )
    return slides


def _category_slide(presentation, study: Study, category, elements, page: int, pages: int, question=None):
    """Chart on the left, element artwork on the right, as in the main report."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    if question is None:
        headline, detail, key = category_story(category)
        stacked = _heading_block(
            slide,
            f"ELEMENT WINNERS: {_kicker(category.label)}{_page_mark(page, pages)}",
            headline,
            detail,
            0.72, 0.26, 7.5,
            headline_size=24,
        )
        banner_y = stacked + 0.10
        banner = add_rect(slide, 0.85, banner_y, 6.9, 0.28, RGBColor(0xF4, 0xE0, 0xE7))
        _shape_label(banner, f"{category.label.upper()} PREFERENCES", 11, NAVY, bold=True, wrap=False)
        chart_y = banner_y + 0.34
        add_lift_chart(slide, elements, 0.46, chart_y, 7.72, max(2.2, 5.98 - chart_y))
        _finding_gallery(slide, study, category, elements, 8.48, 0.28, 4.42, 5.68)
        note = add_rect(slide, 0.72, 6.08, 12.15, 0.68, RGBColor(0xFC, 0xF4, 0xF8), rounded=True, radius=0.1)
        _shape_label(note, key, 12, NAVY, align=PP_ALIGN.LEFT)
        return slide

    headline, detail, key = classification_story(category, question)
    question_mark = f"  ·  {question.text}" if len(study.classifications) > 1 else ""
    # One grid for the chart column and the gallery, so the legend, card, and key bar share edges.
    left = 0.55
    gallery_w = 4.42
    gap = 0.20
    chart_w = SLIDE_W - left - gap - gallery_w - 0.46
    gallery_x = left + chart_w + gap
    column_bottom = 6.06
    stacked = _heading_block(
        slide,
        f"ELEMENT WINNERS: {_kicker(category.label)}{question_mark}{_page_mark(page, pages)}",
        headline,
        detail,
        left, 0.24, chart_w,
    )
    banner_y = stacked + 0.08
    banner = add_rect(slide, left, banner_y, chart_w, 0.28, RGBColor(0xF7, 0xE6, 0xEC))
    _shape_label(banner, category.label.upper(), 11, NAVY, bold=True, wrap=False)
    chart_y = banner_y + 0.34
    add_grouped_lift_chart(slide, elements, question.segments, left, chart_y, chart_w, column_bottom - chart_y)
    _finding_gallery(
        slide, study, category, elements, gallery_x, 0.26, gallery_w, column_bottom - 0.26, question,
    )
    key_h = 0.74 if len(key) > 180 else 0.58
    note = add_rect(
        slide, left, 6.16, gallery_x + gallery_w - left, key_h,
        RGBColor(0xFB, 0xF4, 0xF7), rounded=True, radius=0.12,
    )
    _shape_label(note, key, 12 if len(key) > 180 else 13, NAVY, align=PP_ALIGN.LEFT)
    _FOOTER_NOTES[id(slide)] = _classification_footer(question)
    return slide


def _classification_footer(question) -> str:
    """Sample sizes for the legend, plus the meaning of the starred gallery score."""
    parts = [
        f"{segment.label} basesize = {segment.base_size}" if segment.base_size else segment.label
        for segment in question.segments
    ]
    groups = "; ".join(parts)
    note = f"{groups}. Top-2 scores rounded to whole points. *Average of the groups in the legend."
    if len(note) > 150:
        note = "Top-2 scores rounded to whole points. *Average of the groups in the legend."
    return note


def _group_score(element, question):
    """Gallery number. With a classification cut this is the average of the groups."""
    if question is None:
        return fmt_lift(element.lift), lift_color(element.lift)
    values = [element.segment_lifts.get(segment.key) for segment in question.segments]
    values = [float(value) for value in values if value is not None]
    if not values:
        return fmt_lift(element.lift), lift_color(element.lift)
    # The readout rounds the group average to a whole point, same as the bars.
    average = int(round(sum(values) / len(values)))
    return fmt_lift(average) + "*", lift_color(average)


def _finding_gallery(slide, study, category, elements, x, y, w, h, question=None):
    """Element artwork, name, and score in the card beside the chart."""
    edge = RGBColor(0xE3, 0xE8, 0xEE)
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.08 if question is not None else 0.04, line=edge)
    arts = _gallery_art(study, category, elements)
    if question is not None:
        _caption_gallery(slide, elements, arts, x, y, w, h, question)
        return
    if len(elements) <= 3:
        _gallery_rows(slide, elements, arts, x, y, w, h, question)
        return
    count = len(elements)
    columns = 2 if count <= 6 else 3
    rows = (count + columns - 1) // columns
    pad = 0.14
    cell_w = (w - pad * 2) / columns
    cell_h = (h - pad * 2) / rows
    name_lines = 2 if _visual(study) == "grid" else 1
    name_h = 0.42 if name_lines == 2 else 0.28
    score_h = 0.24
    image_max_h = max(0.55, cell_h - name_h - score_h - 0.12)
    name_size = 12 if columns == 2 else 11
    for index, element in enumerate(elements):
        column = index % columns
        row = index // columns
        left = x + pad + column * cell_w
        top = y + pad + row * cell_h
        art = arts[index]
        img_w, img_h = _fit_size(art, cell_w - 0.12, image_max_h) if art else (0, 0)
        group_h = img_h + 0.06 + name_h + score_h
        origin = top + max(0, (cell_h - group_h) / 2)
        if art:
            _draw_art(slide, art, left + 0.06, origin, cell_w - 0.12, image_max_h)
        name_top = origin + img_h + 0.06
        name_box = _flush(add_text(
            slide, wrap_label(element.label, cell_w - 0.12, name_size, max_lines=name_lines),
            left + 0.02, name_top, cell_w - 0.04, name_h,
            size=name_size, color=NAVY, align=PP_ALIGN.CENTER,
        ))
        name_box.text_frame.word_wrap = name_lines > 1
        score, score_color = _group_score(element, question)
        _flush(add_text(
            slide, score, left, name_top + name_h, cell_w, score_h,
            size=14, bold=True, color=score_color, align=PP_ALIGN.CENTER,
        ))


def _caption_gallery(slide, elements, arts, x, y, w, h, question):
    """Image, then the name, then the score. A row shares one tile height and one baseline."""
    count = max(len(elements), 1)
    columns = 1 if count in (1, 3) else 2
    rows = (count + columns - 1) // columns
    pad_x, pad_y = 0.20, 0.18
    gap_x, gap_y = 0.18, 0.16
    cell_w = (w - pad_x * 2 - gap_x * (columns - 1)) / columns
    cell_h = (h - pad_y * 2 - gap_y * (rows - 1)) / rows
    long_name = any(len(element.label) > 18 for element in elements)
    name_h = 0.40 if long_name else 0.24
    score_h = 0.26
    caption = name_h + score_h + 0.06
    image_cap = max(0.7, cell_h - caption - 0.08)
    fitted = [
        _fit_size(art, cell_w - 0.08, image_cap) if art else (0.0, 0.0)
        for art in arts
    ]
    tile_h = max((height for _width, height in fitted), default=image_cap)
    tile_h = min(tile_h, image_cap)
    block_h = tile_h + caption
    name_size = 13 if columns == 1 else 12
    for index, element in enumerate(elements):
        row = index // columns
        column = index % columns
        left = x + pad_x + column * (cell_w + gap_x)
        top = y + pad_y + row * (cell_h + gap_y)
        origin = top + max(0, (cell_h - block_h) / 2)
        art = arts[index]
        if art:
            _img_w, img_h = fitted[index]
            image_top = origin + max(0, (tile_h - img_h) / 2)
            _draw_art(slide, art, left, image_top, cell_w, img_h)
        name_top = origin + tile_h + 0.06
        name_box = _flush(add_text(
            slide, wrap_label(element.label, cell_w, name_size, max_lines=2),
            left, name_top, cell_w, name_h,
            size=name_size, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        ))
        name_box.text_frame.word_wrap = long_name
        score, score_color = _group_score(element, question)
        _flush(add_text(
            slide, score, left, name_top + name_h, cell_w, score_h,
            size=15, bold=True, color=score_color, align=PP_ALIGN.CENTER,
        ))


def _gallery_rows(slide, elements, arts, x, y, w, h, question=None):
    """Large tiles down the card, for silos with only a few elements."""
    count = max(len(elements), 1)
    pad_x, pad_y = 0.2, 0.22
    inner_h = h - pad_y * 2
    gap = 0.14
    slot_h = (inner_h - gap * (count - 1)) / count
    img_max_w = (w - pad_x * 2) * 0.58
    img_max_h = slot_h - 0.06
    text_x = x + pad_x + img_max_w + 0.14
    text_w = x + w - pad_x - text_x
    for index, element in enumerate(elements):
        row_top = y + pad_y + index * (slot_h + gap)
        art = arts[index]
        img_h = 0
        if art:
            _img_w, img_h = _fit_size(art, img_max_w, img_max_h)
            frame_top = row_top + (slot_h - img_h) / 2
            _draw_art(slide, art, x + pad_x, frame_top, img_max_w, img_max_h)
        text_mid = row_top + slot_h / 2
        _flush(add_text(
            slide, element.label, text_x, text_mid - 0.34, text_w, 0.28,
            size=14, color=NAVY, anchor=MSO_ANCHOR.BOTTOM,
        ))
        score, score_color = _group_score(element, question)
        _flush(add_text(
            slide, score, text_x, text_mid - 0.02, text_w, 0.30,
            size=18, bold=True, color=score_color,
        ))


def _gallery_art(study, category, elements):
    mode = _visual(study)
    if mode == "text":
        return [None for _element in elements]
    if mode == "grid":
        return [_raw_image(element) for element in elements]
    paths = [element.image_path for element in category.elements if element.image_path]
    frame = gallery_frame(paths)
    sample = next((path for path in paths if path.exists()), None)
    aspect = file_aspect(sample) if sample else None
    arts = []
    for element in elements:
        if frame is not None and element.image_path and element.image_path.exists():
            destination = element.image_path.parent.parent / "previews_gallery_fit" / f"{element.code}.png"
            background = study.background_path if study.background_path and study.background_path.exists() else None
            arts.append(element_preview(
                background, element.image_path, destination, width=1200, frame=frame, aspect=aspect,
            ))
        else:
            arts.append(_element_thumb(study, element))
    return arts


def _raw_image(element) -> Path | None:
    path = element.image_path
    if path and Path(path).exists():
        return path
    return None


def _line_count(text: str, width_in: float, size: float) -> int:
    """Estimate wrapped lines so a statement box is tall enough for the whole sentence."""
    chars = max(12, int(width_in * 72 / (max(size, 1) * 0.64)))
    words = str(text).split()
    if not words:
        return 1
    lines = 1
    current = 0
    for word in words:
        extra = len(word) if current == 0 else current + 1 + len(word)
        if extra <= chars:
            current = extra
            continue
        lines += 1
        current = len(word)
    return lines


def _row_height(text: str, width_in: float, size: float) -> float:
    return _line_count(text, width_in, size) * (size / 72 * 1.45) + 0.16


def _text_category_slides(presentation, study: Study, category):
    """Every statement in full. A page breaks before a sentence would be clipped."""
    ordered = sorted(category.elements, key=lambda element: element.lift, reverse=True)
    text_w = 10.3
    top, bottom, size = 1.85, 6.55, 15
    pages: list[list] = []
    current: list = []
    cursor = top
    for element in ordered:
        row_h = _row_height(element.label, text_w, size)
        if current and cursor + row_h > bottom:
            pages.append(current)
            current = []
            cursor = top
        current.append(element)
        cursor += row_h + 0.08
    if current:
        pages.append(current)
    return [
        _text_category_slide(presentation, study, category, page, index, len(pages))
        for index, page in enumerate(pages, start=1)
    ]


def _text_category_slide(presentation, study: Study, category, elements, page: int, pages: int):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    best = max(category.elements, key=lambda element: element.lift)
    worst = min(category.elements, key=lambda element: element.lift)
    _flush(add_text(
        slide, f"ELEMENT WINNERS: {_kicker(category.label)}{_page_mark(page, pages)}",
        0.72, 0.28, 12.0, 0.24, size=12, bold=True, color=TEAL,
    ))
    _flush(add_text(
        slide, f"Every statement in {category.label}",
        0.72, 0.54, 12.0, 0.42, size=26, bold=True, color=NAVY,
    ))
    _flush(add_text(
        slide,
        "Each line is the full statement. The number beside it is that statement’s lift.",
        0.72, 1.05, 12.0, 0.32, size=14, color=GRAY,
    ))
    _statement_rows(slide, [(element.label, element.lift) for element in elements], 0.72, 1.85, 12.0, 6.55, size=15)
    note = (
        f"Strongest {fmt_lift(best.lift)}  ·  weakest {fmt_lift(worst.lift)}  ·  "
        f"gap {fmt_lift(category.spread)} points."
    )
    add_rect(slide, 0.72, 6.62, 12.15, 0.28, RGBColor(0xFC, 0xF4, 0xF8), rounded=True, radius=0.08)
    _flush(add_text(slide, note, 0.88, 6.62, 11.85, 0.28, size=12, color=NAVY, anchor=MSO_ANCHOR.MIDDLE))
    return slide


def _statement_rows(slide, rows, x, y, w, bottom, size: float = 15):
    """rows are (statement, lift). The box is as tall as the wrapped sentence."""
    score_w = 0.9
    text_w = w - score_w - 0.2
    cursor = y
    for statement, lift in rows:
        row_h = _row_height(statement, text_w, size)
        if cursor + row_h > bottom:
            break
        _flush(add_text(
            slide, statement, x, cursor, text_w, row_h,
            size=size, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, fmt_lift(lift), x + w - score_w, cursor, score_w, row_h,
            size=size, bold=True, color=lift_color(lift), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        ))
        cursor += row_h + 0.06
    return cursor


def _lift_column(slide, study, elements, x, y, w, row_h, slot_w, slot_h, lowest, scale):
    name_w = 1.45 if w < 7 else 1.85
    value_w = 0.62
    bar_x = x + slot_w + name_w + 0.14
    bar_w = max(1.1, w - (bar_x - x) - value_w)
    bar_h = min(0.3, row_h * 0.34)
    axis_x = bar_x + ((0 - lowest) / scale) * bar_w
    add_rect(slide, axis_x, y + 0.06, 0.015, row_h * len(elements) - 0.1, RGB_AXIS)

    for index, element in enumerate(elements):
        top = y + index * row_h
        art = _element_art(study, element)
        if art:
            _place_fit(slide, art, x, top + (row_h - slot_h) / 2, slot_w, slot_h)
        label_size = 14 if row_h >= 0.75 else 12
        add_text(
            slide,
            element.label,
            x + slot_w + 0.1,
            top,
            name_w,
            row_h,
            size=label_size,
            bold=True,
            color=NAVY,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        track_top = top + (row_h - bar_h) / 2
        add_rect(slide, bar_x, track_top, bar_w, bar_h, TRACK, rounded=True, radius=0.5)
        end = bar_x + ((element.lift - lowest) / scale) * bar_w
        right_edge = bar_x + bar_w
        start = max(bar_x, min(axis_x, end))
        stop = min(right_edge, max(axis_x, end))
        add_rect(slide, start, track_top, max(stop - start, 0.06), bar_h, bar_color(element.lift), rounded=True, radius=0.5)
        star = "*" if element.significant else ""
        add_text(
            slide,
            fmt_lift(element.lift) + star,
            bar_x + bar_w + 0.04,
            top,
            value_w,
            row_h,
            size=14 if row_h >= 0.7 else 12,
            bold=True,
            color=lift_color(element.lift),
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _ordered_categories(study: Study):
    return sorted(study.categories, key=lambda category: min(element.z_index for element in category.elements))


def _pack_picks(study: Study, choose: str):
    return legal_picks(study, choose)


def _option_label(category, element) -> str:
    """Drop the silo name from the element label. The row already names the silo."""
    text = element.label.strip()
    silo = category.label.lower()
    if silo.endswith("s") and not silo.endswith("ss"):
        silo = silo[:-1]
    if text.lower().startswith(silo + " "):
        text = text[len(silo):].strip()
    if text.lower().startswith("layer "):
        text = text[6:].strip()
    return text


def _place_fit(slide, path, x, y, max_w, max_h):
    """Place an image inside a slot without a colored frame, keeping its proportions."""
    width, height = _fit_size(path, max_w, max_h)
    left = x + max(0, (max_w - width) / 2)
    top = y + max(0, (max_h - height) / 2)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def _fit_size(path, max_w, max_h):
    with Image.open(path) as image:
        aspect = image.width / image.height if image.height else 1
    width = max_h * aspect
    height = max_h
    if width > max_w:
        width = max_w
        height = width / aspect if aspect else max_h
    return width, height


def _draw_art(slide, path, x, y, max_w, max_h):
    """Place the whole image inside the slot. The border sits outside the pixels."""
    width, height = _fit_size(path, max_w, max_h)
    left = x + max(0, (max_w - width) / 2)
    top = y
    edge = RGBColor(0xE2, 0xE8, 0xEE)
    pad = 0.015
    add_rect(slide, left - pad, top - pad, width + pad * 2, height + pad * 2, WHITE, line=edge)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    return width, height


def _element_art(study: Study, element) -> Path | None:
    """Element on the study background, cropped to the layer so it stays readable."""
    return _element_picture(study, element, tight=False)


def _element_on_pack(study: Study, element) -> Path | None:
    """The whole study background, with this one element on top. Nothing cropped."""
    if element.image_path is None or not element.image_path.exists():
        return study.background_path if study.background_path and study.background_path.exists() else None
    destination = element.image_path.parent.parent / "previews_full_fit" / f"{element.code}.png"
    background = study.background_path if study.background_path and study.background_path.exists() else None
    return element_preview(background, element.image_path, destination, width=900, full=True)


def _element_thumb(study: Study, element) -> Path | None:
    """Closer crop for rows where several elements sit side by side."""
    return _element_picture(study, element, tight=True)


def _element_picture(study: Study, element, tight: bool) -> Path | None:
    if element.image_path is None or not element.image_path.exists():
        return study.background_path if study.background_path and study.background_path.exists() else None
    folder = "previews_tight" if tight else "previews"
    destination = element.image_path.parent.parent / f"{folder}_fit" / f"{element.code}.png"
    background = study.background_path if study.background_path and study.background_path.exists() else None
    return element_preview(background, element.image_path, destination, width=1000, tight=tight)


def _place_element_row(slide, study: Study, elements, x, y, width, height):
    """Pack element artwork left to right at a shared height. No color plate."""
    gap = 0.12
    max_w = 2.05
    specs = []
    for element in elements:
        path = _element_thumb(study, element)
        if path is None:
            continue
        with Image.open(path) as image:
            aspect = image.width / image.height if image.height else 1
        slot_w = height * aspect
        slot_h = height
        if slot_w > max_w:
            slot_w = max_w
            slot_h = max_w / aspect if aspect else height
        specs.append((path, slot_w, slot_h))
    if not specs:
        return
    total = sum(slot_w for _, slot_w, _ in specs) + gap * (len(specs) - 1)
    scale = min(1.0, width / total) if total else 1
    cursor = x
    for path, slot_w, slot_h in specs:
        slot_w *= scale
        slot_h *= scale
        top = y + (height - slot_h) / 2
        _place_fit(slide, path, cursor, top, slot_w, slot_h)
        cursor += slot_w + gap * scale


def _render_pack(study: Study, picks) -> Path | None:
    if study.background_path is None and not any(element.image_path for element in picks):
        return None
    cache_root = None
    for element in picks:
        if element.image_path:
            cache_root = element.image_path.parent
            break
    if cache_root is None and study.background_path:
        cache_root = study.background_path.parent
    if cache_root is None:
        return None
    layers = []
    if study.background_path:
        layers.append(study.background_path)
    layers.extend(element.image_path for element in picks if element.image_path)
    key = "-".join(element.code for element in picks)
    return composite_pack(layers, cache_root.parent / "composites" / f"{key}.png")


def _build(presentation, study: Study):
    mode = _visual(study)
    picks = _pack_picks(study, "best")
    if mode == "grid":
        return _grid_combination_slide(
            presentation, study, picks,
            "COMBINATION WINNERS",
            "Highest-lift images, side by side",
            "One image from each set — the strongest in that set. The images are not stacked.",
        )
    if mode == "text":
        return _text_combination_slide(
            presentation, study, picks,
            "COMBINATION WINNERS",
            "Highest-lift statements",
            "The strongest statement in each group, written in full.",
        )
    if study.blocked:
        detail = "The highest coefficient sum that still obeys the pairing rules. A stronger element is left off when the elements it blocks add up to more."
    else:
        detail = "One element from every silo — the strongest in that silo — composited in layer order. The background sits underneath."
    return _combination_slide(
        presentation,
        study,
        picks,
        "COMBINATION WINNERS",
        "Highest coefficient sum, stacked into one pack",
        detail,
    )


def _build_low(presentation, study: Study):
    mode = _visual(study)
    picks = _pack_picks(study, "worst")
    if mode == "grid":
        return _grid_combination_slide(
            presentation, study, picks,
            "COMBINATION WINNERS",
            "Lowest-lift images, side by side",
            "One image from each set — the weakest in that set. The images are not stacked.",
        )
    if mode == "text":
        return _text_combination_slide(
            presentation, study, picks,
            "COMBINATION WINNERS",
            "Lowest-lift statements",
            "The weakest statement in each group, written in full.",
        )
    if study.blocked:
        detail = "The lowest coefficient sum that still fills every silo it can, with no blocked pair on the pack."
    else:
        detail = "The same silos and the same stacking order, using each silo’s weakest element."
    return _combination_slide(
        presentation,
        study,
        picks,
        "COMBINATION WINNERS",
        "Lowest coefficient sum, stacked into one pack",
        detail,
    )


def _build_compare(presentation, study: Study):
    """Higher-lift build against the lower-lift build, one (T) score per layer."""
    categories = _ordered_categories(study)
    high = _pack_picks(study, "best")
    low = _pack_picks(study, "worst")
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, "COMBINATION WINNERS")
    add_text(slide, "The higher-lift build against the lower-lift build", 0.55, 0.58, 12.2, 0.40, size=24, bold=True, color=NAVY)
    high_total = sum(element.lift for element in high)
    low_total = sum(element.lift for element in low)
    gap = high_total - low_total
    _compare_card(
        slide, study, categories, high, 0.42, 1.22, 6.15, 4.72,
        "HIGHER-LIFT BUILD", TEAL, RGBColor(0xF3, 0xF8, 0xF7), high_total, study.base_appeal,
    )
    _compare_card(
        slide, study, categories, low, 6.76, 1.22, 6.15, 4.72,
        "LOWER-LIFT BUILD", ROSE, RGBColor(0xFD, 0xF4, 0xF6), low_total, study.base_appeal,
    )
    if study.base_appeal is not None:
        summary = f"The higher-lift build leads by {fmt_lift(gap)} points, from the same {fmt_pct(study.base_appeal)} base."
    elif _visual(study) == "text":
        summary = f"The higher-lift statements lead by {fmt_lift(gap)} points. This model has no intercept, so the totals are the statement lifts."
    else:
        summary = f"The higher-lift build leads by {fmt_lift(gap)} points. This model has no intercept, so the totals are the element lifts."
    if study.blocked:
        summary += " Both totals are the highest and lowest sums that still obey the pairing rules."
    add_rect(slide, 0.42, 6.08, 12.49, 0.52, RGBColor(0xFC, 0xF4, 0xF8), rounded=True, radius=0.08)
    _flush(add_text(
        slide, summary, 0.62, 6.08, 12.1, 0.52,
        size=14, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    ))
    return slide


def _compare_card(slide, study, categories, picks, x, y, w, h, title, accent, header_fill, total, base):
    edge = RGBColor(0xE4, 0xE9, 0xF0)
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.03, line=edge)
    header_h = 0.78 if _visual(study) != "layer" else 1.36
    add_rect(slide, x + 0.02, y + 0.08, w - 0.04, header_h - 0.08, header_fill)
    add_rect(slide, x, y, w, 0.055, accent)
    _flush(add_text(slide, title, x + 0.16, y + 0.12, w - 0.32, 0.22, size=11, bold=True, color=accent))
    if _visual(study) == "layer":
        image = _render_pack(study, picks)
        pack_w, pack_h = contain_size(file_aspect(image) if image else PACK_ASPECT, 1.15, 0.92)
        if image:
            slide.shapes.add_picture(
                str(image), Inches(x + 0.16), Inches(y + 0.38), Inches(pack_w), Inches(pack_h),
            )
        score_x = x + 0.16 + pack_w + 0.22
        score_w = w - (score_x - x) - 0.16
        _flush(add_text(
            slide, fmt_lift(total), score_x, y + 0.42, score_w, 0.48,
            size=28, bold=True, color=lift_color(total),
        ))
        _flush(add_text(
            slide, _base_line(base), score_x, y + 0.90, score_w, 0.26,
            size=13, color=GRAY,
        ))
    else:
        _flush(add_text(
            slide, fmt_lift(total), x + 0.16, y + 0.36, 1.6, 0.34,
            size=20, bold=True, color=lift_color(total),
        ))
        _flush(add_text(
            slide, _base_line(base), x + 1.85, y + 0.38, w - 2.1, 0.30,
            size=12, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
        ))
    list_top = y + header_h + 0.06
    mode = _visual(study)
    rows = _pick_rows(study, picks)
    if mode == "text":
        _text_compare_rows(slide, rows, x, list_top, w, y + h - 0.08)
        return
    row_h = (y + h - 0.08 - list_top) / max(len(rows), 1)
    rule = RGBColor(0xF0, 0xE6, 0xEA)
    thumb = min(row_h - 0.08, 0.72) if mode == "grid" else 0.0
    silo_w = 1.28
    score_col = 0.72
    name_x = x + 0.16 + silo_w + (thumb + 0.08 if thumb else 0)
    name_w = w - (name_x - x) - score_col - 0.16
    for index, (category, element) in enumerate(rows):
        top = list_top + index * row_h
        if index:
            add_rect(slide, x + 0.14, top, w - 0.28, 0.01, rule)
        _flush(add_text(
            slide, category.label, x + 0.16, top, silo_w - 0.08, row_h,
            size=11, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
        ))
        if thumb:
            art = _raw_image(element)
            if art:
                _place_fit(slide, art, x + 0.16 + silo_w, top + 0.04, thumb, row_h - 0.08)
        label = _option_label(category, element)
        _flush(add_text(
            slide, label, name_x, top, name_w, row_h,
            size=13, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, fmt_lift(element.lift), x + w - score_col - 0.12, top, score_col, row_h,
            size=14, bold=True, color=lift_color(element.lift), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        ))


def _pick_rows(study: Study, picks) -> list[tuple]:
    """Category and element, in layer order. A silo left off the pack is not listed."""
    rows = []
    for category in _ordered_categories(study):
        for element in picks:
            if element in category.elements:
                rows.append((category, element))
                break
    return rows


def _text_compare_rows(slide, rows, x, top, w, bottom):
    """Group name on its own line, full statement beside the score. Rows are not stretched."""
    cursor = top
    score_w = 0.72
    statement_w = w - 0.36 - score_w
    size = 12
    for index, (category, element) in enumerate(rows):
        text_h = max(0.36, _row_height(element.label, statement_w, size))
        block = 0.22 + text_h
        if cursor + block > bottom:
            break
        if index:
            add_rect(slide, x + 0.14, cursor, w - 0.28, 0.01, RGBColor(0xF0, 0xE6, 0xEA))
        _flush(add_text(
            slide, category.label, x + 0.16, cursor, w - 0.32, 0.22,
            size=10, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, element.label, x + 0.16, cursor + 0.22, statement_w, text_h,
            size=size, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, fmt_lift(element.lift), x + w - score_w - 0.12, cursor + 0.22, score_w, text_h,
            size=14, bold=True, color=lift_color(element.lift), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        ))
        cursor += block


def _base_line(base) -> str:
    if base is None:
        return "Lift only, no separate base"
    return f"over {fmt_pct(base)} base"


def _combination_slide(presentation, study, picks, eyebrow, headline, detail):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, eyebrow)
    add_text(slide, headline, 0.48, 0.52, 12.3, 0.38, size=26, bold=True, color=NAVY)
    add_text(slide, detail, 0.48, 0.96, 12.3, 0.46, size=14, color=GRAY)

    stage_x, stage_y, stage_w, stage_h = 0.42, 1.56, 4.42, 5.12
    add_rect(slide, stage_x, stage_y, stage_w, stage_h, RGBColor(0xFB, 0xF4, 0xF7), rounded=True, radius=0.12)
    total = sum(element.lift for element in picks)
    image = _render_pack(study, picks)
    pack_w, pack_h = contain_size(file_aspect(image) if image else PACK_ASPECT, stage_w - 0.7, stage_h - 1.55)
    if image:
        left = stage_x + (stage_w - pack_w) / 2
        top = stage_y + 0.22 + max(0, (stage_h - 1.45 - pack_h) / 2)
        slide.shapes.add_picture(str(image), Inches(left), Inches(top), Inches(pack_w), Inches(pack_h))
    score_y = stage_y + stage_h - 1.05
    add_rect(slide, stage_x + 0.28, score_y, stage_w - 0.56, 0.78, WHITE, rounded=True, radius=0.1)
    _flush(add_text(
        slide, fmt_lift(total), stage_x + 0.42, score_y, stage_w - 0.84, 0.46,
        size=26, bold=True, color=lift_color(total), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM,
    ))
    caption = "coefficient sum" if study.base_appeal is None else f"coefficient sum  ·  base {fmt_pct(study.base_appeal)}"
    _flush(add_text(
        slide, caption, stage_x + 0.42, score_y + 0.42, stage_w - 0.84, 0.28,
        size=11, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
    ))

    _layer_list(slide, study, picks, 5.08, stage_y, 7.82, stage_h)
    return slide


def _grid_combination_slide(presentation, study, picks, eyebrow, headline, detail):
    """Winning images in a row, each with its name and lift. Nothing is stacked."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, eyebrow)
    add_text(slide, headline, 0.55, 0.52, 12.2, 0.4, size=24, bold=True, color=NAVY)
    add_text(slide, detail, 0.55, 0.96, 12.2, 0.36, size=14, color=GRAY)

    count = max(len(picks), 1)
    gap = 0.22
    usable = 12.2
    col_w = (usable - gap * (count - 1)) / count
    box_top = 1.5
    box_h = 3.35
    for index, element in enumerate(picks):
        left = 0.55 + index * (col_w + gap)
        category = next((item for item in study.categories if element in item.elements), None)
        art = _raw_image(element)
        if art:
            _place_fit(slide, art, left, box_top, col_w, box_h)
        else:
            add_rect(slide, left, box_top, col_w, box_h, RGBColor(0xF7, 0xF4, 0xF5), rounded=True, radius=0.04)
        text_top = box_top + box_h + 0.12
        silo = category.label if category else ""
        _flush(add_text(slide, silo, left, text_top, col_w, 0.24, size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER))
        _flush(add_text(slide, element.label, left, text_top + 0.26, col_w, 0.48, size=13, color=NAVY, align=PP_ALIGN.CENTER))
        _flush(add_text(
            slide, fmt_lift(element.lift), left, text_top + 0.76, col_w, 0.32,
            size=16, bold=True, color=lift_color(element.lift), align=PP_ALIGN.CENTER,
        ))
    total = sum(element.lift for element in picks)
    if study.base_appeal is None:
        line = f"Combined lift {fmt_lift(total)}"
    else:
        line = f"Base {fmt_pct(study.base_appeal)}    ·    combined lift {fmt_lift(total)}"
    add_text(slide, line, 0.55, 6.55, 12.2, 0.32, size=14, bold=True, color=NAVY)
    return slide


def _text_combination_slide(presentation, study, picks, eyebrow, headline, detail):
    """Top or bottom statements, grouped under the title. No empty picture column."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, eyebrow)
    add_text(slide, headline, 0.55, 0.50, 12.2, 0.38, size=24, bold=True, color=NAVY)
    add_text(slide, detail, 0.55, 0.92, 12.2, 0.30, size=14, color=GRAY)
    cursor = 1.40
    score_w = 0.95
    statement_w = 8.35
    size = 15
    for element in picks:
        silo = next((category.label for category in study.categories if element in category.elements), "")
        row_h = max(0.64, _row_height(element.label, statement_w, size))
        add_rect(slide, 0.55, cursor, 12.2, row_h, RGBColor(0xFB, 0xF6, 0xF8), rounded=True, radius=0.06)
        _flush(add_text(
            slide, silo, 0.72, cursor, 2.35, row_h,
            size=12, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, element.label, 3.15, cursor + 0.06, statement_w, row_h - 0.12,
            size=size, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        ))
        _flush(add_text(
            slide, fmt_lift(element.lift), 0.55 + 12.2 - score_w - 0.12, cursor, score_w, row_h,
            size=16, bold=True, color=lift_color(element.lift), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        ))
        cursor += row_h + 0.10
    total = sum(element.lift for element in picks)
    if study.base_appeal is None:
        line = f"Combined lift {fmt_lift(total)}"
    else:
        line = f"Base {fmt_pct(study.base_appeal)}    ·    combined lift {fmt_lift(total)}"
    add_rect(slide, 0.55, cursor + 0.06, 12.2, 0.42, RGBColor(0xF4, 0xE0, 0xE7), rounded=True, radius=0.06)
    _flush(add_text(
        slide, line, 0.72, cursor + 0.06, 11.8, 0.42,
        size=16, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    ))
    return slide


def _pack_column(slide, study, picks, x, title, accent):
    add_text(slide, title, x, 1.5, 5.8, 0.28, size=12, bold=True, color=accent)
    image = _render_pack(study, picks)
    pack_w, pack_h = contain_size(file_aspect(image) if image else PACK_ASPECT, 4.6, 3.05)
    left = x + (5.6 - pack_w) / 2
    top = 1.86 + (3.05 - pack_h) / 2
    if image:
        slide.shapes.add_picture(str(image), Inches(left), Inches(top), Inches(pack_w), Inches(pack_h))
    total = sum(element.lift for element in picks)
    if study.base_appeal is None:
        pack_line = "Layer lift, no separate base"
    else:
        pack_line = f"Layer lift over {fmt_pct(study.base_appeal)} base"
    add_text(slide, fmt_lift(total), x, 5.05, 5.8, 0.38, size=26, bold=True, color=lift_color(total), align=PP_ALIGN.CENTER)
    add_text(
        slide,
        pack_line,
        x, 5.42, 5.8, 0.28, size=12, color=GRAY, align=PP_ALIGN.CENTER,
    )
    names = "  ·  ".join(element.label for element in picks[:6])
    if len(picks) > 6:
        names += "  ·  …"
    add_text(slide, names, x, 5.74, 5.8, 0.55, size=11, color=NAVY, align=PP_ALIGN.CENTER)


def _layer_list(slide, study, picks, x, y, w, h):
    """One chip per layer: category, element, and a bar against the strongest coefficient."""
    count = max(len(picks), 1)
    gap = 0.05
    row_h = (h - gap * max(0, count - 1)) / count
    peak = max((abs(element.lift) for element in picks), default=1.0) or 1.0
    score_w = 0.82
    text_w = w - score_w - 0.28
    two_line = row_h >= 0.48
    for index, element in enumerate(picks):
        top = y + index * (row_h + gap)
        silo = next((category.label for category in study.categories if element in category.elements), "")
        add_rect(slide, x, top, w, row_h, RGBColor(0xFB, 0xF6, 0xF8), rounded=True, radius=0.06)
        if two_line:
            _flush(add_text(
                slide, f"{index + 1:02d}    {silo.upper()}", x + 0.16, top + 0.02, text_w, 0.18,
                size=10, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE,
            ))
            _flush(add_text(
                slide, element.label, x + 0.16, top + 0.19, text_w, 0.2,
                size=13, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
            ))
            track_w = text_w
            track_y = top + row_h - 0.07
            add_rect(slide, x + 0.16, track_y, track_w, 0.028, RGBColor(0xF0, 0xE4, 0xEA))
            add_rect(
                slide, x + 0.16, track_y, max(0.06, track_w * min(1.0, abs(element.lift) / peak)), 0.028,
                lift_color(element.lift),
            )
        else:
            _flush(add_text(
                slide, f"{index + 1:02d}    {silo}   ·   {element.label}", x + 0.14, top, text_w, row_h,
                size=12, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
            ))
        _flush(add_text(
            slide, fmt_lift(element.lift), x + w - score_w - 0.1, top, score_w, row_h,
            size=15 if two_line else 13, bold=True, color=lift_color(element.lift),
            align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        ))


def _appendix_intro(presentation, study: Study):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _eyebrow(slide, "APPENDIX  ·  FULL ELEMENT LIFT SUMMARY")
    add_text(slide, "Every element, one silo per page", 0.62, 0.72, 12, 0.5, size=30, bold=True, color=NAVY)
    group_line = ""
    if study.classifications:
        labels = ", ".join(segment.label for segment in study.classifications[0].segments)
        group_line = f" Overall is the full sample. The other columns are {labels}."
    add_text(
        slide,
        f"Scores are {study.model_name} coefficients ({study.rating_bands}), rounded the same way as the charts. "
        f"Green is lift above the baseline. Red is lift below it.{group_line}",
        0.62, 1.35, 12, 0.7, size=16, color=GRAY,
    )
    add_rect(slide, 0.62, 2.3, 3.8, 1.5, BLUSH_CARD, rounded=True)
    add_text(
        slide,
        fmt_pct(study.base_appeal) if study.base_appeal is not None else "No intercept",
        0.85, 2.45, 3.4, 0.7, size=32, bold=True, color=NAVY,
    )
    add_text(slide, "Base appeal", 0.85, 3.2, 3.4, 0.35, size=14, color=GRAY)
    add_rect(slide, 4.65, 2.3, 3.8, 1.5, BLUSH_CARD, rounded=True)
    add_text(slide, str(study.silo_count), 4.9, 2.45, 3.4, 0.7, size=32, bold=True, color=NAVY)
    add_text(slide, "Silos in the tables", 4.9, 3.2, 3.4, 0.35, size=14, color=GRAY)
    add_rect(slide, 8.68, 2.3, 3.95, 1.5, BLUSH_CARD, rounded=True)
    add_text(slide, str(study.element_count), 8.9, 2.45, 3.5, 0.7, size=32, bold=True, color=NAVY)
    add_text(slide, "Elements scored", 8.9, 3.2, 3.5, 0.35, size=14, color=GRAY)
    _FOOTER_NOTES[id(slide)] = _appendix_footer(study)
    return slide


def _appendix_question(study: Study):
    """The classification cut shown beside Overall. Findings slides carry any further cuts."""
    if not study.classifications:
        return None
    return study.classifications[0]


def _appendix_footer(study: Study) -> str:
    """One line under the appendix tables. Edit this sentence to change that line."""
    question = _appendix_question(study)
    if question and any(segment.base_size for segment in question.segments):
        sample = "; ".join(
            f"n={segment.base_size} {segment.label}" if segment.base_size else segment.label
            for segment in question.segments
        )
    else:
        sample = f"n={study.respondents or study.base_size} respondents"
    fractional = False
    for category in study.categories:
        for element in category.elements:
            values = [element.lift]
            if question:
                values.extend(float(element.segment_lifts.get(segment.key, 0.0)) for segment in question.segments)
            if any(abs(float(value) - round(float(value))) >= 0.05 for value in values):
                fractional = True
                break
        if fractional:
            break
    places = "one decimal" if fractional else "whole points"
    return (
        f"{sample}. {study.model_name} model: element lifts show movement above or below the model baseline. "
        f"Values rounded to {places}."
    )


def _lift_tables(presentation, study: Study):
    """One silo per page, as a score table. A very long silo continues on the next page."""
    if _visual(study) == "text":
        return _text_summary_slides(presentation, study)
    slides = []
    for category in study.categories:
        ordered = sorted(category.elements, key=lambda element: element.lift, reverse=True)
        pages = _pages(ordered, 12)
        for index, page in enumerate(pages, start=1):
            slides.append(_appendix_table_slide(presentation, study, category, page, index, len(pages)))
    return slides


def _appendix_table_slide(presentation, study, category, elements, page: int, pages: int):
    """Element, overall lift, and classification columns. Matches the appendix readout."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    question = _appendix_question(study)
    _headline, detail, _key = classification_story(category, question) if question else category_story(category)
    _eyebrow(slide, "FULL ELEMENT LIFT SUMMARY")
    title = f"{category.label} — Element Lift Summary{_page_mark(page, pages)}"
    title_h = _draw_fitted(
        slide, title, 0.62, 0.56, 12.1, 0.72,
        size=28, min_size=18, color=NAVY, bold=True, max_lines=2, one_line_floor=20,
    )
    detail_y = 0.56 + title_h + 0.08
    detail_h = _draw_fitted(
        slide, detail, 0.62, detail_y, 12.1, 0.48,
        size=15, min_size=11, color=GRAY, max_lines=2, one_line_floor=12,
    )
    banner_y = detail_y + detail_h + 0.16
    banner = add_rect(slide, 0.62, banner_y, 12.1, 0.32, RGBColor(0xF7, 0xE6, 0xEC))
    _shape_label(banner, category.label.upper(), 12, NAVY, bold=True, wrap=False)
    _draw_score_table(slide, elements, question, 0.62, banner_y + 0.32, 12.1)
    _FOOTER_NOTES[id(slide)] = _appendix_footer(study)
    return slide


def _draw_score_table(slide, elements, question, x, y, w) -> None:
    """Open table: name on the left, overall and each group centered, green above and red below."""
    segments = list(question.segments) if question else []
    if len(segments) > 5:
        segments = []
    count = 1 + len(segments)
    score_w = 1.55 if count <= 3 else (1.28 if count <= 5 else 1.05)
    element_w = w - score_w * count
    if element_w < 4.4:
        score_w = max(0.95, (w - 4.4) / count)
        element_w = w - score_w * count
    columns = [("Element", element_w, PP_ALIGN.LEFT, None)]
    columns.append(("Overall", score_w, PP_ALIGN.CENTER, "overall"))
    for segment in segments:
        columns.append((segment.label, score_w, PP_ALIGN.CENTER, segment))

    header_h = 0.36
    count = max(len(elements), 1)
    room = (6.80 - y - header_h) / count
    preferred = 0.46 if count <= 6 else (0.42 if count <= 9 else 0.36)
    row_h = min(preferred, max(room, 0.30))
    cursor = x
    rule = RGBColor(0xE7, 0xD5, 0xDE)
    hairline = RGBColor(0xF3, 0xEA, 0xEE)
    for label, width, align, _kind in columns:
        pad = 0.02 if align == PP_ALIGN.CENTER else 0.0
        size, _height = fitted_size(label, width - 0.1, header_h, 12, minimum=9, bold=True, max_lines=1)
        box = add_text(
            slide, label, cursor + pad, y, width - pad * 2, header_h,
            size=size, bold=True, color=NAVY, align=align, anchor=MSO_ANCHOR.MIDDLE,
        )
        _flush(box)
        apply_leading(box, size)
        cursor += width
    add_rect(slide, x, y + header_h, w, 0.012, rule)

    for index, element in enumerate(elements):
        top = y + header_h + index * row_h
        cursor = x
        values = [("name", element.label)]
        values.append(("overall", element.lift))
        for segment in segments:
            values.append(("segment", element.segment_lifts.get(segment.key)))
        for (_label, width, align, kind), (kind_name, value) in zip(columns, values):
            if kind_name == "name":
                size, _height = fitted_size(str(value), width - 0.12, row_h - 0.04, 14, minimum=11, max_lines=2)
                box = add_text(
                    slide, value, cursor, top, width - 0.08, row_h,
                    size=size, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
                )
                _flush(box)
                apply_leading(box, size)
            else:
                number = fmt_lift(value)
                color = GRAY if value is None else lift_color(float(value))
                box = add_text(
                    slide, number, cursor, top, width, row_h,
                    size=14, bold=(kind == "overall"), color=color, align=align, anchor=MSO_ANCHOR.MIDDLE,
                )
                _flush(box)
                apply_leading(box, 14)
            cursor += width
        add_rect(slide, x, top + row_h - 0.008, w, 0.008, hairline)


def _asset_slides(presentation, study: Study):
    mode = _visual(study)
    if mode == "text":
        return _text_asset_slides(presentation, study)
    if mode == "grid":
        return _grid_asset_slides(presentation, study)
    slides = []
    for category in study.categories:
        elements = list(category.elements)
        pages = _pages(elements, _ASSET_PER_PAGE)
        for index, page in enumerate(pages, start=1):
            rows = [page[start:start + _ASSET_PER_ROW] for start in range(0, len(page), _ASSET_PER_ROW)]
            slide = _blank(presentation)
            set_background(slide, WHITE)
            _eyebrow(slide, "ELEMENT ASSETS")
            add_text(
                slide, f"{category.label}  —  element assets{_page_mark(index, len(pages))}",
                0.62, 0.68, 12, 0.42, size=26, bold=True, color=NAVY,
            )
            add_text(
                slide,
                f"{category.code}   ·   {len(elements)} variations",
                0.62, 1.15, 12, 0.3, size=14, color=GRAY,
            )
            _place_asset_rows(slide, study, rows, top=1.55)
            slides.append(slide)
    return slides


def _place_asset_rows(slide, study, rows, top: float):
    """Each asset keeps its own aspect. A slide holds at most two rows."""
    gap = 0.18
    label_h = 0.46
    row_gap = 0.16
    max_h = 2.15
    max_row_w = 12.1
    prepared = []
    for row in rows:
        items = []
        for element in row:
            art = _element_on_pack(study, element)
            aspect = file_aspect(art) if art else 1.0
            items.append((element, art, aspect))
        prepared.append(items)
    row_heights = []
    for items in prepared:
        if not items:
            row_heights.append(max_h)
            continue
        aspect_sum = sum(aspect for _element, _art, aspect in items)
        gaps = gap * (len(items) - 1)
        height = min(max_h, (max_row_w - gaps) / aspect_sum if aspect_sum else max_h)
        row_heights.append(max(height, 0.4))
    block_h = sum(height + label_h for height in row_heights) + row_gap * max(len(prepared) - 1, 0)
    available_h = 6.7 - top
    origin = top + max(0, (available_h - block_h) / 2)
    y = origin
    for items, height in zip(prepared, row_heights):
        widths = [height * aspect for _element, _art, aspect in items]
        total_w = sum(widths) + gap * max(len(items) - 1, 0)
        left = 0.62 + max(0, (max_row_w - total_w) / 2)
        for (element, art, _aspect), width in zip(items, widths):
            if art:
                _place_fit(slide, art, left, y, width, height)
            else:
                add_text(slide, element.code, left, y, width, height, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            label_box = add_text(
                slide, wrap_label(element.label, width, 10),
                left - 0.02, y + height + 0.02, width + 0.04, label_h,
                size=10, color=NAVY, align=PP_ALIGN.CENTER,
            )
            label_box.text_frame.word_wrap = False
            left += width + gap
        y += height + label_h + row_gap


def _text_summary_slides(presentation, study: Study):
    """Appendix list of every statement, in full, with its lift."""
    slides = []
    for category in study.categories:
        ordered = sorted(category.elements, key=lambda element: element.lift, reverse=True)
        text_w = 10.6
        top, bottom, size = 1.35, 6.6, 14
        pages: list[list] = []
        current: list = []
        cursor = top
        for element in ordered:
            row_h = _row_height(element.label, text_w, size)
            if current and cursor + row_h > bottom:
                pages.append(current)
                current = []
                cursor = top
            current.append(element)
            cursor += row_h + 0.06
        if current:
            pages.append(current)
        for index, page in enumerate(pages, start=1):
            slide = _blank(presentation)
            set_background(slide, WHITE)
            _eyebrow(slide, "FULL ELEMENT LIFT SUMMARY")
            add_text(
                slide, f"{category.label}{_page_mark(index, len(pages))}",
                0.55, 0.55, 12.2, 0.42, size=26, bold=True, color=NAVY,
            )
            _statement_rows(
                slide, [(element.label, element.lift) for element in page],
                0.55, 1.25, 12.2, 6.6, size=14,
            )
            slides.append(slide)
    return slides


def _text_asset_slides(presentation, study: Study):
    slides = []
    for category in study.categories:
        elements = list(category.elements)
        text_w = 5.5
        size = 13
        per_column_budget = 4.6
        columns = [[]]
        heights = [0.0]
        for element in elements:
            row_h = _row_height(element.label, text_w, size)
            if columns[-1] and heights[-1] + row_h > per_column_budget and len(columns) % 2 == 1:
                columns.append([])
                heights.append(0.0)
            elif columns[-1] and heights[-1] + row_h > per_column_budget:
                columns.append([])
                heights.append(0.0)
            columns[-1].append(element)
            heights[-1] += row_h + 0.1
        pages = [columns[index:index + 2] for index in range(0, len(columns), 2)]
        for index, page in enumerate(pages, start=1):
            slide = _blank(presentation)
            set_background(slide, WHITE)
            _eyebrow(slide, "ELEMENT ASSETS")
            add_text(
                slide, f"{category.label}  —  statements{_page_mark(index, len(pages))}",
                0.55, 0.55, 12.2, 0.4, size=26, bold=True, color=NAVY,
            )
            for column_index, column in enumerate(page):
                left = 0.55 + column_index * 6.35
                _statement_rows(
                    slide, [(element.label, element.lift) for element in column],
                    left, 1.25, 6.05, 6.6, size=size,
                )
            slides.append(slide)
    return slides


def _grid_asset_slides(presentation, study: Study):
    """Each logo keeps its own shape. The frame does not stretch it."""
    slides = []
    per_row = 4
    for category in study.categories:
        elements = list(category.elements)
        rows = [elements[start:start + per_row] for start in range(0, len(elements), per_row)]
        pages = [rows[start:start + 2] for start in range(0, len(rows), 2)] or [[]]
        for index, page in enumerate(pages, start=1):
            slide = _blank(presentation)
            set_background(slide, WHITE)
            _eyebrow(slide, "ELEMENT ASSETS")
            add_text(
                slide, f"{category.label}  —  images{_page_mark(index, len(pages))}",
                0.55, 0.52, 12.2, 0.4, size=26, bold=True, color=NAVY,
            )
            cell = 2.15
            gap = 0.28
            label_h = 0.5
            for row_index, row in enumerate(page):
                count = len(row)
                total_w = count * cell + (count - 1) * gap
                left0 = 0.55 + (12.2 - total_w) / 2
                top = 1.25 + row_index * (cell + label_h + 0.28)
                for item_index, element in enumerate(row):
                    left = left0 + item_index * (cell + gap)
                    art = _raw_image(element)
                    if art:
                        _place_fit(slide, art, left, top, cell, cell)
                    _flush(add_text(
                        slide, element.label, left - 0.05, top + cell + 0.04, cell + 0.1, label_h,
                        size=12, color=NAVY, align=PP_ALIGN.CENTER,
                    ))
            slides.append(slide)
    return slides


def _thanks(presentation, study: Study):
    """Same left-side marks as the cover, including the MGA logo."""
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    report_name = _cover_marks(slide, study)
    _hairline(slide, 4.2464, 1.3478, 4.2464, 1.3478 + 4.7826)
    _hairline(slide, 4.5996, 3.75, 4.5996 + 6.8551, 3.75)
    add_text(slide, "Thank you", 4.600, 1.975, 7.8, 0.84, size=44, bold=True, color=NAVY)
    add_text(
        slide,
        "We look forward to discussing the full report and next steps.",
        4.600, 2.82, 7.2, 0.40, size=16, color=NAVY,
    )
    when = study.launched_label or ""
    add_text(
        slide,
        f"{report_name}  |  {when}".strip(" |"),
        4.600, 4.244, 7.2, 0.37, size=16, color=NAVY,
    )
    add_text(
        slide,
        _prepared_by_line(study),
        4.600, 4.628, 7.64, 0.37, size=16, color=NAVY,
    )
    add_text(slide, "Mind Genomics Associates, Inc", 4.600, 5.012, 4.2, 0.37, size=16, color=NAVY)
    add_text(slide, study.title, 4.600, 5.711, 8.4, 0.55, size=20, color=NAVY)
    return slide


def _kicker(label: str) -> str:
    parts = []
    for word in label.split():
        if any(character.isdigit() for character in word):
            parts.append(word)
        else:
            parts.append(word.upper())
    return " ".join(parts)


def _eyebrow(slide, text: str):
    add_text(slide, text, 0.62, 0.28, 12, 0.28, size=12, bold=True, color=TEAL)


