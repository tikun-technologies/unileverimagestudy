"""Design Combination Readout driven by saved combination groups.

The slide order follows the Rexona x MGA combination deck: cover, how to
read lifts, then a section for each metric that appears in the file (Top Down,
Bottom Up, Response Time). Inside a section, each saved group is its own
segment: a divider, a comparison of the saved combinations, and one detail slide per design.
"""

from __future__ import annotations

import json
import math
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import Image, ImageChops, ImageDraw

from app.services.reportgen.composite import composite_pack, contain_size, file_aspect
from app.services.reportgen.study import fetch_image
from app.services.reportgen.theme import (
    BLUSH,
    FONT,
    GRAY,
    GREEN,
    NAVY,
    PINK,
    RED,
    TEAL,
    WHITE,
    add_rect,
    add_text,
    set_background,
)

ASSETS = Path(__file__).resolve().parent / "template_assets"
SLIDE_W = 13.333333
SLIDE_H = 7.5
RAIL = RGBColor(0xF3, 0xC9, 0xD4)
INK = RGBColor(0x17, 0x22, 0x38)
CARD = RGBColor(0xFB, 0xEF, 0xF2)
TRACK = RGBColor(0xF4, 0xF0, 0xF2)

TOP_RANKS = (
    "PACK 1 — HIGHEST LIFT BUILD",
    "PACK 2 — SECOND HIGHEST LIFT BUILD",
    "PACK 3 — THIRD HIGHEST LIFT BUILD",
)
BOTTOM_RANKS = (
    "PACK 1 — LOWEST LIFT BUILD",
    "PACK 2",
    "PACK 3",
)
RESPONSE_RANKS = ("PACK 1", "PACK 2", "PACK 3")

_KIND = {
    "top": {
        "family": "COMBINATIONS",
        "prefix": "PACK COMBINATIONS",
        "chip": "COMBINATION",
        "ranks": TOP_RANKS,
        "metric": "Pack",
        "base_note": "average appeal before any element is added",
    },
    "bottom": {
        "family": "COMBINATIONS",
        "prefix": "PACK COMBINATIONS",
        "chip": "COMBINATION",
        "ranks": BOTTOM_RANKS,
        "metric": "Pack",
        "base_note": "average score before any element is added",
    },
    "response": {
        "family": "COMBINATIONS",
        "prefix": "PACK COMBINATIONS",
        "chip": "COMBINATION",
        "ranks": RESPONSE_RANKS,
        "metric": "Pack",
        "base_note": "average time before any element is added",
    },
}

def _metric_kind(metric: str) -> str:
    text = metric.lower()
    if "bottom" in text:
        return "bottom"
    if "response" in text:
        return "response"
    return "top"

def _digits(kind: str) -> int:
    return 3 if kind == "response" else 1

def _named_visual(value) -> str | None:
    """Layer, grid, or text. Hybrid uses the layer layout."""
    text = str(value or "").strip().lower()
    if text == "hybrid":
        return "layer"
    if text in {"layer", "grid", "text"}:
        return text
    return None

def _choose_visual(explicit: str | None, groups: list[Group], fallback: str) -> str:
    """An explicit type wins, then the type stored on each saved item, then the study file."""
    found = _named_visual(explicit)
    if found:
        return found
    for group in groups:
        for build in group.builds:
            found = _named_visual(build.study_type)
            if found:
                return found
    return _named_visual(fallback) or "layer"

def _http_url(value) -> str | None:
    if not isinstance(value, str):
        return None
    text = value.strip()
    if text.lower().startswith(("http://", "https://")):
        return text
    return None

def _optional_float(value) -> float | None:
    if value is None:
        return None
    return float(value)

def _base_for(study: StudyInfo, kind: str) -> float | None:
    if kind == "bottom":
        return study.base_bottom
    if kind == "response":
        return study.base_response
    return study.base_appeal

@dataclass
class ComparisonSegment:
    key: str
    label: str
    base_size: int = 0
    question: str = ""

@dataclass
class Element:
    name: str
    label: str
    lift: float
    category: str
    z_index: int
    image_url: str | None = None
    image_path: Path | None = None
    segment_lifts: dict[str, float] = field(default_factory=dict)

@dataclass
class Build:
    position: int
    name: str
    elements: list[Element]
    total: float
    metric: str = "Top Down"
    study_type: str = ""
    background_url: str | None = None
    background_path: Path | None = None
    comparison_segments: list[ComparisonSegment] = field(default_factory=list)
    comparison_title: str = ""
    audience_label: str = ""

    def signature(self) -> tuple[str, ...]:
        return tuple(element.name for element in self.elements)

@dataclass
class Group:
    name: str
    position: int
    segment: str
    metric: str
    builds: list[Build] = field(default_factory=list)

    def builds_for(self, kind: str) -> list[Build]:
        return [build for build in self.builds if _metric_kind(build.metric) == kind]

@dataclass
class StudyInfo:
    title: str
    country: str
    launched_label: str
    respondents: int
    base_appeal: float | None
    base_bottom: float | None
    base_response: float | None
    model_name: str
    aspect: float
    visual: str = "layer"
    segment_sizes: dict[str, int] = field(default_factory=dict)
    prepared_by: str = ""
    main_question: str = ""

def build_combination_readout(
    combinations_path: str | Path,
    study_path: str | Path,
    output_path: str | Path,
    analysis_path: str | Path | None = None,
    download_images: bool = True,
    study_type: str | None = None,
    cache_dir: str | Path | None = None,
    prepared_by: str | None = None,
) -> Path:
    groups = load_groups(combinations_path)
    study = load_study(study_path, analysis_path)
    study.visual = _choose_visual(study_type, groups, study.visual)
    if prepared_by:
        study.prepared_by = prepared_by
    if download_images and study.visual != "text":
        cache = Path(cache_dir) if cache_dir else Path(combinations_path).parent / ".cache" / "elements"
        _download(groups, cache)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    presentation.core_properties.title = f"Design Combination Readout v2 — {study.title}"
    presentation.core_properties.subject = "Design Combination Readout"
    presentation.core_properties.category = study.model_name

    sections = (
        ("top", "Combinations"),
        ("bottom", "Combinations"),
        ("response", "Combinations"),
    )
    slides: list[tuple[str, object]] = []
    slides.append(("plain", _cover(presentation, study)))
    slides.append(("content", _how_to_read(presentation, study, groups)))
    for kind, title in sections:
        blocks = [(group, group.builds_for(kind)) for group in groups]
        blocks = [(group, builds) for group, builds in blocks if builds]
        if not blocks:
            continue
        slides.append(("plain", _section(presentation, title)))
        for group, builds in blocks:
            builds = sorted(builds, key=lambda build: -build.total)
            slides.extend(_group_slides(presentation, study, group, builds, kind))
    slides.append(("plain", _thanks(presentation, study)))

    for index, (kind, slide) in enumerate(slides, start=1):
        if kind == "content":
            _page_number(slide, index)
    presentation.save(str(destination))
    return destination

def load_groups(path: str | Path) -> list[Group]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    raw_groups = _as_groups(payload)
    groups: list[Group] = []
    for raw in raw_groups:
        items = list(raw.get("items") or [])
        items.sort(key=lambda item: int(item.get("position") or 0))
        builds = [_build_from_item(item) for item in items]
        if not builds:
            continue
        first = items[0]
        item_segment = str(first.get("segment_label") or (first.get("configuration") or {}).get("segment", {}).get("label") or "")
        # The group name is the saved segment (often a classification filter).
        # Designs inside it still say Overall when the scores come from that model.
        segment = str(raw.get("name") or item_segment or "Overall")
        groups.append(
            Group(
                name=segment,
                position=int(raw.get("position") or 0),
                segment=segment,
                metric=str(first.get("metric") or "Top Down"),
                builds=builds,
            )
        )
    groups.sort(key=lambda group: group.position)
    if not groups:
        raise ValueError("No saved combination groups were found.")
    return groups

def load_study(study_path: str | Path, analysis_path: str | Path | None = None) -> StudyInfo:
    study = json.loads(Path(study_path).read_text(encoding="utf-8"))
    config = study.get("study_config") or {}
    analysis = _load_analysis(study_path, analysis_path, str(study.get("id") or ""))
    front = (analysis or {}).get("Front Page") or {}
    info = (analysis or {}).get("Information Block") or {}
    overall = (analysis or {}).get("(T) Overall") or {}
    intercepts = (analysis or {}).get("(T) Intercepts") or {}
    bottom_intercepts = (analysis or {}).get("(B) Intercepts") or {}
    response_intercepts = (analysis or {}).get("(R) Intercepts") or {}
    dashboard = (analysis or {}).get("dashboard_summary") or {}
    settings = (analysis or {}).get("analysis_settings") or {}
    bands = ((settings.get("top") or {}).get("hundred") or [4, 5])
    raw_intercept = intercepts.get("intercept")
    main_question = " ".join(str(
        study.get("main_question")
        or front.get("Main Question")
        or info.get("Main Question")
        or ""
    ).split())
    sizes = {"Overall": int(overall.get("base_size") or dashboard.get("totalRespondents") or config.get("number_of_respondents") or 0)}
    gender = ((analysis or {}).get("(T) Gender") or {}).get("segments") or {}
    for label, segment in gender.items():
        if isinstance(segment, dict) and segment.get("base_size"):
            sizes[str(label)] = int(segment["base_size"])
    return StudyInfo(
        title=str(front.get("Title") or study.get("title") or "Design Study"),
        country=str(config.get("country") or "Overall"),
        launched_label=_launched(front.get("Launched At") or study.get("created_at")),
        respondents=sizes["Overall"],
        base_appeal=float(raw_intercept) if raw_intercept is not None else None,
        base_bottom=_optional_float(bottom_intercepts.get("intercept")),
        base_response=_optional_float(response_intercepts.get("intercept")),
        model_name=f"Top-{len(bands)} Box" if len(bands) > 1 else "Top Box",
        aspect=_aspect(config.get("aspect_ratio") or front.get("Aspect Ratio") or "9:16"),
        visual=_named_visual(study.get("study_type")) or _named_visual(info.get("Study Type")) or "layer",
        segment_sizes=sizes,
        main_question=main_question,
        prepared_by=str(front.get("Prepared By") or study.get("prepared_by") or "").strip(),
    )

def _as_groups(payload) -> list[dict]:
    if isinstance(payload, dict):
        payload = payload.get("groups") or payload.get("combinations") or payload.get("items") or []
    if not isinstance(payload, list):
        raise ValueError("Saved combinations JSON must be a list of groups.")
    if payload and isinstance(payload[0], dict) and "items" in payload[0]:
        return payload
    return [{"name": "Overall", "position": 0, "items": payload}]

def _comparison_from_config(config: dict) -> tuple[list[ComparisonSegment], str, str]:
    raw = config.get("comparison") if isinstance(config.get("comparison"), dict) else {}
    source = str(raw.get("source") or "overall").strip().lower()
    if source == "overall":
        return [], "", _audience_label(raw.get("audience"))
    segments: list[ComparisonSegment] = []
    for item in raw.get("segments") or []:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key") or "").strip()
        label = str(item.get("label") or key).strip()
        if not key or not label:
            continue
        try:
            base_size = int(item.get("base_size") or 0)
        except (TypeError, ValueError):
            base_size = 0
        segments.append(ComparisonSegment(key=key, label=label, base_size=base_size))
    if len(segments) < 2:
        return [], "", _audience_label(raw.get("audience"))
    groups = [item for item in (raw.get("groups") or []) if isinstance(item, dict)]
    by_key: dict[str, str] = {}
    for group in groups:
        asked = _tidy_question(str(group.get("question") or group.get("title") or ""))
        for item in group.get("segments") or []:
            if isinstance(item, dict) and item.get("key") and asked:
                by_key[str(item.get("key"))] = asked
    for segment in segments:
        segment.question = by_key.get(segment.key, "")
        if segment.question and segment.label.lower().startswith(segment.question.lower() + ":"):
            segment.label = segment.label.split(":", 1)[1].strip() or segment.label
    titles = [str(item.get("title") or "").strip() for item in groups]
    titles = [item for item in titles if item]
    question = str(raw.get("question") or "").strip()
    if len(titles) > 1:
        title = " · ".join(titles)
    elif source == "classification":
        title = question or (titles[0] if titles else "Comparison")
    elif titles:
        title = titles[0]
    else:
        title = source.capitalize()
    return segments, title, _audience_label(raw.get("audience"))

def _audience_label(audience) -> str:
    if not isinstance(audience, dict):
        return ""
    parts: list[str] = []
    genders = [str(item).strip() for item in audience.get("genders") or [] if str(item).strip()]
    ages = [str(item).strip() for item in audience.get("age_groups") or [] if str(item).strip()]
    if genders and ages:
        parts.append(f"{', '.join(genders).lower()} respondents aged {', '.join(ages)}")
    elif genders:
        parts.append(f"{', '.join(genders).lower()} respondents")
    elif ages:
        parts.append(f"respondents aged {', '.join(ages)}")
    filters = audience.get("classification_filters") or {}
    if isinstance(filters, dict):
        for question, answers in filters.items():
            chosen = [str(item).strip() for item in answers or [] if str(item).strip()]
            if chosen:
                parts.append(f"{question}: {', '.join(chosen)}")
    return f"Among {'; '.join(parts)}" if parts else ""

def _segment_lifts(raw: dict) -> dict[str, float]:
    values = raw.get("segment_lifts") if isinstance(raw.get("segment_lifts"), dict) else {}
    parsed: dict[str, float] = {}
    for key, value in values.items():
        try:
            parsed[str(key)] = float(value)
        except (TypeError, ValueError):
            continue
    return parsed

def _build_from_item(item: dict) -> Build:
    config = item.get("configuration") or {}
    comparison_segments, comparison_title, audience_label = _comparison_from_config(config)
    elements = []
    for raw in config.get("selected_elements") or []:
        name = str(raw.get("name") or "")
        category = str(raw.get("category") or "")
        content = raw.get("content")
        image_url = _http_url(raw.get("image_url")) or _http_url(content)
        if isinstance(content, str) and content.strip() and image_url is None:
            label = re.sub(r"\s+", " ", content).strip()
        else:
            label = _pretty(name)
        elements.append(
            Element(
                name=name,
                label=label,
                lift=float(raw.get("value") or 0),
                category=category,
                z_index=int(raw.get("z_index") or 0),
                image_url=image_url,
                segment_lifts=_segment_lifts(raw),
            )
        )
    elements.sort(key=lambda element: element.z_index)
    total = item.get("total_coefficient")
    if total is None:
        total = sum(element.lift for element in elements)
    return Build(
        position=int(item.get("position") or 0),
        name=str(item.get("name") or ""),
        elements=elements,
        total=float(total),
        metric=str(item.get("metric") or config.get("metric") or "Top Down"),
        study_type=str(item.get("study_type") or config.get("study_type") or ""),
        background_url=_http_url(config.get("background_url")),
        comparison_segments=comparison_segments,
        comparison_title=comparison_title,
        audience_label=audience_label,
    )

def _load_analysis(study_path: str | Path, analysis_path: str | Path | None, study_id: str) -> dict | None:
    candidates = []
    if analysis_path:
        candidates.append(Path(analysis_path))
    study = Path(study_path)
    candidates.append(study.with_name("analysis_data.json"))
    candidates.append(study.parent.parent / "Dafault-export" / "fresh-harvest" / "analysis_data.json")
    for candidate in candidates:
        if not candidate.exists():
            continue
        payload = json.loads(candidate.read_text(encoding="utf-8"))
        if study_id and payload.get("study_id") and payload.get("study_id") != study_id:
            continue
        analysis = payload.get("analysis")
        if isinstance(analysis, dict):
            return analysis
    return None

def _readable_name(name: str) -> str:
    """Drop study codes so a row reads as a name, not a file path."""
    text = _pretty(name).replace("_", " ").replace("-", " ")
    text = re.sub(r"\s+", " ", text).strip()
    parts = text.split()
    kept = [part for part in parts if not re.fullmatch(r"[A-Za-z]?\d+[A-Za-z]?", part)]
    return " ".join(kept) or text or name

def _fit_lines(text: str, width_in: float, size: float, lines: int, bold: bool = False) -> str:
    """Wrap into a fixed number of lines. A hyphen is a break, and the last line ends in the cell."""
    raw = re.sub(r"\s+", " ", str(text or "")).strip()
    if not raw or lines < 1 or width_in <= 0.15:
        return ""
    factor = 0.66 if bold else 0.58
    per = max(4, int(width_in * 72 / (size * factor)))
    words: list[str] = []
    for word in raw.split(" "):
        parts = [part for part in re.split(r"(?<=-)", word) if part]
        words.extend(parts or [word])
    rows: list[str] = []
    current = ""
    index = 0
    while index < len(words) and len(rows) < lines:
        word = words[index]
        if len(word) > per:
            word = word[: max(1, per - 1)].rstrip("-") + "…"
        trial = word if not current else f"{current} {word}"
        if len(trial) <= per:
            current = trial
            index += 1
            continue
        if len(rows) == lines - 1:
            room = per - (len(current) + 1 if current else 0)
            if room >= 2:
                current = (f"{current} " if current else "") + word[: room - 1].rstrip("-") + "…"
            elif current:
                current = current[: max(1, per - 1)].rstrip() + "…"
            index = len(words)
            break
        if current:
            rows.append(current)
        current = ""
        if len(word) <= per:
            current = word
            index += 1
    if current and len(rows) < lines:
        rows.append(current)
    return "\n".join(rows)


def _tighten(shape) -> None:
    """Keep wrapped lines inside the box. Default paragraph spacing was pushing text into the next cell."""
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    for paragraph in frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)


def _clip_lines(text: str, width_in: float, size: float = 11, lines: int = 2, readable: bool = False) -> str:
    """Wrap into a fixed number of lines so the text stays inside its box."""
    source = _readable_name(text) if readable else str(text or "")
    words = re.sub(r"\s+", " ", source.replace("_", " ")).strip().split() or [""]
    per_line = max(6, int(width_in * 72 / (size * 0.52)))
    rows: list[str] = []
    current = ""
    leftover = False
    for word in words:
        piece = word if len(word) <= per_line else word[: max(1, per_line - 1)].rstrip() + "…"
        trial = piece if not current else f"{current} {piece}"
        if len(trial) <= per_line:
            current = trial
            continue
        if len(rows) >= lines - 1:
            leftover = True
            break
        if current:
            rows.append(current)
        current = piece
    if current and len(rows) < lines:
        rows.append(current)
    if leftover and rows:
        rows[-1] = rows[-1][: max(1, per_line - 1)].rstrip() + "…"
    return "\n".join(rows[:lines])

def _combination_name(build: Build, index: int) -> str:
    """A short name for the chart and the thumbnail. Element codes stay on the detail slide."""
    name = re.sub(r"\s+", " ", (build.name or "").replace("_", " ")).strip()
    if name and "·" not in name and len(name) <= 28:
        return name
    return f"Combination {index + 1}"

def _tidy_question(text: str) -> str:
    """The question a column belongs to, written so the answer under it makes sense."""
    cleaned = re.sub(r"\s+", " ", str(text or "").replace("_", " ")).strip().rstrip(".")
    if not cleaned:
        return ""
    cleaned = re.sub(r"\benglish\b", "English", cleaned, flags=re.IGNORECASE)
    if cleaned[0].islower():
        cleaned = cleaned[0].upper() + cleaned[1:]
    if cleaned.lower() in {"age", "gender", "overall"}:
        return cleaned[:1].upper() + cleaned[1:].lower()
    if not cleaned.endswith("?") and re.match(r"^(can|which|do|does|did|what|where|who|how|is|are)\b", cleaned, re.I):
        cleaned += "?"
    return cleaned

def _human_category(name: str) -> str:
    """Turn a saved category code into words. Age ranges keep their hyphen."""
    text = str(name or "").strip()
    if not text:
        return ""
    text = text.replace("_", " ").replace("=", " · ")
    text = re.sub(r"(?<=\D)-(?=\D)", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    words = []
    for word in text.split(" "):
        if word == "·" or re.fullmatch(r"\d[\d\-–]*", word):
            words.append(word)
        else:
            words.append(word[:1].upper() + word[1:])
    return " ".join(words)

_QUESTION_LEAD = (
    "which of these",
    "which of the",
    "what kind of",
    "what type of",
    "how often do you",
    "how many",
    "how often",
    "do you",
    "can you",
    "are you",
    "have you",
    "did you",
    "would you",
    "which",
    "what",
)
_QUESTION_SKIP = {"of", "the", "these", "those", "a", "an", "do", "you", "your", "please"}

def _condense_question(question: str) -> str:
    """Keep the question's own words. Only the filler at the start comes off."""
    raw = re.sub(r"\s+", " ", str(question or "")).strip()
    if not raw:
        return ""
    asking = raw.endswith("?")
    core = raw.rstrip("?").strip()
    changed = True
    while changed and core:
        changed = False
        lower = core.lower()
        for phrase in _QUESTION_LEAD:
            if lower == phrase or lower.startswith(phrase + " "):
                core = core[len(phrase):].strip(" ?")
                changed = True
                break
    if core.lower().startswith("is "):
        core = core[3:].strip()
    words = [word for word in core.split() if word.lower() not in _QUESTION_SKIP]
    core = " ".join(words).strip() or raw.rstrip("?").strip()
    if core[0].islower():
        core = core[0].upper() + core[1:]
    if asking and not core.endswith("?"):
        core += "?"
    return core

def _question_spans(segments: list[ComparisonSegment], stride: float) -> list[tuple[int, int, str, int, float]]:
    """Each run of columns that share a question, already shortened to the width it has."""
    spans = []
    start = 0
    while start < len(segments):
        question = segments[start].question
        end = start + 1
        while end < len(segments) and segments[end].question == question and question:
            end += 1
        width = max((end - start) * stride - 0.08, 0.7)
        text = _condense_question(question) if question else ""
        per_line = max(8, int(width * 72 / (11 * 0.52)))
        lines = 2 if text and len(text) > per_line else 1
        spans.append((start, end, text, lines, width))
        start = end
    return spans

def _answer_with_count(answer: str, base_size: int, width_in: float, size: float = 10) -> str:
    suffix = f" (n={base_size})" if base_size else ""
    per_line = max(6, int(width_in * 72 / (size * 0.52)))
    if len(answer) + len(suffix) <= per_line:
        return f"{answer}{suffix}"
    room = per_line - len(suffix)
    if room < 2:
        return suffix.strip()
    return answer[: room - 1].rstrip() + "…" + suffix

def _respondent_count(count: int) -> str:
    if count > 0:
        return f"(n={count})"
    return ""

def _segment_answer(segment: ComparisonSegment) -> str:
    label = segment.label.strip()
    question = segment.question.strip()
    if question and label.lower().startswith(question.lower() + ":"):
        label = label.split(":", 1)[1].strip()
    return label or segment.label

def _pretty(name: str) -> str:
    text = name.strip()
    for prefix in (
        "Background Layer ",
        "HEADLINE - ",
        "Headline - ",
        "Sub-Header - ",
        "Sub Header - ",
        "Image Layer_ ",
        "Image Layer ",
        "CTA - ",
    ):
        if text.startswith(prefix):
            text = text[len(prefix):]
            break
    return re.sub(r"\s+", " ", text).strip() or name

def _aspect(value: str) -> float:
    match = re.search(r"(\d+(?:\.\d+)?)\s*[:/]\s*(\d+(?:\.\d+)?)", str(value))
    if not match:
        return 9 / 16
    width, height = float(match.group(1)), float(match.group(2))
    if height == 0:
        return 9 / 16
    return width / height

def _launched(value) -> str:
    if not value:
        return ""
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return ""
    return parsed.strftime("%B %Y")

def _download(groups: list[Group], cache: Path) -> None:
    cache.mkdir(parents=True, exist_ok=True)
    jobs: list[tuple[str, Path, str]] = []
    seen: set[str] = set()

    def add(url: str | None):
        if not url or url in seen:
            return
        seen.add(url)
        leaf = url.rstrip("/").split("/")[-1]
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", leaf)[:120] or "asset"
        if not safe.lower().endswith(".png"):
            safe = f"{safe}.png"
        jobs.append((url, cache / safe, url))

    for group in groups:
        for build in group.builds:
            add(build.background_url)
            for element in build.elements:
                add(element.image_url)

    def fetch(item: tuple[str, Path, str]):
        key, dest, url = item
        fetch_image(dest, url)
        return key, dest

    saved: dict[str, Path] = {}
    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = [pool.submit(fetch, job) for job in jobs]
        for future in as_completed(futures):
            try:
                key, dest = future.result()
            except Exception as exc:
                print(f"  image skipped: {exc}")
                continue
            saved[key] = dest
    for group in groups:
        for build in group.builds:
            if build.background_url:
                build.background_path = saved.get(build.background_url)
            for element in build.elements:
                if element.image_url:
                    element.image_path = saved.get(element.image_url)

def _group_slides(presentation, study: StudyInfo, group: Group, builds: list[Build], kind: str):
    meta = _KIND[kind]
    slides = [("plain", _market(presentation, group, kind))]
    slides.extend(("content", slide) for slide in _compare(presentation, study, group, builds, kind))
    for index, build in enumerate(builds):
        headline = _rank_headline(kind, index, study.visual)
        slides.append(("content", _pack(presentation, study, group, build, headline, kind)))
    return slides

def _blank(presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])

def _flush(box):
    frame = box.text_frame
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    return box

def _text(slide, text, x, y, w, h, **kwargs):
    return _flush(add_text(slide, text, x, y, w, h, **kwargs))

def _hairline(slide, x1, y1, x2, y2, color: RGBColor = RGBColor(0xF2, 0xC6, 0xCF), width: float = 0.75):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line

def _rail(slide):
    add_rect(slide, 0, 0, 0.47, SLIDE_H, RAIL)

def _page_number(slide, number: int):
    _text(slide, str(number), 12.45, 7.16, 0.7, 0.22, size=11, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)

def _report_line(study: StudyInfo) -> str:
    return f"{study.title}  |  {study.launched_label}".strip(" |")

_PREPARED_BY_DEFAULT = "J Brown Fitterman | jbrown@mindgenomicsassociates.com"

def _prepared_by_line(study: StudyInfo) -> str:
    credit = (study.prepared_by or "").strip() or _PREPARED_BY_DEFAULT
    return f"Prepared by: {credit}"

def _cover_headline(study: StudyInfo) -> str:
    question = (study.main_question or "").strip()
    return question or "Design Combination Readout v2"

def _marks(slide):
    mga = ASSETS / "mga_logo.png"
    _text(slide, "MGA", 1.15, 2.45, 2.4, 0.5, size=28, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    if mga.exists():
        slide.shapes.add_picture(str(mga), Inches(1.377), Inches(3.55), Inches(1.795), Inches(1.795))

def _cover(presentation, study: StudyInfo):
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    _marks(slide)
    _hairline(slide, 4.2464, 1.3478, 4.2464, 1.3478 + 4.7826)
    _hairline(slide, 4.5996, 3.75, 4.5996 + 6.8551, 3.75)
    headline = _cover_headline(study)
    headline_size = 28 if len(headline) > 70 else 32 if len(headline) > 40 else 36
    headline_height = 1.55 if len(headline) > 40 else 0.85
    _text(slide, headline, 4.60, 1.98, 8.48, headline_height, size=headline_size, bold=True, color=NAVY)
    _text(slide, _report_line(study), 4.60, 4.24, 8.2, 0.37, size=16, color=NAVY)
    _text(
        slide,
        _prepared_by_line(study),
        4.60, 4.63, 7.64, 0.37, size=16, color=NAVY,
    )
    _text(slide, "Mind Genomics Associates, Inc", 4.60, 5.01, 4.2, 0.37, size=16, color=NAVY)
    _text(slide, study.country, 4.60, 5.71, 8.0, 0.44, size=16, color=NAVY)
    return slide

def _thanks(presentation, study: StudyInfo):
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    _marks(slide)
    _hairline(slide, 4.2464, 1.3478, 4.2464, 1.3478 + 4.7826)
    _hairline(slide, 4.5996, 3.75, 4.5996 + 6.8551, 3.75)
    _text(slide, "Thank you", 4.60, 1.98, 7.8, 0.7, size=36, bold=True, color=NAVY)
    _text(slide, "We look forward to exploring more combinations", 4.61, 2.82, 7.6, 0.4, size=16, color=NAVY)
    _text(slide, _report_line(study), 4.60, 4.24, 8.2, 0.37, size=16, color=NAVY)
    _text(
        slide,
        _prepared_by_line(study),
        4.60, 4.63, 7.64, 0.37, size=16, color=NAVY,
    )
    _text(slide, "Mind Genomics Associates, Inc", 4.60, 5.01, 4.2, 0.37, size=16, color=NAVY)
    _text(slide, study.country, 4.60, 5.71, 8.0, 0.44, size=16, color=NAVY)
    return slide

def _section(presentation, title: str):
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    add_rect(slide, 0, 0, 0.18, SLIDE_H, PINK)
    _text(slide, title, 1.15, 2.15, 11.4, 3.2, size=48, bold=True, color=NAVY)
    return slide

def _market(presentation, group: Group, kind: str):
    slide = _blank(presentation)
    set_background(slide, BLUSH)
    title = _human_category(group.segment) or "Saved combinations"
    size = 26 if len(title) > 42 else 36
    _text(
        slide, title,
        0.9, 2.35, 11.5, 1.8,
        size=size, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
    )
    _text(
        slide,
        "Next: the combinations compared as wholes, then each combination opened element by element.",
        1.6, 4.4, 10.1, 0.7,
        size=16, color=GRAY, align=PP_ALIGN.CENTER,
    )
    return slide

def _how_to_read(presentation, study: StudyInfo, groups: list[Group]):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _rail(slide)
    _text(slide, "HOW TO READ THIS REPORT: UNDERSTANDING ELEMENT LIFTS", 0.83, 0.33, 12.0, 0.28, size=14, bold=True, color=TEAL)
    add_rect(slide, 0.83, 0.68, 12.0, 0.46, CARD, rounded=True, radius=0.08)
    _text(
        slide,
        "Combinations are created by selecting the strongest option in each element category",
        0.98, 0.74, 11.7, 0.34, size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, 0.83, 1.32, 6.05, 5.45, WHITE, rounded=True, radius=0.04, line=RGBColor(0xF0, 0xE4, 0xEA))
    _text(slide, "How the model works", 1.05, 1.48, 5.6, 0.32, size=16, bold=True, color=NAVY)
    _text(slide, _model_copy(study.visual), 1.05, 1.90, 5.6, 4.5, size=14, color=INK)
    add_rect(slide, 7.08, 1.32, 5.75, 5.45, CARD, rounded=True, radius=0.04)
    add_rect(slide, 7.08, 1.32, 5.75, 0.1, TEAL)
    _text(slide, "Example from the saved builds", 7.28, 1.55, 5.35, 0.4, size=16, bold=True, color=NAVY)
    _text(slide, _example_copy(groups, study.visual), 7.28, 2.15, 5.35, 4.3, size=14, color=INK)
    return slide

def _slot_words(visual: str) -> tuple[str, str]:
    if visual == "text":
        return "statement", "statements"
    if visual == "grid":
        return "image", "images"
    return "layer", "layers"

def _model_copy(visual: str) -> str:
    one, many = _slot_words(visual)
    if visual == "grid":
        opening = "Respondents rated grids of images — one image from each set. "
        closing = (
            "A saved combination keeps one image from each set. The images sit side by side. They are not stacked. "
            "Swapping one image only changes that set. Everything else stays, because each score was measured on its own."
        )
    elif visual == "text":
        opening = "Respondents rated mixes of statements. "
        closing = (
            "A saved combination keeps one statement from each group, written in full. "
            "Swapping one statement only changes that line. Everything else stays, because each score was measured on its own."
        )
    else:
        opening = "Respondents rated a large number of sign combinations — each showing different design elements. "
        closing = (
            "A saved combination keeps one option from each layer. Swapping one element only changes that slot. "
            "Everything else stays, because each score was measured on its own."
        )
    return (
        f"{opening}"
        "The model then separates those ratings into an individual lift score.\n\n"
        f"Each lift score shows how much that {one} raises or lowers appeal relative to the model baseline.\n\n"
        "The comparison slide shows each saved combination and its total lift: every element in that combination added together. A short list shows the packs side by side. A longer list ranks them, highest first.\n\n"
        "The following slides open one combination. Each row is one element. A column is one respondent group, named with the question, the answer, and how many people are in it.\n\n"
        f"{closing}"
    )

def _example_copy(groups: list[Group], visual: str = "layer") -> str:
    one, many = _slot_words(visual)
    pair = _swap_example(groups)
    if pair is None:
        return (
            f"Each row in the build slides is one {one}. The number beside it is that {one}’s lift. "
            "The total is the sum of the selected lifts."
        )
    group, higher, lower, index = pair
    changed_high = higher.elements[index]
    changed_low = lower.elements[index]
    return (
        f"In {group.segment}, one build uses {changed_high.label} ({_lift(changed_high.lift)}) in {changed_high.category}.\n\n"
        f"Swapping to {changed_low.label} ({_lift(changed_low.lift)}) only changes that slot. "
        f"The other {many} stay as they are.\n\n"
        f"The result: {changed_low.label} totals {_lift(lower.total)} against {_lift(higher.total)} "
        f"for {changed_high.label} — a difference of {abs(higher.total - lower.total):.1f} points."
    )

def _swap_example(groups: list[Group]):
    found = None
    for group in groups:
        for left_index, left in enumerate(group.builds):
            for right in group.builds[left_index + 1:]:
                if _metric_kind(left.metric) != _metric_kind(right.metric):
                    continue
                if len(left.elements) != len(right.elements):
                    continue
                diffs = [i for i, (a, b) in enumerate(zip(left.elements, right.elements)) if a.name != b.name]
                if len(diffs) != 1:
                    continue
                gap = abs(left.total - right.total)
                higher, lower = (left, right) if left.total >= right.total else (right, left)
                candidate = (gap > 0, gap, group, higher, lower, diffs[0])
                if found is None or candidate[:2] > found[:2]:
                    found = candidate
    if found is None:
        return None
    _, _, group, higher, lower, index = found
    return group, higher, lower, index

def _compare(presentation, study: StudyInfo, group: Group, builds: list[Build], kind: str) -> list:
    """A ranked bar chart with the packs, then build cards. Each card slide holds one or two combinations."""
    indexed = list(enumerate(builds))
    if not indexed:
        return [_compare_stage(presentation, study, group, [], builds, kind, 0, 1)]
    slides = []
    ranked = sorted(indexed, key=lambda item: (-item[1].total, item[0]))
    board_size = 6
    board_count = max(1, (len(ranked) + board_size - 1) // board_size)
    for page_index, start in enumerate(range(0, len(ranked), board_size)):
        slides.append(
            _score_board(
                presentation, study, group, ranked[start:start + board_size], builds, kind, page_index, board_count,
            )
        )
    page_size = 2
    page_count = max(1, (len(indexed) + page_size - 1) // page_size)
    for page_index, start in enumerate(range(0, len(indexed), page_size)):
        page = indexed[start:start + page_size]
        slides.append(
            _compare_stage(presentation, study, group, page, builds, kind, page_index, page_count)
        )
    return slides


def _compare_stage(
    presentation, study: StudyInfo, group: Group, page: list[tuple[int, Build]],
    builds: list[Build], kind: str, page_index: int, page_count: int,
):
    """One slide of the pair. `page` is (saved index, combination)."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _rail(slide)
    digits = _digits(kind)
    context = _compare_context(group, builds)
    if page_count > 1 and page:
        start = page[0][0] + 1
        end = page[-1][0] + 1
        if start == end:
            place = f"Combination {start} of {len(builds)}"
        else:
            place = f"Combinations {start}–{end} of {len(builds)}"
        context = f"{context}. {place}" if context else place
    _buildup_heading(slide, kind, builds, digits, context)
    _buildup_cards(slide, study, page, builds, kind, digits)
    _buildup_key(slide, builds, kind, digits, page_index < page_count - 1)
    return slide


_BAR_COLORS = (
    RGBColor(0x1A, 0x6A, 0x78),
    RGBColor(0xE0, 0x7A, 0x2F),
    RGBColor(0x1F, 0x8A, 0x4C),
    RGBColor(0x6B, 0x3F, 0xA0),
    RGBColor(0x2F, 0x8F, 0x57),
    RGBColor(0x7D, 0xC8, 0x6A),
)


def _score_board(
    presentation, study: StudyInfo, group: Group, page: list[tuple[int, Build]],
    builds: list[Build], kind: str, page_index: int, page_count: int,
):
    """Ranked bars on the left, the same combinations as a pack grid on the right."""
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _rail(slide)
    digits = _digits(kind)
    title = _human_category(group.segment) or "Saved combinations"
    _text(
        slide, f"KEY FINDINGS: {title.upper()}",
        0.62, 0.16, 7.7, 0.22, size=11, bold=True, color=NAVY,
    )
    heading = _text(slide, title, 0.62, 0.40, 7.7, 0.62, size=26, bold=True, color=NAVY)
    heading.text_frame.word_wrap = True
    _tighten(heading)
    _score_chart(slide, page, digits, 0.42, 1.22, 7.7, 5.15, title)
    _score_thumbs(slide, study, page, digits, 8.32, 1.22, 4.70, 5.15)
    places = "three decimals" if digits >= 3 else "one decimal"
    sample = f"n={study.respondents} {study.country} respondents. " if study.respondents else ""
    note = f"{sample}{study.model_name} model: estimated total combination lift above baseline. Values rounded to {places}."
    if page_count > 1:
        note = f"{note} {page_index + 1} of {page_count}."
    footer = _text(slide, note, 0.62, 6.95, 7.7, 0.36, size=10, color=GRAY)
    footer.text_frame.word_wrap = True
    _tighten(footer)
    return slide


def _score_chart(slide, page: list[tuple[int, Build]], digits: int, x, y, w, h, banner: str) -> None:
    banner_h = 0.28
    add_rect(slide, x + 1.35, y, max(2.2, w - 1.55), banner_h, _BUILD_HEADER)
    label = _text(
        slide, banner.upper(),
        x + 1.35, y, max(2.2, w - 1.55), banner_h,
        size=11, bold=True, color=_BUILD_ROLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(label)
    if not page:
        return
    data = CategoryChartData()
    data.categories = [_keep_line(_combination_name(build, index), 1.15, 10, lines=2) for index, build in page]
    data.add_series("Lift", [build.total for _index, build in page])
    chart_top = y + banner_h + 0.10
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(chart_top), Inches(w), Inches(h - (chart_top - y)), data,
    )
    chart = chart_shape.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.vary_by_categories = True
    plot.has_data_labels = True
    plot.gap_width = 50
    labels = plot.data_labels
    labels.font.size = Pt(11)
    labels.font.bold = True
    labels.font.name = FONT
    labels.font.color.rgb = NAVY
    labels.number_format = "0.000" if digits >= 3 else "0.0"
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = chart.series[0]
    for index, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _BAR_COLORS[index % len(_BAR_COLORS)]
    chart.category_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.color.rgb = NAVY
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.has_minor_gridlines = False
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.color.rgb = GRAY
    peak = max((build.total for _index, build in page), default=1)
    _set_axis_bounds(chart.value_axis, 0 if min(build.total for _index, build in page) >= 0 else None, _nice_ceiling(max(peak * 1.18, 1)))
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.paragraphs[0].text = "Estimated total lift above baseline"
    for paragraph in chart.value_axis.axis_title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(11)
            run.font.color.rgb = GRAY
            run.font.name = FONT
    _quiet_chart_frame(chart)


def _set_axis_bounds(axis, minimum, maximum) -> None:
    scaling = axis._element.find(qn("c:scaling"))
    if scaling is None:
        return
    for tag, value in (("c:min", minimum), ("c:max", maximum)):
        if value is None:
            continue
        node = scaling.find(qn(tag))
        if node is None:
            node = scaling.makeelement(qn(tag), {})
            scaling.append(node)
        node.set("val", str(value))


def _quiet_chart_frame(chart) -> None:
    chart_element = chart._element
    plot_area = chart_element.find(qn("c:chart")).find(qn("c:plotArea"))
    sp_pr = plot_area.find(qn("c:spPr"))
    if sp_pr is None:
        return
    for child in list(sp_pr):
        if child.tag == qn("a:ln"):
            sp_pr.remove(child)


def _score_thumbs(slide, study: StudyInfo, page: list[tuple[int, Build]], digits: int, x, y, w, h) -> None:
    """Packs in two columns, centered in the card. One combination sits alone."""
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.04, line=_BUILD_LINE)
    if not page:
        return
    count = len(page)
    columns = 1 if count == 1 else 2
    rows = math.ceil(count / columns)
    pad_x, pad_y = 0.18, 0.16
    gap_x, gap_y = 0.16, 0.14
    inner_w = w - pad_x * 2
    inner_h = h - pad_y * 2
    cell_w = (inner_w - gap_x * (columns - 1)) / columns
    name_h, score_h = 0.42, 0.30
    row_budget = (inner_h - gap_y * (rows - 1)) / rows
    image_h = min(cell_w * 1.85, max(0.95, row_budget - name_h - score_h))
    block_h = rows * (image_h + name_h + score_h) + gap_y * (rows - 1)
    origin_y = y + pad_y + max(0.0, (inner_h - block_h) / 2)
    shown = [build for _index, build in page]
    for offset, (index, build) in enumerate(page):
        row = offset // columns
        column = offset % columns
        row_count = min(columns, count - row * columns)
        row_w = row_count * cell_w + (row_count - 1) * gap_x
        origin_x = x + pad_x + (inner_w - row_w) / 2
        left = origin_x + column * (cell_w + gap_x)
        top = origin_y + row * (image_h + name_h + score_h + gap_y)
        _card_visual(slide, study, shown, build, left, top, cell_w, image_h, anchor="bottom")
        name = _keep_line(_combination_name(build, index), cell_w - 0.06, 12, lines=2, bold=True)
        name_box = _text(
            slide, name, left, top + image_h, cell_w, name_h,
            size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        _tighten(name_box)
        _fill, ink, _bar = _score_tone(build.total, digits)
        _text(
            slide, _lift(build.total, digits), left, top + image_h + name_h, cell_w, score_h,
            size=16, bold=True, color=ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )


def _compare_list(
    presentation, study, group, builds, kind, page, start_rank: int, total: int, page_number: int, pages: int,
):
    slide = _blank(presentation)
    set_background(slide, WHITE)
    _rail(slide)
    digits = _digits(kind)
    insight = _compare_insight(builds, kind, digits)
    context = _compare_context(group, builds)
    end_rank = start_rank + len(page) - 1
    if pages > 1:
        order = f"{start_rank}–{end_rank} of {total}, highest first"
    else:
        order = "Highest first"
    context = f"{context}. {order}" if context else order
    top = _draw_finding(slide, kind, insight, context)
    _ranked_rows(slide, study, builds, page, kind, 0.83, top, 12.20, max(2.15, 6.90 - top), start_rank)
    _compare_footer(slide)
    return slide


def _compare_footer(slide) -> None:
    footer = _text(
        slide,
        "A + score is green. A − score is red. The next slides open each combination.",
        0.83, 7.08, 11.35, 0.28, size=12, color=GRAY,
    )
    footer.text_frame.word_wrap = True

def _compare_context(group: Group, builds: list[Build]) -> str:
    """Where these combinations were saved, in words a reader does not have to decode."""
    parts: list[str] = []
    category = _human_category(group.segment)
    if category and category.lower() not in {"overall", "saved combinations"}:
        parts.append(f"Saved in {category}")
    audience = next((build.audience_label for build in builds if build.audience_label), "")
    if audience:
        parts.append(audience.rstrip("."))
    return ". ".join(parts)

def _segment_total(build: Build, key: str) -> float:
    if not build.elements:
        return 0.0
    return sum(float(element.segment_lifts.get(key, 0.0)) for element in build.elements)

_MINT = RGBColor(0xF3, 0xFA, 0xF6)
_CARD_LINE = RGBColor(0xF0, 0xE4, 0xEA)
_PILL = RGBColor(0xF6, 0xEE, 0xF2)
_MAT_LINE = RGBColor(0xE4, 0xEE, 0xE8)

def _tie_cutoff(digits: int) -> float:
    return 0.0005 if digits >= 3 else 0.05

def _leader_indexes(builds: list[Build], digits: int) -> set[int]:
    if not builds:
        return set()
    best = max(build.total for build in builds)
    cutoff = _tie_cutoff(digits)
    return {index for index, build in enumerate(builds) if abs(build.total - best) <= cutoff}

def _gap_figure(gap: float, digits: int) -> str:
    return f"{abs(gap):.{digits}f}"

def _gap_words(gap: float, digits: int) -> str:
    figure = _gap_figure(gap, digits)
    if digits >= 3:
        return figure
    unit = "point" if abs(abs(gap) - 1) < 0.05 else "points"
    return f"{figure} {unit}"

def _name_list(builds: list[Build], indexes: set[int]) -> str:
    names = [_combination_name(builds[index], index) for index in sorted(indexes)]
    if len(names) <= 1:
        return names[0] if names else ""
    if len(names) == 2:
        return f"{names[0]} and {names[1]}"
    return ", ".join(names[:-1]) + f", and {names[-1]}"

def _compare_insight(builds: list[Build], kind: str, digits: int) -> str:
    """The one sentence a reader should remember from this slide."""
    if not builds:
        return "No combinations are saved in this group."
    if len(builds) == 1:
        name = _combination_name(builds[0], 0)
        return f"{name} is the only combination saved here, at {_lift(builds[0].total, digits)}."
    leaders = _leader_indexes(builds, digits)
    best = max(builds, key=lambda build: build.total)
    worst = min(builds, key=lambda build: build.total)
    gap = best.total - worst.total
    if len(leaders) == len(builds):
        if len(builds) == 2:
            return f"The two combinations are tied at {_lift(best.total, digits)}."
        return f"All {len(builds)} combinations are tied at {_lift(best.total, digits)}."
    if len(leaders) > 1:
        return f"{_name_list(builds, leaders)} are tied for the highest total, at {_lift(best.total, digits)}."
    index = next(iter(leaders))
    name = _combination_name(builds[index], index)
    ordered = sorted(builds, key=lambda item: item.total, reverse=True)
    lead = ordered[0].total - ordered[1].total
    span = gap
    if len(builds) > 2 and span - lead > _tie_cutoff(digits):
        words = _gap_words(span, digits)
        if kind == "bottom":
            return f"{name} is the highest of this group. {words} separate the highest and lowest."
        if kind == "response":
            return f"{name} has the highest total. {words} separate the highest and lowest."
        return f"{name} is highest. {words} separate the highest and lowest."
    words = _gap_words(lead, digits)
    if kind == "response":
        return f"{name} has the highest total, ahead by {words}."
    if kind == "bottom":
        return f"{name} is the highest of this group, by {words}."
    return f"{name} leads by {words}."

def _score_caption(build: Build, builds: list[Build], digits: int, index: int, leaders: set[int]) -> str:
    if len(builds) < 2:
        return "Saved combination"
    if index in leaders and len(leaders) == 1:
        return "Highest"
    if index in leaders:
        return "Tied for highest"
    best = max(item.total for item in builds)
    return f"{_gap_words(best - build.total, digits)} behind"

def _draw_finding(slide, kind: str, insight: str, context: str) -> float:
    """Eyebrow, title, the finding, and where the combinations were saved."""
    meta = _KIND[kind]
    _text(
        slide, f"KEY FINDINGS: {meta['family']}",
        0.83, 0.24, 12.0, 0.22, size=12, bold=True, color=TEAL,
    )
    title = _text(
        slide, "How the combinations compare",
        0.83, 0.46, 12.0, 0.40, size=26, bold=True, color=NAVY,
    )
    title.text_frame.word_wrap = True
    y = 0.92
    insight_h = 0.56 if len(insight) > 74 else 0.36
    finding = _text(slide, insight, 0.83, y, 12.2, insight_h, size=16, bold=True, color=NAVY)
    finding.text_frame.word_wrap = True
    if context:
        note = _text(slide, context, 0.83, y + insight_h + 0.02, 12.2, 0.26, size=13, color=GRAY)
        note.text_frame.word_wrap = True
        return y + insight_h + 0.36
    return y + insight_h + 0.14

def _scores_are_close(values: list[float]) -> bool:
    """True when a column from zero would make the totals look the same."""
    if len(values) < 2:
        return False
    low, high = min(values), max(values)
    if math.isclose(low, high, abs_tol=1e-9):
        return True
    if low >= 0 and high > 0:
        return low / high >= 0.92
    if high <= 0 and low < 0:
        return abs(high) / abs(low) >= 0.92
    return False

def _window_step(span: float) -> float:
    span = abs(span)
    if span <= 0.02:
        return 0.005
    if span <= 0.08:
        return 0.01
    if span <= 0.4:
        return 0.05
    if span <= 1.2:
        return 0.2
    if span <= 3:
        return 0.5
    if span <= 8:
        return 1
    if span <= 20:
        return 2
    return 5

def _snap_floor(value: float, step: float) -> float:
    return round(math.floor(value / step + 1e-9) * step, 6)

def _snap_ceil(value: float, step: float) -> float:
    return round(math.ceil(value / step - 1e-9) * step, 6)

def _score_window(values: list[float]) -> tuple[float, float, bool]:
    """(low, high, zoomed). Zoomed windows start near the scores, not at zero."""
    if not values:
        return 0.0, 1.0, False
    low, high = min(values), max(values)
    if _scores_are_close(values):
        span = high - low
        pad = span * 0.7 if span > 0 else max(abs(high) * 0.04, 0.4)
        raw_low = low - pad
        raw_high = high + pad
        if low >= 0:
            raw_low = max(0.0, raw_low)
        if high <= 0:
            raw_high = min(0.0, raw_high)
        if raw_high <= raw_low:
            raw_high = raw_low + max(pad, 0.5)
        step = _window_step(raw_high - raw_low)
        snapped_low = _snap_floor(raw_low, step)
        snapped_high = _snap_ceil(raw_high, step)
        if low >= 0:
            snapped_low = max(0.0, snapped_low)
        if snapped_high <= snapped_low:
            snapped_high = snapped_low + step
        return snapped_low, snapped_high, True
    if low >= 0:
        ceiling = _nice_ceiling(high * 1.12) if high > 0 else 1.0
        return 0.0, max(ceiling, high, 1.0), False
    if high <= 0:
        floor = _nice_ceiling(abs(low) * 1.12) if low < 0 else 1.0
        return -max(floor, abs(low), 1.0), 0.0, False
    span = max(high - low, 1.0)
    pad = max(span * 0.08, 0.5)
    step = _window_step(span + pad * 2)
    return _snap_floor(low - pad, step), _snap_ceil(high + pad, step), False

def _needs_ruler(values: list[float], zoomed: bool) -> bool:
    """A position scale, used when two or three totals are close or cross zero."""
    if len(values) < 2 or len(values) > 3:
        return False
    if zoomed:
        return True
    low, high = min(values), max(values)
    return low < 0 < high

def _axis_label(value: float, digits: int) -> str:
    if abs(value - round(value)) < 1e-6:
        return str(int(round(value)))
    return f"{value:.{max(digits, 1)}f}".rstrip("0").rstrip(".")

def _oval(slide, x, y, w, h, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape

def _draw_meter(slide, x, y, w, h, value: float, low: float, high: float, fill: RGBColor) -> None:
    """Length is how high the score sits between low and high. Longer means higher."""
    if w <= 0 or h <= 0:
        return
    add_rect(slide, x, y, w, h, TRACK, rounded=True, radius=0.5)
    span = high - low
    if span <= 0:
        return
    fraction = min(1.0, max(0.0, (value - low) / span))
    if fraction <= 0.012:
        return
    add_rect(slide, x, y, max(0.08, w * fraction), h, fill, rounded=True, radius=0.5)

_BUILD_HEADER = RGBColor(0xF7, 0xE4, 0xEC)
_BUILD_LINE = RGBColor(0xE2, 0xE7, 0xF0)
_BUILD_BAND = RGBColor(0xFB, 0xF6, 0xF8)
_BUILD_ROLE = RGBColor(0x8A, 0x45, 0x60)
_PILL_POS = RGBColor(0x1F, 0xA3, 0x4A)
_PILL_NEG = RGBColor(0xDC, 0x3D, 0x3D)
_PILL_ZERO = RGBColor(0x9A, 0xA6, 0xB4)
_KEY_FILL = RGBColor(0xFC, 0xF0, 0xF4)


def _buildup_heading(slide, kind: str, builds: list[Build], digits: int, context: str) -> None:
    meta = _KIND[kind]
    _text(
        slide, f"KEY FINDINGS: {meta['family']}",
        0.68, 0.16, 12.2, 0.20, size=12, bold=True, color=TEAL,
    )
    title = _text(
        slide, _buildup_title(builds, digits),
        0.68, 0.36, 12.2, 0.38, size=24, bold=True, color=NAVY,
    )
    title.text_frame.word_wrap = True
    _tighten(title)
    note = "Each card adds the measured lift from the elements in that combination."
    if context:
        note = f"{note} {context}."
    subtitle = _text(slide, note, 0.68, 0.76, 12.2, 0.40, size=13, color=GRAY)
    subtitle.text_frame.word_wrap = True
    _tighten(subtitle)


def _buildup_title(builds: list[Build], digits: int) -> str:
    if len(builds) <= 1:
        return "The elements in this combination"
    if len(_leader_indexes(builds, digits)) == len(builds):
        return "These combinations land on the same total"
    return "Higher-lift elements make the stronger combination"


def _build_role(index: int, builds: list[Build], digits: int) -> str:
    if len(builds) <= 1:
        return "SAVED COMBINATION"
    leaders = _leader_indexes(builds, digits)
    if len(leaders) == len(builds):
        return "TIED BUILD"
    if len(builds) == 2:
        return "HIGHER-LIFT BUILD" if index in leaders else "ALTERNATIVE BUILD"
    if index in leaders and len(leaders) == 1:
        return "HIGHER-LIFT BUILD"
    if index in leaders:
        return "TIED FOR HIGHEST"
    cutoff = _tie_cutoff(digits)
    place = 1 + len({other.total for other in builds if other.total > builds[index].total + cutoff})
    return {2: "SECOND BUILD", 3: "THIRD BUILD"}.get(place, f"BUILD {place}")


def _category_label(element: Element) -> str:
    return _human_category(element.category)


def _choice_label(element: Element) -> str:
    return _readable_name(element.label or element.name)


def _varied_categories(builds: list[Build]) -> bool:
    names = {_category_label(element) for build in builds for element in build.elements}
    names.discard("")
    return len(names) > 1


def _grid_slots(count: int) -> tuple[float, list[float]]:
    """One card is centered and wide. Two share the row. Three use equal columns."""
    left, right = 0.56, 0.26
    width = SLIDE_W - left - right
    if count <= 1:
        card_w = min(7.15, width)
        return card_w, [left + (width - card_w) / 2]
    if count == 2:
        gap = 0.18
        card_w = (width - gap) / 2
        return card_w, [left, left + card_w + gap]
    gap = 0.12
    card_w = (width - gap * 2) / 3
    return card_w, [left + index * (card_w + gap) for index in range(count)]


def _buildup_cards(
    slide, study: StudyInfo, page: list[tuple[int, Build]], builds: list[Build], kind: str, digits: int,
) -> None:
    card_w, origins = _grid_slots(len(page))
    top, height = 1.22, 5.28
    shown = [build for _, build in page]
    show_category = _varied_categories(shown)
    for origin, (index, build) in zip(origins, page):
        _buildup_card(
            slide, study, shown, builds, build, index, kind, digits,
            origin, top, card_w, height, show_category,
        )


def _buildup_card(
    slide, study, shown, builds, build, index, kind, digits, x, y, w, h, show_category: bool,
) -> None:
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.04, line=_BUILD_LINE)
    pad = 0.10
    inner_x = x + pad
    inner_w = w - pad * 2
    header_h = 0.58 if w >= 4.4 else 0.52
    add_rect(slide, inner_x, y + 0.10, inner_w, header_h, _BUILD_HEADER, rounded=True, radius=0.12)
    role = _text(
        slide, _build_role(index, builds, digits),
        inner_x, y + 0.12, inner_w, 0.20,
        size=10 if w < 4.4 else 11, bold=True, color=_BUILD_ROLE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(role)
    name_size = 13 if w >= 4.4 else 12
    name = _keep_line(_combination_name(build, index), inner_w - 0.12, name_size, lines=1, bold=True)
    name_box = _text(
        slide, name,
        inner_x + 0.06, y + 0.32, inner_w - 0.12, header_h - 0.26,
        size=name_size, bold=True, color=NAVY,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(name_box)

    cursor = y + 0.10 + header_h + 0.06
    base = _base_for(study, kind)
    if base is not None:
        cursor = _buildup_base(slide, kind, digits, base, inner_x, cursor, inner_w)
        cursor += 0.04

    formula_h = 0.34
    panel_top = cursor
    panel_bottom = y + h - 0.08 - formula_h
    panel_h = max(1.2, panel_bottom - panel_top)
    _buildup_panel(
        slide, study, shown, build, digits, show_category,
        inner_x, panel_top, inner_w, panel_h, base is not None, w,
    )
    _buildup_formula(slide, build, builds, inner_x, panel_top + panel_h + 0.02, inner_w, formula_h - 0.02)


def _buildup_base(slide, kind: str, digits: int, base: float, x, y, w) -> float:
    row_h = 0.40
    if kind == "response":
        label, figure, note = "Base time", f"{base:.2f}", "starting point"
    elif kind == "bottom":
        label, figure, note = "Base score", f"{int(round(base))}%", "starting point"
    else:
        note = "starting point (overall)" if w >= 5.2 else "starting point"
        label, figure = "Base appeal", f"{int(round(base))}%"
    label_w = 1.20 if w >= 5.2 else 0.98
    figure_w = 0.72 if w >= 5.2 else 0.58
    _text(slide, label, x + 0.06, y, label_w, row_h, size=12, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    _text(
        slide, figure, x + 0.06 + label_w, y, figure_w, row_h,
        size=18 if w >= 4.6 else 15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _text(
        slide, note, x + 0.10 + label_w + figure_w, y, max(0.4, w - label_w - figure_w - 0.16), row_h,
        size=11, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
    )
    return y + row_h


def _buildup_panel(
    slide, study, builds, build, digits, show_category, x, y, w, h, has_base: bool, card_w: float,
) -> None:
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.03, line=_BUILD_LINE)
    band_h = 0.50 if card_w >= 4.6 else 0.44
    rows_top = y + 0.04
    rows_h = max(0.4, h - band_h - 0.06)
    elements = sorted(build.elements, key=lambda element: element.lift, reverse=True)
    count = max(len(elements), 1)
    row_h = rows_h / count
    visual_w = 0.0 if study.visual == "text" else (2.35 if card_w >= 5.6 else 1.28)
    visual_w = min(visual_w, w * 0.42)
    rows_w = w - visual_w - (0.08 if visual_w else 0.0)
    if elements:
        for index, element in enumerate(elements):
            _buildup_row(
                slide, element, digits, show_category,
                x + 0.08, rows_top + index * row_h, rows_w - 0.10, row_h,
                draw_line=index > 0,
            )
    else:
        _text(
            slide, "No elements in this combination",
            x + 0.10, rows_top, rows_w - 0.16, rows_h,
            size=12, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
        )
    if visual_w:
        _card_visual(
            slide, study, builds, build,
            x + rows_w, rows_top, visual_w - 0.06, rows_h,
            anchor="middle",
        )
    band_y = y + h - band_h
    add_rect(slide, x + 0.04, band_y, w - 0.08, band_h - 0.04, _BUILD_BAND, rounded=True, radius=0.08)
    total_label = "TOTAL LIFT OVER BASE" if has_base and card_w >= 5.2 else "TOTAL LIFT"
    _text(
        slide, total_label,
        x + 0.14, band_y, w * 0.52, band_h - 0.04,
        size=10 if card_w < 4.6 else 11, bold=True, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _fill, ink, _bar = _score_tone(build.total, digits)
    _text(
        slide, _lift(build.total, digits),
        x + w * 0.54, band_y, w * 0.42 - 0.10, band_h - 0.04,
        size=22 if card_w >= 5.2 else 16, bold=True, color=ink,
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
    )


def _buildup_row(slide, element: Element, digits: int, show_category: bool, x, y, w, h, draw_line: bool) -> None:
    if draw_line:
        add_rect(slide, x, y, w, 0.01, RGBColor(0xF0, 0xE6, 0xEA))
    # A three-up card is too narrow for category, score, and name on one line.
    if show_category and w < 3.15 and h >= 0.40:
        pill_w = 0.64 if digits < 3 else 0.78
        top_h = min(0.26, h * 0.46)
        cat_w = max(0.6, w - pill_w - 0.06)
        category = _keep_line(_category_label(element), cat_w, 10, lines=1)
        box = _text(slide, category, x, y, cat_w, top_h, size=10, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
        _tighten(box)
        _lift_pill(slide, element.lift, digits, x + cat_w, y + max(0, (top_h - 0.22) / 2), pill_w, min(0.22, top_h))
        label = _keep_line(_choice_label(element), w, 11, lines=1)
        box = _text(
            slide, label, x, y + top_h, w, h - top_h,
            size=11, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        )
        _tighten(box)
        return
    pill_w = 0.62 if digits < 3 else 0.78
    if w >= 3.3:
        pill_w = 0.78 if digits < 3 else 0.96
    cat_w = 0.0
    if show_category:
        cat_w = 0.78 if w < 2.9 else min(1.45 if w >= 3.3 else 1.05, w * 0.38)
    label_w = max(0.45, w - cat_w - pill_w - 0.08)
    size = 12 if h >= 0.36 else 10
    if show_category:
        category = _keep_line(_category_label(element), cat_w - 0.04, size - 1, lines=1)
        box = _text(
            slide, category, x, y, cat_w, h,
            size=size - 1, color=GRAY, anchor=MSO_ANCHOR.MIDDLE,
        )
        _tighten(box)
    pill_x = x + cat_w
    pill_h = min(0.28, max(0.18, h - 0.08))
    pill_y = y + (h - pill_h) / 2
    _lift_pill(slide, element.lift, digits, pill_x, pill_y, pill_w, pill_h)
    label = _keep_line(_choice_label(element), label_w - 0.06, size, lines=2 if h >= 0.46 else 1)
    box = _text(
        slide, label, pill_x + pill_w + 0.08, y, label_w, h,
        size=size, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(box)


def _lift_pill(slide, value: float, digits: int, x, y, w, h) -> None:
    cutoff = _tie_cutoff(digits)
    if value > cutoff:
        fill = _PILL_POS
    elif value < -cutoff:
        fill = _PILL_NEG
    else:
        fill = _PILL_ZERO
    add_rect(slide, x, y, w, h, fill, rounded=True, radius=0.35)
    box = _text(
        slide, _lift(value, digits), x, y, w, h,
        size=11 if h >= 0.24 else 9, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(box)


def _buildup_formula(slide, build: Build, builds: list[Build], x, y, w, h) -> None:
    elements = sorted(build.elements, key=lambda element: element.lift, reverse=True)
    labels = [_choice_label(element) for element in elements]
    if not labels:
        return
    other_names: set[str] = set()
    if len(builds) == 2:
        other = next(item for item in builds if item is not build)
        other_names = {element.name for element in other.elements}
    differs = [element.name not in other_names for element in elements] if other_names or len(builds) == 2 else [False] * len(labels)
    if len(builds) != 2:
        differs = [False] * len(labels)
    limit = max(18, int(w * 72 / (11 * 0.52)))
    shown: list[tuple[str, bool]] = []
    used = 0
    for label, differ in zip(labels, differs):
        piece = label
        extra = len(piece) + (3 if shown else 0)
        if used + extra > limit:
            room = limit - used - (3 if shown else 0)
            if room >= 4:
                shown.append((piece[: room - 1].rstrip() + "…", differ))
            elif shown:
                text, flag = shown[-1]
                shown[-1] = (text[: max(1, len(text) - 1)].rstrip() + "…", flag)
            break
        shown.append((piece, differ))
        used += extra
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for index, (label, differ) in enumerate(shown):
        if index:
            plus = paragraph.add_run()
            plus.text = " + "
            plus.font.size = Pt(11)
            plus.font.bold = False
            plus.font.color.rgb = GRAY
            plus.font.name = FONT
        run = paragraph.add_run()
        run.text = label
        run.font.size = Pt(11)
        run.font.bold = differ
        run.font.underline = differ
        run.font.color.rgb = NAVY
        run.font.name = FONT


def _keep_line(text: str, width_in: float, size: float, lines: int = 1, bold: bool = False) -> str:
    """Fit text on a few lines. A hyphen stays inside the word."""
    raw = re.sub(r"\s+", " ", str(text or "")).strip()
    if not raw or width_in <= 0.2 or lines < 1:
        return ""
    factor = 0.62 if bold else 0.54
    per = max(6, int(width_in * 72 / (size * factor)))
    words = raw.split(" ")
    rows: list[str] = []
    current = ""
    leftover = False
    for word in words:
        piece = word if len(word) <= per else word[: max(1, per - 1)].rstrip("-") + "…"
        trial = piece if not current else f"{current} {piece}"
        if len(trial) <= per:
            current = trial
            continue
        if len(rows) >= lines - 1:
            leftover = True
            break
        if current:
            rows.append(current)
        current = piece
    if current and len(rows) < lines:
        rows.append(current)
    if leftover and rows:
        rows[-1] = rows[-1][: max(1, per - 1)].rstrip() + "…"
    return "\n".join(rows[:lines])


def _buildup_key(slide, builds: list[Build], kind: str, digits: int, more: bool = False) -> None:
    x, y, w, h = 0.62, 6.62, SLIDE_W - 0.62 - 0.32, 0.46
    add_rect(slide, x, y, w, h, _KEY_FILL, rounded=True, radius=0.1)
    insight = _compare_insight(builds, kind, digits)
    tail = "More combinations are on the next slide." if more else "The next slides open each combination."
    body = f"{insight} {tail}"
    box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y), Inches(w - 0.32), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    lead = paragraph.add_run()
    lead.text = "Key read: "
    lead.font.size = Pt(13)
    lead.font.bold = True
    lead.font.color.rgb = NAVY
    lead.font.name = FONT
    rest = paragraph.add_run()
    rest.text = body
    rest.font.size = Pt(13)
    rest.font.bold = False
    rest.font.color.rgb = NAVY
    rest.font.name = FONT


def _stage_cards(slide, study, builds, kind, x, y, w, h) -> None:
    count = len(builds)
    gutter = 0.72 if count == 2 else 0.0
    gap = 0.28 if count > 2 else 0.0
    if count <= 1:
        card_w = min(6.4, w)
        slots = [x + (w - card_w) / 2]
    elif count == 2:
        card_w = (w - gutter) / 2
        slots = [x, x + card_w + gutter]
    else:
        card_w = (w - gap * (count - 1)) / count
        slots = [x + index * (card_w + gap) for index in range(count)]
    digits = _digits(kind)
    leaders = _leader_indexes(builds, digits)
    unique = len(leaders) == 1
    pad = 0.08
    image_h = _shared_image_height(study, builds, max(0.4, card_w - pad * 2), min(3.35, h * 0.68))
    metrics = _stage_metrics(y, h, image_h, card_w >= 4.6)
    for index, build in enumerate(builds):
        _stage_card(
            slide, study, builds, build, index, digits,
            slots[index], y, card_w, h,
            index in leaders, unique, metrics,
        )
    if count == 2:
        _gap_mark(
            slide, abs(builds[0].total - builds[1].total), digits,
            x + card_w, y, gutter, h, metrics["image_top"], metrics["image_h"],
        )

def _shared_image_height(study: StudyInfo, builds: list[Build], inner_w: float, max_h: float) -> float:
    """One picture height for the row, so the names line up."""
    max_h = max(0.7, max_h)
    heights: list[float] = []
    for build in builds:
        if study.visual == "text":
            heights.append(min(1.35, max_h))
        elif study.visual == "grid":
            heights.append(min(1.25, max_h))
        else:
            image = _render(study, build)
            aspect = file_aspect(image) if image else (study.aspect or 0.62)
            _width, height = contain_size(aspect, inner_w, max_h)
            heights.append(height)
    return max(heights) if heights else min(1.6, max_h)


def _stage_metrics(y: float, h: float, image_h: float, wide: bool) -> dict:
    """Positions for the picture, the name, and the score, centered as one block."""
    name_h = 0.34
    caption_h = 0.22
    number_h = 0.48 if wide else 0.40
    kicker_h = 0.18
    stack = image_h + 0.08 + name_h + 0.04 + kicker_h + number_h + caption_h
    origin = y + max(0.08, (h - stack) / 2)
    image_top = origin
    name_y = image_top + image_h + 0.08
    score_top = name_y + name_h + 0.02
    return {
        "image_top": image_top,
        "image_h": image_h,
        "name_y": name_y,
        "name_h": name_h,
        "score_top": score_top,
        "kicker_h": kicker_h,
        "number_h": number_h,
        "caption_h": caption_h,
    }


def _stage_card(
    slide, study, builds, build, index, digits, x, y, w, h, is_leader, unique, metrics,
) -> None:
    pad = 0.04
    inner_x = x + pad
    inner_w = max(0.4, w - pad * 2)
    _card_visual(
        slide, study, builds, build,
        inner_x, metrics["image_top"], inner_w, metrics["image_h"],
        anchor="bottom",
    )
    _text(
        slide, _clip_lines(_combination_name(build, index), inner_w, size=15 if w >= 4.6 else 13, lines=2),
        inner_x, metrics["name_y"], inner_w, metrics["name_h"],
        size=15 if w >= 4.6 else 13, bold=True, color=NAVY,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _text(
        slide, "Total lift",
        inner_x, metrics["score_top"], inner_w, metrics["kicker_h"],
        size=11, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tone_fill, ink, _bar = _score_tone(build.total, digits)
    _text(
        slide, _lift(build.total, digits),
        inner_x, metrics["score_top"] + metrics["kicker_h"], inner_w, metrics["number_h"],
        size=32 if w >= 4.6 else 26, bold=True, color=ink,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _text(
        slide, _score_caption(build, builds, digits, index, _leader_indexes(builds, digits)),
        inner_x, metrics["score_top"] + metrics["kicker_h"] + metrics["number_h"], inner_w, metrics["caption_h"],
        size=12, bold=is_leader, color=TEAL if is_leader and unique else GRAY,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

def _gap_mark(slide, gap: float, digits: int, x, y, w, h, image_top: float, image_h: float) -> None:
    """The difference, set between the two packs."""
    tied = gap <= _tie_cutoff(digits)
    figure = "TIE" if tied else _gap_figure(gap, digits)
    caption = "same total" if tied else "apart"
    block_top = image_top + max(0.0, image_h / 2 - 0.28)
    size = 15 if len(figure) > 4 else 20
    _text(
        slide, figure,
        x, block_top, w, 0.36,
        size=size, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _text(
        slide, caption,
        x, block_top + 0.32, w, 0.22,
        size=11, color=GRAY, align=PP_ALIGN.CENTER,
    )

def _card_visual(slide, study: StudyInfo, builds: list[Build], build: Build, x, y, w, h, anchor: str = "middle") -> None:
    if w <= 0.2 or h <= 0.2:
        return
    if study.visual == "text":
        body = _clip_lines(_variant_label(builds, build), w - 0.08, size=13, lines=6, readable=True)
        note = _text(
            slide, body, x, y, w, h,
            size=13, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        note.text_frame.word_wrap = True
        return
    if study.visual == "grid":
        _grid_thumb_row(slide, build, x, y, w, h)
        return
    image = _clear_pack_mat(_render(study, build))
    if image is None:
        labels = " · ".join(element.label for element in build.elements[:5]) or build.name
        note = _text(
            slide, _clip_lines(labels, w, size=12, lines=5, readable=True),
            x, y, w, h, size=12, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        note.text_frame.word_wrap = True
        return
    width, height = contain_size(file_aspect(image), w, h)
    left = x + (w - width) / 2
    if anchor == "bottom":
        top = y + max(0, h - height)
    elif anchor == "top":
        top = y
    else:
        top = y + (h - height) / 2
    _place_picture(slide, _clear_pack_mat(image), left, top, width, height)

def _closeness_ruler(slide, builds, digits, low, high, zoomed, x, y, w) -> None:
    if zoomed:
        caption = (
            f"The scale runs from {_axis_label(low, digits)} to {_axis_label(high, digits)}, "
            "not from zero, so a small gap stays visible."
        )
    else:
        caption = f"The scale runs from {_axis_label(low, digits)} to {_axis_label(high, digits)}."
    note = _text(slide, caption, x, y, w, 0.26, size=12, color=GRAY)
    note.text_frame.word_wrap = True
    track_y = y + 0.36
    track_h = 0.07
    add_rect(slide, x, track_y, w, track_h, TRACK, rounded=True, radius=0.5)
    span = high - low or 1.0
    leaders = _leader_indexes(builds, digits)
    dot = 0.18
    centers = []
    for build in builds:
        fraction = min(1.0, max(0.0, (build.total - low) / span))
        centers.append(x + fraction * w)
    label_w = 2.35
    labels = _stagger_labels(centers, x, w, label_w)
    for index, build in enumerate(builds):
        center = centers[index]
        leader = index in leaders
        _oval(slide, center - dot / 2, track_y + (track_h - dot) / 2, dot, dot, TEAL if leader else NAVY)
        left, above = labels[index]
        text_y = track_y - 0.26 if above else track_y + 0.12
        label = f"{_combination_name(build, index)}  {_lift(build.total, digits)}"
        _text(
            slide, _clip_lines(label, label_w, size=11, lines=1),
            left, text_y, label_w, 0.22,
            size=11, bold=leader, color=NAVY if leader else GRAY, align=PP_ALIGN.CENTER,
        )
    if low < 0 < high:
        zero_x = x + (0 - low) / span * w
        _hairline(slide, zero_x, track_y - 0.05, zero_x, track_y + track_h + 0.05, color=GRAY, width=1.0)

def _stagger_labels(centers: list[float], x: float, w: float, label_w: float) -> list[tuple[float, bool]]:
    """(left, above) for each dot. Above is only used when two names would overlap."""
    order = sorted(range(len(centers)), key=lambda index: centers[index])
    placed: list[tuple[float, bool] | None] = [None] * len(centers)
    for index in order:
        left = min(max(x, centers[index] - label_w / 2), x + w - label_w)
        above = False
        for other in order:
            previous = placed[other]
            if previous is None:
                continue
            other_left, other_above = previous
            overlaps = left < other_left + label_w + 0.1 and other_left < left + label_w + 0.1
            if overlaps and above == other_above:
                above = not other_above
        placed[index] = (left, above)
    return [item if item is not None else (x, False) for item in placed]

def _ranked_rows(slide, study, builds, entries, kind, x, y, w, h, start_rank: int) -> None:
    """A ranked list. One column up to six, two columns after that. No boxes."""
    count = max(len(entries), 1)
    columns = 2 if count > 6 else 1
    col_gap = 0.48
    col_w = w if columns == 1 else (w - col_gap) / 2
    per_col = (count + columns - 1) // columns
    row_h = h / per_col
    digits = _digits(kind)
    leaders = _leader_indexes(builds, digits)
    unique = len(leaders) == 1
    for offset, (original_index, build) in enumerate(entries):
        column = offset // per_col
        row = offset % per_col
        _ranked_row(
            slide, study, builds, build, original_index, start_rank + offset, digits,
            x + column * (col_w + col_gap), y + row * row_h, col_w, row_h,
            original_index in leaders, unique,
        )


def _ranked_row(
    slide, study, builds, build, index, rank, digits, x, y, w, h, is_leader, unique,
) -> None:
    rank_w = 0.38
    score_w = 1.55 if digits >= 3 else 1.28
    thumb = 0.0
    if study.visual != "text" and h >= 0.70:
        thumb = min(0.92, h - 0.10)
    _text(
        slide, str(rank),
        x, y, rank_w, h,
        size=14, bold=is_leader, color=TEAL if is_leader else GRAY, anchor=MSO_ANCHOR.MIDDLE,
    )
    text_x = x + rank_w
    if thumb:
        _card_visual(slide, study, builds, build, text_x, y + (h - thumb) / 2, thumb, thumb)
        text_x += thumb + 0.12
    text_w = max(1.1, w - (text_x - x) - score_w)
    name_size = 14 if h >= 0.78 else 12
    name_h = min(0.32, h * 0.42)
    caption_h = 0.20
    block = name_h + caption_h
    name_y = y + max(0.0, (h - block) / 2)
    _text(
        slide, _clip_lines(_combination_name(build, index), text_w, size=name_size, lines=1),
        text_x, name_y, text_w, name_h,
        size=name_size, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _text(
        slide, _score_caption(build, builds, digits, index, _leader_indexes(builds, digits)),
        text_x, name_y + name_h, text_w, caption_h,
        size=11, bold=is_leader and unique, color=TEAL if is_leader and unique else GRAY,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _tone_fill, ink, _bar = _score_tone(build.total, digits)
    _text(
        slide, _lift(build.total, digits),
        x + w - score_w, y, score_w, h,
        size=20 if h >= 0.78 else 16, bold=True, color=ink,
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
    )

def _nice_ceiling(value: float) -> float:
    if value <= 0:
        return 1
    if value <= 2:
        step = 0.2
    elif value <= 8:
        step = 1
    elif value <= 40:
        step = 4
    else:
        step = 10
    count = int(value / step)
    if abs(value - count * step) > 1e-9:
        count += 1
    return round(step * count, 3)

# Colors sampled from the Rexona combination detail slide.
_INK = RGBColor(0x1A, 0x2A, 0x40)
_TITLE = RGBColor(0x0B, 0x1F, 0x4B)
_EYEBROW = RGBColor(0x1C, 0x25, 0x35)
_SEGMENT = RGBColor(0x02, 0x80, 0x90)
_MUTED = RGBColor(0x8A, 0x9B, 0xB0)
_BAR = RGBColor(0xF4, 0xE0, 0xE7)
_RULE = RGBColor(0xD0, 0xD8, 0xE4)
_PILL_UP = RGBColor(0x1A, 0x7A, 0x4A)
_PILL_DOWN = RGBColor(0xCC, 0x33, 0x00)
_RAIL_BLUSH = RGBColor(0xF5, 0xE7, 0xEE)
# One meaning for color: a plus is green, a minus is red.
_STRONG_FILL = RGBColor(0xE3, 0xF6, 0xEA)
_STRONG_INK = RGBColor(0x0E, 0x6B, 0x3C)
_STRONG_BAR = RGBColor(0x1F, 0x8A, 0x4C)
_WEAK_FILL = RGBColor(0xFD, 0xE8, 0xE4)
_WEAK_INK = RGBColor(0xB4, 0x23, 0x18)
_WEAK_BAR = RGBColor(0xD6, 0x45, 0x3D)
_MID_FILL = RGBColor(0xEE, 0xF1, 0xF4)
_MID_INK = RGBColor(0x3D, 0x4A, 0x5C)
_MID_BAR = RGBColor(0xC5, 0xCE, 0xD6)

def _score_tone(value: float, digits: int = 1) -> tuple[RGBColor, RGBColor, RGBColor]:
    """Signed number color. A plus is green, a minus is red, and a zero stays neutral."""
    cutoff = 0.0005 if digits >= 3 else 0.05
    if value < -cutoff:
        return _WEAK_FILL, _WEAK_INK, _WEAK_BAR
    if value > cutoff:
        return _STRONG_FILL, _STRONG_INK, _STRONG_BAR
    return _MID_FILL, _MID_INK, _MID_BAR

_CANVAS = RGBColor(0xF5, 0xF7, 0xFB)
_SAMPLE_FILL = RGBColor(0xE3, 0xF1, 0xFB)
_SAMPLE_INK = RGBColor(0x1E, 0x4E, 0x89)
_HEADER_TONES = (
    (RGBColor(0xF8, 0xE7, 0xF0), RGBColor(0x9D, 0x3B, 0x6A)),
    (RGBColor(0xE5, 0xF6, 0xEC), RGBColor(0x1C, 0x7A, 0x45)),
    (RGBColor(0xFB, 0xF3, 0xD8), RGBColor(0x8A, 0x5A, 0x10)),
    (RGBColor(0xEE, 0xE7, 0xF8), RGBColor(0x5C, 0x3F, 0x8C)),
)
_HEAT_ZERO = (RGBColor(0xF4, 0xF6, 0xF8), RGBColor(0x8A, 0x97, 0xA6))


def _column_family(segment: ComparisonSegment) -> str:
    """A role for the header tint. The words come from the saved question, not a fixed list of columns."""
    question = (segment.question or "").lower()
    label = (segment.label or "").lower()
    if "gender" in question or "gender" in (segment.key or "").lower() or label in {"female", "male", "woman", "man", "women", "men"}:
        return "gender"
    if "age" in question or re.search(r"\d{1,2}\s*[-–]\s*\d{1,2}", label):
        return "age"
    return "other"


def _header_tone(segment: ComparisonSegment, index: int) -> tuple[RGBColor, RGBColor]:
    family = _column_family(segment)
    if family == "gender":
        return _HEADER_TONES[0]
    if family == "age":
        return _HEADER_TONES[1]
    return _HEADER_TONES[2 + (index % 2)]


def _column_icon(segment: ComparisonSegment) -> str:
    """Preset mark for age, gender, male, female, and country. Other questions share one icon."""
    question = f"{segment.question} {segment.key}".lower()
    answer = _segment_answer(segment).lower()
    if answer in {"female", "woman", "women"}:
        return "female"
    if answer in {"male", "man", "men"}:
        return "male"
    if "gender" in question:
        return "gender"
    if "age" in question or re.search(r"\d{1,2}\s*[-–]\s*\d{1,2}", answer):
        return "age"
    if any(word in question for word in ("country", "countries", "nation")) or "live in" in question:
        return "country"
    return "generic"


def _preset_icon(kind: str, rgb: tuple[int, int, int]) -> Path:
    """A small transparent icon. Sigma is the overall mark; the rest sit in the column headers."""
    folder = Path("/tmp/ms-report-icons")
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / f"{kind}-{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}.png"
    if dest.exists():
        return dest
    size = 128
    image = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    ink = (*rgb, 255)
    stroke = 10
    if kind == "sigma":
        draw.ellipse((6, 6, 122, 122), fill=(0xE3, 0xF1, 0xFB, 255))
        draw.line((86, 36, 40, 36), fill=ink, width=stroke)
        draw.line((40, 36, 78, 64), fill=ink, width=stroke)
        draw.line((78, 64, 40, 92), fill=ink, width=stroke)
        draw.line((40, 92, 86, 92), fill=ink, width=stroke)
    elif kind == "female":
        draw.ellipse((36, 14, 92, 70), outline=ink, width=stroke)
        draw.line((64, 70, 64, 112), fill=ink, width=stroke)
        draw.line((44, 94, 84, 94), fill=ink, width=stroke)
    elif kind == "male":
        draw.ellipse((22, 40, 80, 98), outline=ink, width=stroke)
        draw.line((74, 46, 110, 12), fill=ink, width=stroke)
        draw.line((86, 12, 112, 12), fill=ink, width=stroke)
        draw.line((112, 12, 112, 38), fill=ink, width=stroke)
    elif kind == "gender":
        draw.ellipse((10, 28, 52, 70), outline=ink, width=8)
        draw.line((31, 70, 31, 108), fill=ink, width=8)
        draw.line((16, 92, 46, 92), fill=ink, width=8)
        draw.ellipse((58, 36, 100, 78), outline=ink, width=8)
        draw.line((94, 42, 116, 18), fill=ink, width=8)
        draw.line((100, 16, 118, 16), fill=ink, width=8)
        draw.line((116, 16, 116, 34), fill=ink, width=8)
    elif kind == "age":
        draw.ellipse((44, 10, 84, 50), outline=ink, width=stroke)
        draw.arc((18, 54, 110, 132), 200, 340, fill=ink, width=stroke)
    elif kind == "country":
        draw.ellipse((14, 14, 114, 114), outline=ink, width=stroke)
        draw.ellipse((40, 14, 88, 114), outline=ink, width=8)
        draw.line((18, 64, 110, 64), fill=ink, width=8)
    elif kind == "sample":
        draw.ellipse((18, 18, 52, 52), outline=ink, width=8)
        draw.arc((4, 56, 66, 110), 200, 340, fill=ink, width=8)
        draw.ellipse((70, 14, 110, 54), outline=ink, width=8)
        draw.arc((58, 56, 124, 114), 200, 340, fill=ink, width=8)
    else:
        draw.rounded_rectangle((22, 18, 106, 110), radius=14, outline=ink, width=stroke)
        draw.line((40, 46, 88, 46), fill=ink, width=8)
        draw.line((40, 66, 88, 66), fill=ink, width=8)
        draw.line((40, 86, 70, 86), fill=ink, width=8)
    image.save(dest)
    return dest


def _heading_parts(segment: ComparisonSegment) -> tuple[str, str]:
    question = _condense_question(segment.question).rstrip("?").strip() if segment.question else ""
    answer = _segment_answer(segment)
    if not question:
        return answer or segment.label, ""
    words = question.split()
    if len(words) == 1 and words[0].lower() == "gender":
        return answer, ""
    if len(words) <= 2 and len(question) <= 18:
        return f"{question} {answer}".strip(), ""
    title = question + ("?" if (segment.question or "").strip().endswith("?") else "")
    if answer and answer.lower() not in title.lower():
        return title, answer
    return title, ""


def _ramp(magnitude: float, stops: tuple) -> tuple[tuple[int, int, int], tuple[int, int, int]]:
    if magnitude >= stops[-1][0]:
        return stops[-1][1], stops[-1][2]
    for index in range(1, len(stops)):
        low, high = stops[index - 1], stops[index]
        if magnitude <= high[0]:
            span = high[0] - low[0] or 1.0
            amount = (magnitude - low[0]) / span
            fill = tuple(int(low[1][channel] + (high[1][channel] - low[1][channel]) * amount) for channel in range(3))
            return fill, high[2] if amount >= 0.55 else low[2]
    return stops[0][1], stops[0][2]


# Nearby scores stay apart. A small lift stays pale, the teens turn bright, and the totals keep going darker.
_POS_STOPS = (
    (0.0, (0xF3, 0xFB, 0xF6), (0x1D, 0x6B, 0x3C)),
    (3.0, (0xC6, 0xED, 0xD4), (0x14, 0x5C, 0x32)),
    (7.0, (0x7D, 0xD4, 0x9A), (0x0E, 0x4E, 0x2A)),
    (12.0, (0x3E, 0xC0, 0x72), (0xFF, 0xFF, 0xFF)),
    (18.0, (0x1C, 0x9A, 0x4E), (0xFF, 0xFF, 0xFF)),
    (26.0, (0x12, 0x7A, 0x3C), (0xFF, 0xFF, 0xFF)),
    (40.0, (0x0E, 0x64, 0x32), (0xFF, 0xFF, 0xFF)),
    (70.0, (0x0C, 0x50, 0x28), (0xFF, 0xFF, 0xFF)),
    (100.0, (0x08, 0x3C, 0x1E), (0xFF, 0xFF, 0xFF)),
    (130.0, (0x05, 0x2C, 0x16), (0xFF, 0xFF, 0xFF)),
)
_ROW_BAND = RGBColor(0xF4, 0xF7, 0xFB)
_NEG_STOPS = (
    (0.0, (0xF8, 0xE4, 0xE0), (0x9C, 0x2A, 0x22)),
    (6.0, (0xF0, 0xB4, 0xAA), (0x9C, 0x2A, 0x22)),
    (14.0, (0xE0, 0x6A, 0x62), (0xFF, 0xFF, 0xFF)),
    (28.0, (0xD0, 0x3C, 0x34), (0xFF, 0xFF, 0xFF)),
    (50.0, (0xA8, 0x1E, 0x18), (0xFF, 0xFF, 0xFF)),
)


def _heat(value: float, digits: int) -> tuple[RGBColor, RGBColor]:
    """Color deepens with the score, so +17 reads lighter than +80."""
    cutoff = 0.0005 if digits >= 3 else 0.05
    if abs(value) <= cutoff:
        return _HEAT_ZERO
    fill, ink = _ramp(abs(value), _NEG_STOPS if value < 0 else _POS_STOPS)
    return RGBColor(*fill), RGBColor(*ink)


def _detail_takeaways(build: Build, digits: int) -> list[tuple[float, str, str]]:
    """Two or three findings from this combination: score, name, and a short line."""
    elements = list(build.elements)
    if not elements:
        return [(0.0, "No elements", "Nothing is saved in this combination.")]
    ranked = sorted(elements, key=lambda element: element.lift, reverse=True)
    best = ranked[0]
    points = [(best.lift, _readable_name(best.label), f"Leads the full sample at {_lift(best.lift, digits)}")]
    cutoff = 0.0005 if digits >= 3 else 0.05
    if len(ranked) > 1 and ranked[1].lift > cutoff:
        nxt = ranked[1]
        points.append((nxt.lift, _readable_name(nxt.label), f"Next in the full sample, at {_lift(nxt.lift, digits)}"))
    weakest = None
    for element in elements:
        for segment in build.comparison_segments:
            value = float(element.segment_lifts.get(segment.key, 0.0))
            if weakest is None or value < weakest[0]:
                weakest = (value, element, segment)
    if weakest is not None and weakest[0] < -cutoff:
        value, element, segment = weakest
        who = _segment_answer(segment) or segment.label
        points.append((value, _readable_name(element.label), f"{_lift(value, digits)} among {who}"))
    elif len(points) < 3 and ranked[-1].lift < ranked[0].lift:
        last = ranked[-1]
        points.append((last.lift, _readable_name(last.label), f"Smallest full-sample lift, at {_lift(last.lift, digits)}"))
    return points[:3]


def _element_noun(visual: str) -> str:
    if visual == "text":
        return "Statements"
    if visual == "grid":
        return "Images"
    return "Pack elements"


def _pack(presentation, study: StudyInfo, group: Group, build: Build, headline: str, kind: str):
    """One combination: a colored score table, the pack, and takeaways from these scores."""
    slide = _blank(presentation)
    set_background(slide, _CANVAS)

    digits = _digits(kind)
    sample = study.respondents if study.respondents > 0 else 0
    title = _text(slide, _prefix(kind, study.visual).title(), 0.46, 0.22, 6.3, 0.46, size=28, bold=True, color=NAVY)
    title.text_frame.word_wrap = False
    columns = len(build.comparison_segments) + (1 if build.comparison_segments else 0)
    columns = max(columns, 1)
    board_w, side_x, side_w = _detail_frame(columns)
    _text(slide, headline, 0.48, 0.86, board_w - 0.2, 0.24, size=12, bold=True, color=_MUTED)
    pill = "Pack view"
    pill_w = min(2.3, max(1.45, 0.36 + len(pill) * 0.075))
    pill_x = 0.46 + board_w - pill_w
    add_rect(slide, pill_x, 0.42, pill_w, 0.34, WHITE, rounded=True, radius=0.5, line=RGBColor(0xD5, 0xDE, 0xEA))
    _text(slide, pill, pill_x, 0.42, pill_w, 0.34, size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    _detail_base_bar(slide, study, kind, digits, sample, 0.46, 1.20, board_w)
    _detail_board(slide, study, build, digits, sample, 0.46, 1.70, board_w, 5.05)
    _detail_side(slide, study, build, digits, side_x, 0.22, side_w, 6.55)

    places = "three decimals" if digits == 3 else "one decimal"
    footer = _text(slide, _detail_footer(build), 0.48, 6.88, board_w, 0.42, size=11, color=_MUTED)
    footer.text_frame.word_wrap = True
    _text(slide, f"Values rounded to {places}.", side_x, 7.12, side_w, 0.22, size=10, color=_MUTED, align=PP_ALIGN.RIGHT)
    return slide


def _detail_base_bar(slide, study: StudyInfo, kind: str, digits: int, sample: int, x, y, w) -> None:
    add_rect(slide, x, y, w, 0.40, RGBColor(0xEE, 0xE8, 0xF8), rounded=True, radius=0.4)
    _text(slide, "Base size", x + 0.18, y, 1.15, 0.40, size=13, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    if sample:
        _text(slide, f"n = {sample}", x + 1.36, y, 1.25, 0.40, size=16, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        note = "respondents in the full sample"
    else:
        note = "full-sample size was not saved"
    base_value = _base_for(study, kind)
    if base_value is not None:
        note = f"{note}  ·  base appeal {base_value:.{digits}f}"
    _text(slide, note, x + 2.70, y, w - 2.88, 0.40, size=12, color=_MUTED, anchor=MSO_ANCHOR.MIDDLE)


def _detail_frame(columns: int) -> tuple[float, float, float]:
    """Table width, then the pack's left edge and width. The pack stays on the slide."""
    left, gap, right = 0.46, 0.14, 0.32
    pack_w = 3.70 if columns <= 4 else 3.25 if columns <= 5 else 2.90
    board_w = SLIDE_W - left - gap - pack_w - right
    return board_w, left + board_w + gap, pack_w


def _measure_header(title: str, subtitle: str, cell_w: float) -> tuple[str, str, float, float, float, float, float]:
    """Question, answer, sizes, and the height that keeps both inside the header."""
    narrow = cell_w < 1.20
    title_size = 8 if narrow else 10
    sub_size = 8 if narrow else 9
    inner = max(0.36, cell_w - 0.10)
    title_text = _fit_lines(title, inner, title_size, 3 if narrow else 2, bold=True)
    sub_text = _fit_lines(subtitle, inner, sub_size, 2) if subtitle else ""
    title_n = title_text.count("\n") + 1 if title_text else 1
    sub_n = sub_text.count("\n") + 1 if sub_text else 0
    title_h = title_n * (title_size / 72 * 1.28) + 0.02
    sub_h = sub_n * (sub_size / 72 * 1.28) + (0.03 if sub_n else 0.0)
    block_h = 0.32 + title_h + sub_h + 0.18
    return title_text, sub_text, title_size, sub_size, title_h, sub_h, block_h


def _detail_board(slide, study: StudyInfo, build: Build, digits: int, sample: int, x, y, w, h) -> None:
    segments = list(build.comparison_segments)
    columns = len(segments) + (1 if segments else 0)
    columns = max(columns, 1)
    name_floor = 2.40 if columns <= 4 else 2.10 if columns <= 5 else 1.90
    name_w = min(3.05, max(name_floor, w - columns * 1.08 - 0.12))
    score_w = (w - name_w - 0.12) / columns
    add_rect(slide, x, y, w, h, WHITE, rounded=True, radius=0.06, line=RGBColor(0xE1, 0xE7, 0xF0))
    cell_w = score_w - 0.08
    sample_header = _measure_header("Full sample", "", cell_w)
    measured = [_measure_header(*_heading_parts(segment), cell_w) for segment in segments]
    overall_h = 0.46
    bottom_gap = 0.08
    elements = sorted(build.elements, key=lambda element: element.lift, reverse=True)
    count = max(len(elements), 1)
    needed = max((item[6] for item in (sample_header, *measured)), default=0.58)
    header_h = min(needed, h - 0.08 - overall_h - bottom_gap - count * 0.36)
    header_h = max(0.50, header_h)
    row_h = (h - 0.08 - header_h - overall_h - bottom_gap) / count
    head_y = y + 0.04
    _detail_header(
        slide, study, segments, sample_header, measured, sample,
        x + 0.06, head_y, name_w, score_w, header_h - 0.02,
    )
    body = head_y + header_h
    for index, element in enumerate(elements):
        _detail_row(
            slide, element, segments, digits,
            x + 0.06, body + index * row_h, name_w, score_w, row_h, index,
        )
    _detail_overall(
        slide, build, segments, digits, sample,
        x + 0.06, y + h - overall_h - bottom_gap, name_w, score_w, overall_h,
    )


def _detail_header(slide, study, segments, sample_header, measured, sample, x, y, name_w, score_w, h) -> None:
    _text(
        slide, _element_noun(study.visual),
        x + 0.08, y, name_w - 0.16, h,
        size=13 if score_w < 1.15 else 14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _header_score(
        slide, x + name_w, y, score_w - 0.08, h,
        sample_header, sample, _SAMPLE_FILL, _SAMPLE_INK, "sample",
    )
    for index, segment in enumerate(segments):
        fill, ink = _header_tone(segment, index)
        _header_score(
            slide, x + name_w + (index + 1) * score_w, y, score_w - 0.08, h,
            measured[index], segment.base_size, fill, ink, _column_icon(segment),
        )


def _header_score(slide, x, y, w, h, fitted, count, fill, ink, icon: str) -> None:
    title, subtitle, title_size, sub_size, _title_h, _sub_h, _block = fitted
    add_rect(slide, x, y, w, h, fill)
    count_h = 0.16 if count else 0.0
    icon_box = min(0.24, w * 0.32, max(0.14, h - count_h - 0.30))
    _place_picture(
        slide, _preset_icon(icon, (int(ink[0]), int(ink[1]), int(ink[2]))),
        x + (w - icon_box) / 2, y + 0.03, icon_box, icon_box,
    )
    text_top = y + 0.05 + icon_box
    room = max(0.14, h - (text_top - y) - count_h - 0.04)
    inner = max(0.3, w - 0.10)
    title_lines = 2 if room > 0.46 and "\n" in title else 1
    title_draw = _fit_lines(title.replace("\n", " "), inner, title_size, title_lines, bold=True)
    sub_draw = ""
    if subtitle and room > 0.30:
        sub_draw = _fit_lines(subtitle.replace("\n", " "), inner, sub_size, 1)
    if sub_draw:
        title_h = min(room * 0.55, 0.16 * title_lines + 0.02)
        sub_h = max(0.14, room - title_h)
    else:
        title_h = room
        sub_h = 0.0
    label = _text(
        slide, title_draw,
        x + 0.04, text_top, w - 0.08, title_h,
        size=title_size, bold=True, color=ink, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE if not sub_draw else MSO_ANCHOR.BOTTOM,
    )
    _tighten(label)
    if sub_draw:
        sub = _text(
            slide, sub_draw,
            x + 0.04, text_top + title_h, w - 0.08, sub_h,
            size=sub_size, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
        )
        _tighten(sub)
    if count:
        count_label = _text(
            slide, f"(n={count})",
            x + 0.04, y + h - count_h - 0.02, w - 0.08, count_h,
            size=8 if w < 1.05 else 9, color=_MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        _tighten(count_label)


def _detail_row(slide, element: Element, segments, digits, x, y, name_w, score_w, row_h, index: int) -> None:
    if index % 2 == 1:
        add_rect(slide, x, y, name_w - 0.02, row_h, _ROW_BAND)
    thumb = 0.0
    if element.image_path and Path(element.image_path).exists() and row_h > 0.24:
        # Stay inside the row, and leave the rest of the column for the name.
        thumb = min(row_h - 0.06, 0.40, max(0.18, name_w * 0.22))
        thumb = min(thumb, max(0.18, name_w - 1.15))
        _place_picture(
            slide, _element_thumb(Path(element.image_path)),
            x + 0.04, y + (row_h - thumb) / 2, thumb, thumb,
        )
    text_x = x + (0.04 + thumb + 0.08 if thumb else 0.06)
    text_w = max(0.7, name_w - (text_x - x) - 0.04)
    lines = 1 if row_h < 0.46 else 2
    shown = _clip_lines(element.label, text_w, size=12, lines=lines, readable=True)
    name = _text(
        slide, shown, text_x, y, text_w, row_h,
        size=12, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(name)
    score_size = 12 if score_w < 1.15 else 13
    pad_x, pad_y = 0.015, 0.012
    _heat_cell(
        slide, element.lift, digits,
        x + name_w + pad_x, y + pad_y, score_w - pad_x * 2, row_h - pad_y * 2, size=score_size,
    )
    for segment_index, segment in enumerate(segments):
        value = float(element.segment_lifts.get(segment.key, 0.0))
        _heat_cell(
            slide, value, digits,
            x + name_w + (segment_index + 1) * score_w + pad_x, y + pad_y,
            score_w - pad_x * 2, row_h - pad_y * 2, size=score_size,
        )


def _detail_overall(slide, build: Build, segments, digits, sample, x, y, name_w, score_w, row_h) -> None:
    add_rect(slide, x, y, name_w - 0.02, row_h, _SAMPLE_FILL)
    mark = min(0.22, row_h - 0.10)
    _place_picture(
        slide, _preset_icon("sigma", (0x1E, 0x4E, 0x89)),
        x + 0.08, y + (row_h - mark) / 2, mark, mark,
    )
    label = "Overall" if not sample else f"Overall  (n={sample})"
    _text(
        slide, label, x + 0.12 + mark, y, name_w - mark - 0.22, row_h,
        size=12 if name_w < 2.3 else 13, bold=True, color=_SAMPLE_INK, anchor=MSO_ANCHOR.MIDDLE,
    )
    score_size = 12 if score_w < 1.15 else 14
    pad_x, pad_y = 0.015, 0.012
    _heat_cell(
        slide, build.total, digits,
        x + name_w + pad_x, y + pad_y, score_w - pad_x * 2, row_h - pad_y * 2, size=score_size,
    )
    for index, segment in enumerate(segments):
        _heat_cell(
            slide, _segment_total(build, segment.key), digits,
            x + name_w + (index + 1) * score_w + pad_x, y + pad_y,
            score_w - pad_x * 2, row_h - pad_y * 2, size=score_size,
        )


def _heat_cell(slide, value: float, digits: int, x, y, w, h, size: float = 12) -> None:
    fill, ink = _heat(value, digits)
    add_rect(slide, x, y, w, h, fill)
    label = _text(
        slide, _lift(value, digits),
        x, y, w, h,
        size=size, bold=True, color=ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    _tighten(label)


def _detail_side(slide, study: StudyInfo, build: Build, digits: int, x, y, w, h) -> None:
    points = _detail_takeaways(build, digits)
    row_h = 0.62
    card_h = min(2.35, 0.42 + row_h * len(points))
    card_y = y + h - card_h
    image_h = max(2.4, card_y - y - 0.16)
    if study.visual == "grid":
        _grid_thumb_row(slide, build, x, y, w, image_h)
    elif study.visual != "text":
        image = _render(study, build)
        _place_picture(slide, _clear_pack_mat(image), x, y, w, image_h)
    else:
        card_y = y
        card_h = min(h, 0.42 + row_h * len(points))
    _text(slide, "KEY POINTS", x, card_y, w, 0.22, size=11, bold=True, color=TEAL)
    add_rect(slide, x, card_y + 0.26, 0.42, 0.035, TEAL)
    cursor = card_y + 0.40
    for value, name, detail in points:
        fill, _ink = _heat(value, digits)
        add_rect(slide, x, cursor + 0.06, 0.07, row_h - 0.18, fill)
        title = _text(
            slide, _clip_lines(name, w - 0.28, size=13, lines=1, readable=True),
            x + 0.18, cursor, w - 0.18, 0.28,
            size=13, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        )
        title.text_frame.word_wrap = False
        note = _text(slide, detail, x + 0.18, cursor + 0.28, w - 0.18, 0.26, size=11, color=_MUTED)
        note.text_frame.word_wrap = True
        cursor += row_h

def _rank_headline(kind: str, index: int, visual: str) -> str:
    if visual == "layer":
        ranks = _KIND[kind]["ranks"]
        if index < len(ranks):
            return ranks[index]
        return f"PACK {index + 1}"
    noun = "STATEMENT" if visual == "text" else "SET"
    if kind == "bottom":
        ranks = (
            f"{noun} 1 — LOWEST LIFT",
            f"{noun} 2",
            f"{noun} 3",
        )
    elif kind == "response":
        ranks = (f"{noun} 1", f"{noun} 2", f"{noun} 3")
    else:
        ranks = (
            f"{noun} 1 — HIGHEST LIFT",
            f"{noun} 2 — SECOND HIGHEST LIFT",
            f"{noun} 3 — THIRD HIGHEST LIFT",
        )
    if index < len(ranks):
        return ranks[index]
    return f"{noun} {index + 1}"

def _prefix(kind: str, visual: str) -> str:
    if visual == "text":
        return {"top": "STATEMENTS", "bottom": "STATEMENTS", "response": "STATEMENTS"}[kind]
    if visual == "grid":
        return {"top": "IMAGE SETS", "bottom": "IMAGE SETS", "response": "IMAGE SETS"}[kind]
    return _KIND[kind]["prefix"]

def _grid_thumb_row(slide, build: Build, x, y, w, h):
    paths = [element.image_path for element in build.elements if element.image_path and element.image_path.exists()]
    if not paths:
        names = _clip_lines(" · ".join(element.label for element in build.elements[:4]), w - 0.08, size=10, lines=3, readable=True)
        note = _text(slide, names, x, y, w, h, size=10, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        note.text_frame.word_wrap = True
        return
    gap = 0.06
    cell_w = (w - gap * (len(paths) - 1)) / len(paths)
    for index, path in enumerate(paths):
        _place_fit(slide, path, x + index * (cell_w + gap), y, cell_w, h)

def _place_grid_images(slide, build: Build):
    """Each selected image in its own cell. Nothing is stacked into a pack."""
    paths = [element.image_path for element in build.elements if element.image_path and element.image_path.exists()]
    if not paths:
        return
    x, y, w, h = 9.15, 1.85, 3.85, 5.05
    cols = 2 if len(paths) > 3 else 1
    rows = (len(paths) + cols - 1) // cols
    gap = 0.08
    cell_w = (w - gap * (cols - 1)) / cols
    cell_h = (h - gap * (rows - 1)) / rows
    for index, path in enumerate(paths):
        col = index % cols
        row = index // cols
        _place_fit(slide, path, x + col * (cell_w + gap), y + row * (cell_h + gap), cell_w, cell_h)

def _place_statements(slide, build: Build, digits: int, sample: int = 0):
    """Full statements across the slide. There is no picture column."""
    elements = build.elements
    if not elements:
        _total_row(slide, build.total, 1.85, 0.83, 8.2, 9.2, digits, sample=sample)
        return
    top = 1.82
    bottom = 6.72
    gap = 0.08
    cat_w = 1.9
    text_x = 0.83 + cat_w + 0.1
    pill_x = 11.85
    text_w = pill_x - text_x - 0.12
    size = 13
    heights = [_statement_height(element.label, text_w, size) for element in elements]
    block = sum(heights) + gap * (len(elements) - 1) + 0.50
    if block > bottom - top:
        size = 11
        gap = 0.05
        heights = [_statement_height(element.label, text_w, size) for element in elements]
    room = bottom - top - 0.50
    body = sum(heights) + gap * (len(elements) - 1)
    if body > room and body > 0:
        scale = room / body
        heights = [max(0.34, height * scale) for height in heights]
    cursor = top
    for element, row_h in zip(elements, heights):
        _text(slide, element.category, 0.83, cursor, cat_w, row_h, size=11, bold=True, color=_SEGMENT, anchor=MSO_ANCHOR.MIDDLE)
        statement = _text(slide, element.label, text_x, cursor, text_w, row_h, size=size, color=_INK, anchor=MSO_ANCHOR.MIDDLE)
        statement.text_frame.word_wrap = True
        _score_pill(slide, element.lift, pill_x, cursor, row_h, size=12, digits=digits)
        cursor += row_h + gap
    if cursor > bottom:
        cursor = bottom
    _total_row(slide, build.total, cursor, 0.83, 7.4, pill_x, digits, sample=sample)

def _statement_height(text: str, width_in: float, size: float) -> float:
    chars = max(8, int(width_in * 11 * (11 / max(size, 1))))
    words = str(text).split() or [""]
    lines = 1
    current = 0
    for word in words:
        extra = len(word) if current == 0 else current + 1 + len(word)
        if extra <= chars:
            current = extra
            continue
        lines += 1
        current = len(word)
    return max(0.42, lines * (size / 72 * 1.4) + 0.1)

def _clear_pack_mat(path: Path | None) -> Path | None:
    """Remove the white studio box around a pack. White printed on the wrapper stays."""
    if path is None or not Path(path).exists():
        return path
    source = Path(path)
    dest = source.with_name(f"{source.stem}-clear3.png")
    try:
        if dest.exists() and dest.stat().st_mtime >= source.stat().st_mtime:
            return dest
        image = Image.open(source).convert("RGBA")
        width, height = image.size
        seeds = (
            (0, 0), (width - 1, 0), (0, height - 1), (width - 1, height - 1),
            (width // 2, 0), (width // 2, height - 1), (0, height // 2), (width - 1, height // 2),
        )
        for point in seeds:
            red, green, blue, alpha = image.getpixel(point)
            if alpha < 20 or (red >= 242 and green >= 242 and blue >= 242):
                # Stay under the bottle's own white. A wider threshold eats the wrapper.
                ImageDraw.floodfill(image, point, (255, 255, 255, 0), thresh=12)
        _drop_edge_hairline(image)
        box = image.getchannel("A").getbbox()
        if box:
            image = image.crop(box)
        image.save(dest)
        return dest
    except Exception:
        return source


def _drop_edge_hairline(image: Image.Image) -> None:
    """Remove a one-pixel gray stroke on the bright side. The shaded side stays.

    The bottle art draws a hard gray pixel where the wrapper is white, and a
    soft shadow on the other side. Clearing every light edge pixel would shave
    the white wrapper. This only drops a pixel when the bottle just inside it
    is already white.
    """
    pixels = image.load()
    width, height = image.size
    clear: list[tuple[int, int]] = []
    for y in range(height):
        for x in range(width):
            red, green, blue, alpha = pixels[x, y]
            if alpha < 16:
                continue
            if min(red, green, blue) < 176 or max(red, green, blue) > 236:
                continue
            if max(red, green, blue) - min(red, green, blue) > 14:
                continue
            inward = None
            if x > 0 and pixels[x - 1, y][3] < 16 and x + 1 < width:
                inward = pixels[x + 1, y]
            elif x + 1 < width and pixels[x + 1, y][3] < 16 and x > 0:
                inward = pixels[x - 1, y]
            elif y > 0 and pixels[x, y - 1][3] < 16 and y + 1 < height:
                inward = pixels[x, y + 1]
            elif y + 1 < height and pixels[x, y + 1][3] < 16 and y > 0:
                inward = pixels[x, y - 1]
            if inward is None or inward[3] < 16:
                continue
            if min(inward[0], inward[1], inward[2]) < 236:
                continue
            clear.append((x, y))
    for x, y in clear:
        pixels[x, y] = (255, 255, 255, 0)


def _content_box(image: Image.Image) -> tuple[int, int, int, int] | None:
    """Where the artwork actually is. Empty margin, transparent or white, does not count."""
    width, height = image.size
    alpha = image.getchannel("A")
    transparent = sum(alpha.histogram()[:24]) / max(width * height, 1)
    if transparent > 0.08:
        return alpha.point(lambda value: 255 if value > 24 else 0).getbbox()
    bands = [band.point(lambda value: 0 if value > 242 else 255) for band in image.convert("RGB").split()]
    mask = bands[0]
    for band in bands[1:]:
        mask = ImageChops.lighter(mask, band)
    return mask.getbbox()


def _element_thumb(path: Path) -> Path:
    """Zoom a small element so the artwork fills the row instead of sitting in empty margin."""
    dest = path.with_name(f"{path.stem}-zoom2.png")
    try:
        if dest.exists() and dest.stat().st_mtime >= path.stat().st_mtime:
            return dest
        image = Image.open(path).convert("RGBA")
        width, height = image.size
        box = _content_box(image) or (0, 0, width, height)
        content_w = max(1, box[2] - box[0])
        content_h = max(1, box[3] - box[1])
        small = (content_w * content_h) / max(width * height, 1) < 0.62
        if small:
            pad_x = content_w * 0.1
            pad_y = content_h * 0.1
            left = max(0, int(box[0] - pad_x))
            top = max(0, int(box[1] - pad_y))
            right = min(width, int(box[2] + pad_x))
            bottom = min(height, int(box[3] + pad_y))
            image = image.crop((left, top, max(left + 1, right), max(top + 1, bottom)))
        else:
            side = min(width, height)
            left = max(0, (width - side) // 2)
            top = max(0, (height - side) // 2)
            image = image.crop((left, top, left + side, top + side))
        resample = getattr(getattr(Image, "Resampling", Image), "LANCZOS")
        longest = max(image.size)
        scale = 420 / max(longest, 1)
        resized = (
            max(1, round(image.size[0] * scale)),
            max(1, round(image.size[1] * scale)),
        )
        image.resize(resized, resample).save(dest)
        return dest
    except Exception:
        return path


def _place_picture(slide, path: Path | None, x, y, max_w, max_h) -> None:
    """Fit a picture in the box. Width and height always keep the file's aspect."""
    if path is None or not Path(path).exists():
        return
    width, height = contain_size(file_aspect(path), max_w, max_h)
    left = x + max(0, (max_w - width) / 2)
    top = y + max(0, (max_h - height) / 2)
    picture = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    picture.line.fill.background()
    element = picture._element
    frame = element.spPr
    border = frame.find(qn("a:ln"))
    if border is not None:
        frame.remove(border)
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)

def _place_fit(slide, path: Path, x, y, max_w, max_h) -> None:
    _place_picture(slide, path, x, y, max_w, max_h)

def _detail_legend(build: Build) -> str:
    parts: list[str] = []
    if build.audience_label:
        parts.append(build.audience_label.rstrip("."))
    if len(build.comparison_segments) >= 2 and build.comparison_title:
        parts.append(build.comparison_title)
    return ". ".join(parts)

def _legend_columns(build: Build, digits: int) -> tuple[list[ComparisonSegment], float, float] | None:
    """Name, then one score column for the full sample, then one per group."""
    segments = list(build.comparison_segments)
    if len(segments) < 2:
        return None
    table_right = 8.72
    left = 0.83
    columns = len(segments) + 1
    avail = table_right - left - 0.12
    name_w = min(3.45, max(2.35, avail - columns * (1.36 if digits < 3 else 1.5)))
    stride = (avail - name_w) / columns
    if stride < 1.02:
        stride = 1.02 if digits < 3 else 1.12
        name_w = avail - columns * stride
    if stride < 0.92 or name_w < 2.05:
        return None
    return segments, name_w, stride

def _place_element_list(slide, build: Build, digits: int = 1, sample: int = 0):
    """Element rows beside the pack. Space separates the rows. Each column shows its base size."""
    elements = build.elements
    count = max(len(elements), 1)
    fitted = _legend_columns(build, digits)
    limit = 6.55
    if fitted:
        segments, name_w, stride = fitted
        spans = _question_spans(segments, stride)
        question_lines = max((item[3] for item in spans if item[2]), default=1)
        answer_lines = _answer_lines(segments, stride)
        header_h = (0.36 if question_lines == 2 else 0.20) + (0.36 if answer_lines == 2 else 0.22) + 0.20
        total_h = 0.46
        start = 1.64 + header_h
        room = 6.58 - start - total_h - 0.10
        row_h = min(0.48, max(0.36, room / count))
        if row_h >= 0.32:
            pill_x = 0.83 + name_w + 0.10
            _comparison_headers(slide, segments, pill_x, start - header_h, stride, header_h, sample)
            _element_column(
                slide, elements, 0.83, name_w, pill_x, pill_x, 0.0,
                start, row_h, digits, segments, stride, sample_column=True,
            )
            _total_row(
                slide, build.total, start + count * row_h + 0.08,
                0.83, name_w, pill_x, digits, segments, stride, build, sample,
            )
            return

    chip_h = _legend_chip_row(slide, build, digits, 1.74, 0.83, 8.1) if len(build.comparison_segments) >= 2 else 0.0
    start = 1.78 + chip_h
    row_h = 0.46
    fits = start + count * row_h + 0.50 <= limit
    if fits:
        _element_column(slide, elements, 0.95, 3.15, 4.20, 5.10, 3.7, start, row_h, digits)
        _total_row(slide, build.total, start + count * row_h + 0.14, 0.95, 3.15, 4.20, digits, sample=sample)
        return

    split = (count + 1) // 2
    columns = [elements[:split], elements[split:]]
    tallest = max(len(column) for column in columns)
    row_h = min(row_h, max(0.32, (limit - start - 0.50) / tallest))
    _element_column(slide, columns[0], 0.95, 1.85, 2.85, 3.60, 1.05, start, row_h, digits)
    if columns[1]:
        _element_column(slide, columns[1], 4.85, 1.85, 6.75, 7.50, 1.05, start, row_h, digits)
    _total_row(slide, build.total, start + tallest * row_h + 0.14, 0.95, 3.15, 4.20, digits, sample=sample)


def _legend_chip_row(slide, build: Build, digits: int, y: float, x: float, width: float) -> float:
    """One chip per legend group when the columns will not fit beside the pack."""
    cursor_x = x
    cursor_y = y
    row_h = 0.52
    for segment in build.comparison_segments:
        label = _clip_lines(_column_title(segment), 1.7, size=8, lines=2)
        if segment.base_size:
            label = f"{label}\n{_respondent_count(segment.base_size)}"
        chip_w = min(2.8, 0.28 + max(1.3, len(label) * 0.055) + 0.78)
        if cursor_x > x and cursor_x + chip_w > x + width:
            cursor_x = x
            cursor_y += row_h
        _text(
            slide, label,
            cursor_x, cursor_y, chip_w - 0.78, row_h,
            size=8, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE,
        )
        _score_pill(
            slide, _segment_total(build, segment.key),
            cursor_x + chip_w - 0.72, cursor_y, row_h,
            size=9, digits=digits,
        )
        cursor_x += chip_w + 0.08
    return cursor_y + row_h - y + 0.06

def _answer_lines(segments: list[ComparisonSegment], stride: float) -> int:
    width = max(0.7, stride - 0.12)
    per_line = max(8, int(width * 72 / (11 * 0.52)))
    for segment in segments:
        if len(_segment_answer(segment)) > per_line:
            return 2
    return 1


def _comparison_headers(
    slide, segments: list[ComparisonSegment], x: float, y: float, stride: float, height: float, sample: int = 0,
) -> None:
    """Full sample first, then each group. Question and answer wrap instead of being cut off."""
    spans = _question_spans(segments, stride)
    question_lines = max((item[3] for item in spans if item[2]), default=1)
    question_h = 0.36 if question_lines == 2 else 0.20
    answer_h = 0.36 if _answer_lines(segments, stride) == 2 else 0.22
    _text(
        slide, "Full sample",
        x, y, stride - 0.08, question_h,
        size=11, bold=True, color=_SEGMENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM,
    )
    answer_y = y + question_h + 0.02
    if sample:
        _text(
            slide, f"(n={sample})",
            x, answer_y + answer_h - 0.02, stride - 0.08, 0.20,
            size=11, color=_MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
    for start, end, text, lines, width in spans:
        if not text:
            continue
        label = _text(
            slide, text,
            x + (start + 1) * stride, y, width, question_h,
            size=11, bold=True, color=_SEGMENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM,
        )
        label.text_frame.word_wrap = lines == 2
    for index, segment in enumerate(segments):
        cell = x + (index + 1) * stride
        box = _text(
            slide, _segment_answer(segment),
            cell, answer_y, stride - 0.10, answer_h,
            size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        box.text_frame.word_wrap = True
        if segment.base_size:
            _text(
                slide, f"(n={segment.base_size})",
                cell, answer_y + answer_h - 0.02, stride - 0.10, 0.20,
                size=11, color=_MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            )

def _column_title(segment: ComparisonSegment) -> str:
    answer = _segment_answer(segment)
    question = _condense_question(segment.question) if segment.question else ""
    if question and answer and question.lower().rstrip("?") != answer.lower():
        return f"{question} {answer}"
    return answer or segment.label

def _detail_footer(build: Build) -> str:
    audience = build.audience_label.rstrip(".")
    lead = f"{audience}. " if audience else ""
    if len(build.comparison_segments) >= 2:
        return (
            f"{lead}A + score is green. A − score is red. "
            "The first score is the full sample. A group column does not add up to it."
        )
    return (
        f"{lead}A + score is green. A − score is red. "
        "Overall adds the element lifts together."
    )

def _element_column(slide, elements, name_x, name_w, pill_x, cat_x, cat_w, top, row_h, digits: int, segments: list[ComparisonSegment] | None = None, stride: float = 0.86, sample_column: bool = False):
    series = segments or []
    name_size = 12 if name_w >= 3.2 else 11
    name_lines = 2 if row_h >= 0.40 else 1
    name_line = RGBColor(0xD5, 0xDC, 0xE5)
    for index, element in enumerate(elements):
        row_top = top + index * row_h
        box_h = max(0.30, row_h - 0.06)
        add_rect(slide, name_x, row_top + 0.03, name_w - 0.08, box_h, WHITE, rounded=True, radius=0.08, line=name_line)
        shown = _clip_lines(element.label, name_w - 0.28, size=name_size, lines=name_lines, readable=True)
        name = _text(
            slide, shown, name_x + 0.1, row_top + 0.03, name_w - 0.26, box_h,
            size=name_size, bold=True, color=_INK, anchor=MSO_ANCHOR.MIDDLE,
        )
        name.text_frame.word_wrap = True
        if sample_column:
            _score_pill(slide, element.lift, pill_x, row_top, row_h, size=13, digits=digits, max_w=0.98, cell_w=stride)
            for series_index, segment in enumerate(series):
                _score_pill(
                    slide,
                    float(element.segment_lifts.get(segment.key, 0.0)),
                    pill_x + (series_index + 1) * stride,
                    row_top,
                    row_h,
                    size=13,
                    digits=digits,
                    max_w=0.98,
                    cell_w=stride,
                )
        elif series:
            for series_index, segment in enumerate(series):
                _score_pill(
                    slide,
                    float(element.segment_lifts.get(segment.key, 0.0)),
                    pill_x + series_index * stride,
                    row_top,
                    row_h,
                    size=13,
                    digits=digits,
                    max_w=0.98,
                    cell_w=stride,
                )
        else:
            _score_pill(slide, element.lift, pill_x, row_top, row_h, size=14, digits=digits)
        if cat_w > 0.2:
            category = _clip_lines(element.category, cat_w - 0.06, size=10, lines=2, readable=True)
            note = _text(slide, category, cat_x, row_top, cat_w, row_h, size=10, color=_MUTED, anchor=MSO_ANCHOR.MIDDLE)
            note.text_frame.word_wrap = True

def _total_row(
    slide,
    total: float,
    top: float,
    label_x: float,
    label_w: float,
    pill_x: float,
    digits: int,
    segments: list[ComparisonSegment] | None = None,
    stride: float = 0.86,
    build: Build | None = None,
    sample: int = 0,
):
    row_h = 0.42
    label = "Overall" if not sample else f"Overall  (n={sample})"
    add_rect(slide, label_x, top + 0.02, min(label_w, 2.7) - 0.08, row_h - 0.04, _BAR, rounded=True, radius=0.08)
    _text(
        slide, label,
        label_x + 0.1, top + 0.02, min(label_w, 2.7) - 0.24, row_h - 0.04,
        size=12, bold=True, color=_TITLE, anchor=MSO_ANCHOR.MIDDLE,
    )
    if segments:
        _score_pill(slide, total, pill_x, top, row_h, size=13, digits=digits, max_w=0.98, cell_w=stride)
        for index, segment in enumerate(segments):
            value = _segment_total(build, segment.key) if build is not None else 0.0
            _score_pill(
                slide, value,
                pill_x + (index + 1) * stride, top, row_h,
                size=13, digits=digits, max_w=0.98, cell_w=stride,
            )
    else:
        _score_pill(slide, total, pill_x, top, row_h, size=14, digits=digits)

def _score_pill(slide, value: float, x: float, row_top: float, row_h: float, size: float, digits: int = 1, fill: RGBColor | None = None, max_w: float | None = None, cell_w: float | None = None, ink: RGBColor | None = None):
    """A solid block. Green carries a plus, red carries a minus, zero stays neutral."""
    pill_h = min(0.34, max(0.26, row_h - 0.08))
    pill_w = 1.05 if digits >= 3 else 0.92
    if max_w is not None:
        pill_w = min(pill_w, max(0.62, max_w))
    if cell_w is not None:
        pill_w = min(pill_w, max(0.62, cell_w - 0.12))
        x = x + max(0.0, (cell_w - pill_w) / 2)
    _tone_fill, tone_ink, bar = _score_tone(value, digits)
    cutoff = 0.0005 if digits >= 3 else 0.05
    if fill is None:
        fill = _MID_FILL if abs(value) <= cutoff else bar
    if ink is None:
        ink = tone_ink if abs(value) <= cutoff else WHITE
    top = row_top + (row_h - pill_h) / 2
    add_rect(slide, x, top, pill_w, pill_h, fill, rounded=True, radius=0.12)
    _text(
        slide, _lift(value, digits),
        x, top, pill_w, pill_h,
        size=size, bold=True, color=ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

def _variant_label(builds: list[Build], build: Build) -> str:
    if len(builds) < 2:
        return build.elements[0].label if build.elements else build.name
    shared = set(builds[0].elements[i].name for i in range(len(builds[0].elements)))
    for other in builds[1:]:
        names = {element.name for element in other.elements}
        shared &= names
    parts = [element.label for element in build.elements if element.name not in shared]
    return " · ".join(parts) if parts else (build.elements[0].label if build.elements else build.name)

def _render(study: StudyInfo, build: Build) -> Path | None:
    layers: list[Path] = []
    if build.background_path:
        layers.append(build.background_path)
    layers.extend(element.image_path for element in build.elements if element.image_path)
    if not layers:
        return None
    cache_root = layers[0].parent
    key = "-".join(re.sub(r"[^A-Za-z0-9]+", "_", part)[:40] for part in build.signature())
    return composite_pack(layers, cache_root.parent / "composites" / f"{key}.png")

def _sample_line(study: StudyInfo, group: Group, kind: str, short: bool = False) -> str:
    known = study.segment_sizes.get(group.segment)
    digits = _digits(kind)
    places = "three decimals" if digits == 3 else "one decimal"
    base_value = _base_for(study, kind)
    base = f" Base appeal {base_value:.{digits}f}." if base_value is not None else ""
    if short:
        if known:
            return f"n={known} {group.segment} respondents."
        return "Saved segment. Scores are from the overall model."
    if known:
        who = study.country if group.segment.lower() == "overall" else group.segment
        sample = f"n={known} {who} respondents."
    else:
        sample = f"Saved segment: {group.segment}. Scores are from the overall model."
    return (
        f"{sample}{base} "
        f"{study.model_name} model: estimated total combination lift above baseline. Values rounded to {places}."
    )

def _lift(value: float, digits: int = 1) -> str:
    cutoff = 0.0005 if digits >= 3 else 0.05
    if abs(value) <= cutoff:
        return f"{0:.{digits}f}"
    return f"{value:+.{digits}f}"

def _tone(value: float) -> RGBColor:
    if value > 0.05:
        return GREEN
    if value < -0.05:
        return RED
    return GRAY
