"""Horizontal lift charts styled to the appeal report."""

from __future__ import annotations

import math

from lxml import etree
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.services.reportgen.theme import FONT, GRAY, NAVY, add_rect, add_text, bar_color, fmt_lift, lift_color

# Room kept clear of the bar tip so the score never sits on the bar or the name.
_LABEL_ROOM = 0.62


def add_lift_chart(slide, elements, x, y, w, h, name_w: float = 2.05, pictures=None):
    """Drawn bars. Labels stay on the left; the score sits just past the bar tip.

    ``pictures`` maps an element code to an image path. When it is set, that
    image replaces the name. An Excel bar chart places category names on the
    zero line. When most scores are negative, that line sits inside the plot
    and the names land on the bars. Shapes keep the same layout in PowerPoint
    and Keynote.
    """
    ordered = sorted(elements, key=lambda element: element.lift, reverse=True)
    if not ordered:
        return None

    axis_h = 0.30
    plot_x = x + name_w
    plot_w = max(w - name_w - 0.06, 2.4)
    plot_y = y
    plot_h = max(h - axis_h, 1.2)
    axis_min, axis_max = _axis_bounds([element.lift for element in ordered], plot_w)
    span = axis_max - axis_min or 1

    def at(value: float) -> float:
        return plot_x + (value - axis_min) / span * plot_w

    grid = RGBColor(0xE6, 0xEA, 0xEE)
    axis_ink = RGBColor(0xB9, 0xC1, 0xCA)
    for tick, gx in _ticks(axis_min, axis_max, plot_x, plot_w):
        add_rect(slide, gx, plot_y, 0.01, plot_h, grid)
        label_left = min(max(gx - 0.32, plot_x - 0.08), plot_x + plot_w - 0.56)
        _fit_text(
            slide, _tick_label(tick), label_left, plot_y + plot_h + 0.02, 0.64, 0.24,
            size=10, bold=False, color=GRAY, align=PP_ALIGN.CENTER,
        )
    zero = at(0)
    add_rect(slide, zero, plot_y, 0.015, plot_h, axis_ink)
    add_rect(slide, plot_x, plot_y + plot_h, plot_w, 0.012, axis_ink)

    row_h = plot_h / len(ordered)
    bar_h = min(0.24, row_h * 0.42)
    for index, element in enumerate(ordered):
        row_top = plot_y + index * row_h
        bar_top = row_top + (row_h - bar_h) / 2
        end = at(element.lift)
        left = min(zero, end)
        width = max(abs(end - zero), 0.035)
        add_rect(slide, left, bar_top, width, bar_h, bar_color(element.lift))
        _score(slide, element.lift, end, row_top, row_h, plot_x, plot_x + plot_w)
        picture = pictures.get(element.code) if pictures else None
        if picture:
            _label_picture(slide, picture, x, row_top, name_w - 0.10, row_h)
        else:
            size = _name_size(element.label)
            _fit_text(
                slide, wrap_label(element.label, name_w - 0.18, size), x, row_top, name_w - 0.14, row_h,
                size=size, bold=False, color=NAVY, align=PP_ALIGN.RIGHT,
            )
    return None


def add_column_chart(slide, pairs: list[tuple[str, float]], x, y, w, h, series_name: str = "Count"):
    data = CategoryChartData()
    data.categories = [label for label, _value in pairs]
    data.add_series(series_name, [value for _label, value in pairs])
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data
    )
    chart = chart_shape.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.gap_width = 60
    labels = plot.data_labels
    labels.font.size = Pt(11)
    labels.font.bold = True
    labels.font.name = FONT
    labels.font.color.rgb = NAVY
    labels.number_format = "0"
    labels.position = XL_LABEL_POSITION.OUTSIDE_END

    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = NAVY

    chart.category_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.color.rgb = NAVY
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.visible = False
    chart.value_axis.format.line.fill.background()
    _clear_chart_border(chart)
    return chart


def _score(slide, lift, end, row_top, row_h, plot_left, plot_right) -> None:
    label_w = 0.56
    gap = 0.07
    if lift < 0:
        left = end - gap - label_w
        align = PP_ALIGN.RIGHT
    else:
        left = end + gap
        align = PP_ALIGN.LEFT
    left = min(max(left, plot_left), plot_right - label_w)
    _fit_text(
        slide, fmt_lift(lift), left, row_top, label_w, row_h,
        size=12, bold=True, color=lift_color(lift), align=align,
    )


def _label_picture(slide, path, x, y, w, h) -> None:
    """Fit the element in the label column, right against the plot."""
    with Image.open(path) as image:
        aspect = image.width / image.height if image.height else 1
    max_h = h * 0.86
    width = max_h * aspect
    height = max_h
    if width > w:
        width = w
        height = width / aspect if aspect else max_h
    left = x + w - width
    top = y + (h - height) / 2
    edge = RGBColor(0xE2, 0xE8, 0xEE)
    pad = 0.01
    add_rect(slide, left - pad, top - pad, width + pad * 2, height + pad * 2, RGBColor(0xFF, 0xFF, 0xFF), line=edge)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def _fit_text(slide, text, x, y, w, h, size, bold, color, align) -> None:
    box = add_text(
        slide, text, x, y, w, h,
        size=size, bold=bold, color=color, align=align, anchor=MSO_ANCHOR.MIDDLE,
    )
    frame = box.text_frame
    frame.word_wrap = False
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    for paragraph in frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)


def wrap_label(label: str, width_in: float, size_pt: int = 11, max_lines: int = 2) -> str:
    """Break a name onto a fixed number of lines so it stays inside its row."""
    chars = max(8, int(width_in * 11 * (11 / max(size_pt, 1))))
    words = str(label).split()
    if not words:
        return ""
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if len(trial) <= chars:
            current = trial
            continue
        if current:
            lines.append(current)
        current = word
    if current:
        lines.append(current)
    if len(lines) <= max_lines:
        return "\n".join(lines)
    kept = lines[:max_lines]
    kept[-1] = kept[-1].rstrip(" .,;:—-") + "…"
    return "\n".join(kept)


def _name_size(label: str) -> int:
    if len(label) > 36:
        return 10
    if len(label) > 18:
        return 11
    return 12


def _axis_bounds(values, plot_w: float) -> tuple[float, float]:
    lo = min(min(values), 0)
    hi = max(max(values), 0)
    span = max(hi - lo, 2)
    extra = _LABEL_ROOM / max(plot_w - _LABEL_ROOM, 1) * span
    axis_min = lo - extra if lo < 0 else 0
    axis_max = hi + extra if hi > 0 else 0
    step = _nice_step(max(axis_max - axis_min, 2) / 4)
    if lo < 0:
        axis_min = math.floor(axis_min / step) * step
    if hi > 0:
        axis_max = math.ceil(max(axis_max, step) / step) * step
    else:
        # Keep zero on the right edge, with the label room already in axis_min.
        axis_max = 0
    if axis_max <= axis_min:
        axis_max = axis_min + step
    return axis_min, axis_max


def _ticks(axis_min: float, axis_max: float, plot_x: float, plot_w: float):
    step = _nice_step(max(axis_max - axis_min, 1) / 4)
    span = axis_max - axis_min or 1
    values = []
    cursor = math.ceil((axis_min - 1e-6) / step) * step
    while cursor <= axis_max + 1e-6:
        values.append(round(cursor, 6))
        cursor += step
    if axis_min <= 0 <= axis_max and not any(abs(value) < 1e-6 for value in values):
        values.append(0.0)
        values.sort()
    kept = []
    for value in values:
        gx = plot_x + (value - axis_min) / span * plot_w
        if kept and abs(gx - kept[-1][1]) < 0.48 and abs(value) > 1e-6:
            continue
        kept.append((value, gx))
    return kept


def _tick_label(value: float) -> str:
    if abs(value) < 1e-6:
        return "0"
    number = int(round(value))
    return str(number)


def _nice_step(rough: float) -> float:
    if rough <= 0:
        return 1
    magnitude = 10 ** math.floor(math.log10(rough))
    residual = rough / magnitude
    if residual <= 1:
        nice = 1
    elif residual <= 2:
        nice = 2
    elif residual <= 5:
        nice = 5
    else:
        nice = 10
    return nice * magnitude


def _set_plot_layout(chart, x: float, y: float, w: float, h: float) -> None:
    plot_area = chart._element.find(qn("c:chart")).find(qn("c:plotArea"))
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        plot_area.insert(0, layout)
    for child in list(layout):
        layout.remove(child)
    manual = etree.SubElement(layout, qn("c:manualLayout"))

    def add(tag: str, value: str) -> None:
        node = etree.SubElement(manual, qn(f"c:{tag}"))
        node.set("val", value)

    add("xMode", "edge")
    add("yMode", "edge")
    add("x", f"{x:.3f}")
    add("y", f"{y:.3f}")
    add("w", f"{w:.3f}")
    add("h", f"{h:.3f}")


def _clear_chart_border(chart) -> None:
    chart_element = chart._element
    plot_area = chart_element.find(qn("c:chart")).find(qn("c:plotArea"))
    sp_pr = plot_area.find(qn("c:spPr"))
    if sp_pr is None:
        return
    for child in list(sp_pr):
        if child.tag == qn("a:ln"):
            sp_pr.remove(child)
